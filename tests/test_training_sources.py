"""Model_Training_Datasets sheet: the classifier's training sources CSV."""

import json

import pytest

CSV_TEXT = (
    "model_version,source_id,class,dataset_title,url,license,training_policy,role,"
    "student_train_n,student_test_n,synthetic,notes\n"
    "m1,zenodo-1,CFP,Fundus A,https://zenodo.org/records/1,CC-BY-4.0,allowed,train,350,0,no,\n"
    "m1,synth,FAF,Synthetic FAF,https://doi.org/10.5522/x,CC0-1.0,allowed,train_synthetic,200,0,yes,\n"
    "m1,bank,FAF,Image bank FAF,https://example.org/,Copyright,eval_only,eval_only,0,155,no,never train\n"
    "m1,nc-1,OCTA,NC OCTA,https://zenodo.org/records/2,CC-BY-NC-4.0,eval_only_nc,train,350,0,no,\n"
    ",,,,,,,,,,,\n"
    "m2,zenodo-3,OCT,New OCT,https://zenodo.org/records/3,CC-BY-4.0,allowed,train,100,20,no,retrain\n"
)


def _results(tmp_path):
    out = tmp_path / "res"
    out.mkdir()
    res = out / "survey_results.jsonl"
    res.write_text(json.dumps({"record_id": "1", "status": "ok"}) + "\n", encoding="utf-8")
    return out, res


def test_training_sources_sheet_and_readme(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    from envision_eye_actionable.survey.excel import TRAINING_SHEET, build_workbook
    out, res = _results(tmp_path)
    csv_path = tmp_path / "training_sources.csv"
    csv_path.write_text(CSV_TEXT, encoding="utf-8")
    stats = build_workbook(res, out / "w.xlsx", training_sources=csv_path)
    assert stats["training_sources"] == 5                     # the blank line is skipped
    wb = load_workbook(out / "w.xlsx")
    assert wb.sheetnames[-1] == TRAINING_SHEET
    ws = wb[TRAINING_SHEET]
    head = [c.value for c in ws[1]]
    assert head[:3] == ["model_version", "source_id", "class"] and "notes" in head
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 5 and [r[0] for r in rows] == ["m1"] * 4 + ["m2"]
    n_col = head.index("student_train_n")
    assert rows[0][n_col] == 350 and isinstance(rows[0][n_col], int)
    assert ws.cell(row=2, column=head.index("url") + 1).hyperlink.target == "https://zenodo.org/records/1"
    readme = {r[0].value: r[1].value for r in wb["README"].iter_rows()}
    note = readme["Model training datasets"]
    assert TRAINING_SHEET in note
    assert "m1: 4 source rows over 3 classes (CFP, FAF, OCTA), 900 training images, 155 test images" in note
    assert "synthetic: synth" in note and "evaluation only (never trained on): bank" in note
    assert "not cleared by license policy: nc-1" in note
    assert "m2: 1 source rows" in note


def test_training_sources_optional(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    from envision_eye_actionable.survey.excel import TRAINING_SHEET, build_workbook
    out, res = _results(tmp_path)
    stats = build_workbook(res, out / "w.xlsx")
    wb = load_workbook(out / "w.xlsx")
    assert TRAINING_SHEET not in wb.sheetnames and stats["training_sources"] is None
    readme = {r[0].value: r[1].value for r in wb["README"].iter_rows()}
    assert readme["Model training datasets"].startswith("not included")


def test_training_sources_needs_model_version(tmp_path):
    from envision_eye_actionable.survey.excel import load_training_sources
    p = tmp_path / "t.csv"
    p.write_text("source_id,class\nx,CFP\n", encoding="utf-8")
    with pytest.raises(ValueError, match="model_version"):
        load_training_sources(p)


def test_cli_training_sources_errors(tmp_path):
    from envision_eye_actionable.survey.cli import main
    out, _ = _results(tmp_path)
    with pytest.raises(SystemExit):
        main(["excel", "--results-dir", str(out), "--out", str(out / "w.xlsx"),
              "--training-sources", str(tmp_path / "missing.csv")])
    bad = tmp_path / "bad.csv"
    bad.write_text("source_id\nx\n", encoding="utf-8")
    with pytest.raises(SystemExit):
        main(["excel", "--results-dir", str(out), "--out", str(out / "w.xlsx"), "--training-sources", str(bad)])
    assert not (out / "w.xlsx").exists()


def test_cli_training_sources_builds_sheet(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    from envision_eye_actionable.survey.cli import main
    from envision_eye_actionable.survey.excel import TRAINING_SHEET
    out, _ = _results(tmp_path)
    csv_path = tmp_path / "training_sources.csv"
    csv_path.write_text(CSV_TEXT, encoding="utf-8")
    assert main(["excel", "--results-dir", str(out), "--out", str(out / "w.xlsx"),
                 "--training-sources", str(csv_path)]) == 0
    assert TRAINING_SHEET in load_workbook(out / "w.xlsx").sheetnames

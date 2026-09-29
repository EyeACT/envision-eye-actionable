"""Review list of likely false positives in the workbook: ok records whose
only eye classes are OCTA and/or PSC, with SetFit NEGATIVE or an eye image
fraction under 0.5."""

import json

import pytest


def _row(rid, present, label="EYE_IMAGING", ef=0.9, status="ok"):
    return {"record_id": rid, "status": status, "present_eye_classes": present, "setfit_label": label,
            "eye_image_fraction": ef, "n_classified": 10, "title": f"t{rid}", "url": f"https://zenodo.org/records/{rid}"}


def test_likely_false_positive_rule():
    from envision_eye_actionable.survey.excel import likely_false_positive as fp
    assert fp(_row("1", "OCTA", label="NEGATIVE")) == "only OCTA; discovery SetFit label NEGATIVE"
    assert fp(_row("1", "OCTA, PSC", ef=0.2)) == "only OCTA + PSC; eye image fraction 0.20 < 0.5"
    assert fp(_row("1", "PSC", label="NEGATIVE", ef=0.49)).count(";") == 2
    assert fp(_row("1", "OCTA")) == ""                              # eye metadata and most images eye
    assert fp(_row("1", "OCTA", ef=0.5)) == ""                      # 0.5 is not under 0.5
    assert fp(_row("1", "OCTA, OCT", label="NEGATIVE")) == ""       # another eye class too
    assert fp(_row("1", "", label="NEGATIVE", ef=0.0)) == ""        # no eye class
    assert fp(_row("1", "OCTA", label="NEGATIVE", status="no_eye_images")) == ""
    assert fp(_row("1", "PSC", ef=0.1, status="ok_partial_download")) != ""
    assert fp(_row("1", "OCTA", ef=None)) == ""


def test_workbook_review_sheet_lists_only_likely_false_positives(tmp_path):
    pytest.importorskip("openpyxl")
    from envision_eye_actionable.survey.excel import REVIEW_FP_SHEET, build_workbook
    from openpyxl import load_workbook
    out = tmp_path / "res"
    out.mkdir()
    rows = [_row("11", "OCTA", label="NEGATIVE"), _row("12", "CFP", label="NEGATIVE", ef=0.1),
            _row("13", "PSC", ef=0.3), _row("14", "OCTA"), _row("15", "OCTA", label="NEGATIVE", status="no_eye_images")]
    res = out / "survey_results.jsonl"
    res.write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
    stats = build_workbook(res, out / "w.xlsx")
    assert stats["review_likely_fp"] == 2
    wb = load_workbook(out / "w.xlsx")
    ws = wb[REVIEW_FP_SHEET]
    got = [(r[0].value, r[4].value) for r in ws.iter_rows(min_row=2)]
    assert got == [("11", "only OCTA; discovery SetFit label NEGATIVE"),
                   ("13", "only PSC; eye image fraction 0.30 < 0.5")]
    head = [c.value for c in wb["Records"][1]]
    col = head.index("Likely false positive (review list): why")
    by_id = {r[0].value: r[col].value for r in wb["Records"].iter_rows(min_row=2)}
    assert by_id["11"] and by_id["13"] and not by_id["12"] and not by_id["14"] and not by_id["15"]
    readme = {r[0].value: r[1].value for r in wb["README"].iter_rows()}
    assert str(readme[f"Records on the {REVIEW_FP_SHEET} review list"]).startswith("2 ")

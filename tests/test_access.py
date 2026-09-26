"""Offline tests of enrich-access (survey/access.py) and its workbook supplement."""

from __future__ import annotations

import hashlib
import json

import pytest

from envision_eye_actionable.survey import access, cmds, excel


def _invenio(rid, files="restricted", embargo=None, settings=True, rights=("cc-by-4.0",), title="Retinal OCT set",
             description="<p>Fundus photographs.</p>"):
    doc = {
        "id": rid,
        "access": {"record": "public", "files": files, "status": "embargoed" if embargo else files,
                   "embargo": embargo or {"active": False, "reason": None}},
        "parent": {"id": "1", "access": {"owned_by": {"user": "1"}}},
        "metadata": {"title": title, "resource_type": {"id": "dataset"},
                     "rights": [{"id": r, "title": {"en": r.upper()}} for r in rights],
                     "description": description, "publication_date": "2024-01-01"},
        "links": {"access_request": f"https://zenodo.org/api/records/{rid}/access/request"},
        "files": {"enabled": True},
    }
    if settings:
        doc["parent"]["access"]["settings"] = {"allow_user_requests": True, "allow_guest_requests": False,
                                               "accept_conditions_text": "<p>Academic use only.</p>",
                                               "secret_link_expiration": 0}
    return doc


def _datacite():
    return {"doi": "10.5281/zenodo.11", "titles": [{"title": "Retinal OCT set"}], "publicationYear": "2024",
            "creators": [{"name": "Doe, Jane", "nameType": "Personal"}], "publisher": {"name": "Zenodo"}}


class _Resp:
    def __init__(self, doc):
        self.doc = doc

    def json(self):
        return self.doc

    def close(self):
        pass


class FakeClient:
    """Stands in for ZenodoClient: InvenioRDM docs by id, a request counter."""

    def __init__(self, docs, legacy=None, datacite=None):
        self.docs, self._legacy, self._datacite = docs, legacy or {}, datacite or {}
        self.n_requests = 0
        self.urls = []

    def get(self, url, headers=None, timeout=None):
        self.n_requests += 1
        self.urls.append((url, (headers or {}).get("Accept")))
        rid = url.rsplit("/", 1)[1]
        if rid not in self.docs:
            raise access.RemoteError("HTTP 410", 410)
        return _Resp(self.docs[rid])

    def legacy(self, rid):
        self.n_requests += 1
        return self._legacy.get(rid)

    def datacite(self, rid):
        self.n_requests += 1
        return self._datacite.get(rid)


def test_parse_access_reads_what_zenodo_exposes():
    a = access.parse_access(_invenio("11"))
    assert a["access_files"] == "restricted" and a["access_record"] == "public"
    assert a["access_files_public_now"] is False
    assert a["access_allow_user_requests"] is True and a["access_allow_guest_requests"] is False
    assert a["access_accept_conditions_text"] == "Academic use only."
    assert a["access_request_url"] == "https://zenodo.org/records/11"
    assert a["access_request_api"].endswith("/records/11/access/request")
    assert a["access_license_ids"] == "cc-by-4.0" and a["access_license_use"] == "train"
    assert a["access_resource_type"] == "dataset"
    assert "accepts access requests" in a["access_request_note"]
    # settings not exposed: None, never guessed
    b = access.parse_access(_invenio("12", settings=False))
    assert b["access_settings_exposed"] is False and b["access_allow_user_requests"] is None
    assert "does not expose" in b["access_request_note"]
    # an active embargo is never public now, even with files "public"
    c = access.parse_access(_invenio("13", files="public",
                                     embargo={"active": True, "until": "2027-01-01", "reason": "paper"}))
    assert c["access_files_public_now"] is False and c["access_embargo_until"] == "2027-01-01"
    assert "embargoed until 2027-01-01" in c["access_request_note"]
    d = access.parse_access(_invenio("14", files="public"))
    assert d["access_files_public_now"] is True and "no request needed" in d["access_request_note"]


@pytest.mark.parametrize("ids,use", [(["cc-by-4.0"], "train"), (["cc0-1.0"], "train"), (["mit"], "train"),
                                     (["cc-by-sa-4.0"], "train"), (["cc-by-nc-4.0"], "evaluate_only_nc_nd"),
                                     (["cc-by-nd-4.0"], "evaluate_only_nc_nd"),
                                     (["cc-by-4.0", "cc-by-nc-sa-4.0"], "evaluate_only_nc_nd"),
                                     ([], "no_license"), (["notspecified"], "no_license"),
                                     (["gpl-3.0-or-later"], "check_license"), (["other-closed"], "check_license")])
def test_license_use(ids, use):
    assert access.license_use(ids) == use


def test_keyword_terms():
    assert access.keyword_terms("Industrial OCT of ceramics") == ["OCT"]
    assert access.keyword_terms("Soil maps of Bavaria", ["soil"]) == []
    assert access.keyword_terms("Retinal fundus images", ["AMD", "glaucoma"]) == ["retinal", "fundus", "glaucoma",
                                                                                 "AMD"]
    assert access.keyword_terms("Octopus doctrine") == []          # acronyms need a word boundary


def test_select_targets_uses_scrape_and_rows():
    scrape = [{"source_id": "1", "access_type": "restricted"}, {"source_id": "2", "access_type": "open"},
              {"source_id": "3", "access_type": "open"}, {"source_id": "4", "access_type": "embargoed"},
              {"source_id": "5", "access_type": "open"}]
    rows = {"2": {"status": "restricted"}, "3": {"status": "no_image_files", "access_right": "closed"},
            "5": {"status": "ok", "access_right": "open"}, "9": {"status": "restricted"}}
    assert access.select_targets(scrape, rows) == ["1", "2", "3", "4", "9"]
    assert access.select_targets(scrape, rows, ids=["4", "5"]) == ["4"]


def test_out_dir_never_overlaps_the_survey(tmp_path):
    survey = tmp_path / "out"
    survey.mkdir()
    (survey / "survey_results.jsonl").write_text("", encoding="utf-8")
    for bad in (survey, survey / "access", tmp_path):
        with pytest.raises(ValueError):
            access.check_out_dir(bad, survey)
    other = tmp_path / "x"
    other.mkdir()
    (other / "survey_results.jsonl").write_text("", encoding="utf-8")
    with pytest.raises(ValueError):
        access.check_out_dir(other, None)
    access.check_out_dir(tmp_path / "access", survey)       # a sibling is fine


def _setup(tmp_path):
    scrape = [
        {"source_id": "11", "title": "Retinal OCT set", "access_type": "restricted", "label": "EYE_IMAGING",
         "prob_eye_imaging": 0.97, "license": "cc-by-4.0", "keywords": ["retina"]},
        {"source_id": "12", "title": "Soil maps", "access_type": "restricted", "label": "NEGATIVE",
         "prob_eye_imaging": 0.01, "license": None, "keywords": []},
        {"source_id": "13", "title": "Fundus embargo", "access_type": "embargoed", "label": "EYE_IMAGING",
         "prob_eye_imaging": 0.8, "license": "cc-by-nc-4.0", "keywords": []},
        {"source_id": "14", "title": "Now open", "access_type": "restricted", "label": "EYE_IMAGING",
         "prob_eye_imaging": 0.9, "license": "cc-by-4.0", "keywords": []},
        {"source_id": "15", "title": "Gone", "access_type": "restricted", "label": "NEGATIVE",
         "prob_eye_imaging": 0.2},
        {"source_id": "20", "title": "Open", "access_type": "open", "label": "EYE_IMAGING", "prob_eye_imaging": 0.99},
    ]
    sp = tmp_path / "scrape.json"
    sp.write_text(json.dumps(scrape), encoding="utf-8")
    survey = tmp_path / "out"
    (survey / "cache" / "legacy").mkdir(parents=True)
    (survey / "cache" / "datacite").mkdir(parents=True)
    rows = [{"record_id": "11", "status": "restricted", "access_right": "restricted", "title": "Retinal OCT set"},
            {"record_id": "14", "status": "ok", "access_right": "open", "n_CFP": 5, "present_eye_classes": "CFP",
             "cmds_dir": "cmds/14"},
            {"record_id": "20", "status": "ok", "access_right": "open"}]
    res = survey / "survey_results.jsonl"
    res.write_text("".join(json.dumps(r) + "\n" for r in rows), encoding="utf-8")
    (survey / "cache" / "legacy" / "11.json").write_text(json.dumps(
        {"metadata": {"access_right": "restricted", "title": "Retinal OCT set"}, "files": []}), encoding="utf-8")
    (survey / "cache" / "datacite" / "11.json").write_text(json.dumps(_datacite()), encoding="utf-8")
    docs = {"11": _invenio("11"), "12": _invenio("12", settings=False, rights=(), title="Soil maps", description="Soil maps of Bavaria"),
            "13": _invenio("13", files="restricted", rights=("cc-by-nc-4.0",), title="Fundus embargo",
                           embargo={"active": True, "until": "2027-03-01", "reason": None}),
            "14": _invenio("14", files="public", title="Now open")}
    client = FakeClient(docs, legacy={"13": {"metadata": {"access_right": "embargoed", "embargo_date": "2027-03-01"}}},
                        datacite={"13": _datacite()})
    return sp, survey, client


def _sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def test_enricher_end_to_end(tmp_path):
    pytest.importorskip("jsonschema")
    sp, survey, client = _setup(tmp_path)
    before = _sha(survey / "survey_results.jsonl")
    out = tmp_path / "access"
    e = access.AccessEnricher(sp, survey, out, client=client)
    stats = e.run()
    assert _sha(survey / "survey_results.jsonl") == before           # the survey's results are only read
    assert not (survey / "cache" / "invenio").exists()               # nor its cache written
    assert stats["targets"] == 5 and stats["written"] == 5 and stats["fetch_failed"] == 1
    assert all(a == access.INVENIO_ACCEPT for _, a in client.urls)
    rows = {r["record_id"]: r for r in map(json.loads, (out / access.SUPPLEMENT_NAME).read_text().splitlines())}
    r11 = rows["11"]
    assert r11["access_eye_relevant"] and r11["access_keyword_match"] and r11["access_setfit_prob"] == 0.97
    assert r11["access_metadata_sources"] == "legacy cache, datacite cache"     # from the survey's cache
    dd = json.loads((out / r11["cmds_dir"] / "dataset_description.json").read_text(encoding="utf-8"))
    text = dd["datasetDeIdentLevel"]["deIdentDetails"]
    assert "publicly released on zenodo as image files" not in text.lower()
    assert "restricted by the depositor" in text.lower() and "de-identification is not reported" in text
    assert "accepts access requests" in dd["accessDetails"]["description"]
    assert dd["accessType"] == "CaseByCaseDownload"
    assert r11["dd_valid"] and r11["dsd_valid"] and cmds.validate(dd, "dataset_description") == []
    # embargoed: the end date reaches the de-identification and access texts
    dd13 = json.loads((out / rows["13"]["cmds_dir"] / "dataset_description.json").read_text(encoding="utf-8"))
    assert "until 2027-03-01" in dd13["datasetDeIdentLevel"]["deIdentDetails"]
    assert "until 2027-03-01" in dd13["accessDetails"]["description"]
    assert rows["13"]["access_license_use"] == "evaluate_only_nc_nd"
    # files public now: access facts only, the survey's own CMDS stays
    assert rows["14"]["access_files_public_now"] and "cmds_dir" not in rows["14"]
    # not eye relevant, no license on record
    assert not rows["12"]["access_eye_relevant"] and rows["12"]["access_license_use"] == "no_license"
    assert rows["15"]["access_fetch_ok"] is False and "410" in rows["15"]["access_error"]
    assert "20" not in rows
    summary = json.loads((out / access.SUMMARY_NAME).read_text())
    assert summary["state"] == "done" and summary["counts"]["eye_relevant_not_public"] == 2
    # resumable: a second run fetches only the failed record again
    n = client.n_requests
    stats2 = access.AccessEnricher(sp, survey, out, client=client).run()
    assert stats2["todo"] == 1 and client.n_requests == n + 1


def test_workbook_supplement_sets_access_fields_only(tmp_path):
    pytest.importorskip("openpyxl")
    pytest.importorskip("jsonschema")
    from openpyxl import load_workbook
    sp, survey, client = _setup(tmp_path)
    out = tmp_path / "access"
    access.AccessEnricher(sp, survey, out, client=client).run()
    # a supplement row that tries to set survey fields: they must be ignored
    with open(out / access.SUPPLEMENT_NAME, "a", encoding="utf-8") as fp:
        fp.write(json.dumps({"record_id": "14", "status": "restricted", "n_CFP": 99, "present_eye_classes": "",
                             "access_files": "public", "access_files_public_now": True,
                             "access_eye_relevant": True}) + "\n")
    src, sup, empty = excel.split_results_dirs([survey, out])
    assert src == [survey / "survey_results.jsonl"] and sup == [out / access.SUPPLEMENT_NAME] and not empty
    stats = excel.build_workbook(src, tmp_path / "w.xlsx", supplements=sup)
    assert stats["supplement_rows"] == 5 and stats["supplement_rows_matched"] == 2
    assert stats["access_requests"] == 2
    wb = load_workbook(tmp_path / "w.xlsx")
    ws = wb["Records"]
    head = [c.value for c in ws[1]]
    rows = {str(r[0]): dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["14"]["Survey status"] == "ok" and rows["14"]["n CFP"] == 5          # survey fields kept
    assert rows["11"]["Survey status"] == "restricted"
    assert rows["11"]["Zenodo files access (enrich-access)"] == "restricted"
    assert rows["11"]["Owner accepts access requests (enrich-access; empty: not exposed)"] is True
    assert rows["11"]["CMDS JSON folder (results dir + cmds/<id>)"].endswith("access/cmds/11")
    ar = wb["Access_Requests"]
    ah = [c.value for c in ar[1]]
    arows = [dict(zip(ah, r)) for r in ar.iter_rows(min_row=2, values_only=True)]
    assert [r["Zenodo record id"] for r in arows] == ["11", "13"]                # by SetFit prob, public left out
    assert arows[1]["Embargo until"] == "2027-03-01"
    assert arows[0]["Request access on (landing page)"] == "https://zenodo.org/records/11"
    cj = wb["CMDS_JSON"]
    cjh = [c.value for c in cj[1]]
    cjrows = {str(r[0]): dict(zip(cjh, r)) for r in cj.iter_rows(min_row=2, values_only=True)}
    assert "restricted by the depositor" in cjrows["11"]["dataset_description.json"].lower()


def test_cli_guards(tmp_path):
    from envision_eye_actionable.survey.cli import main
    sp, survey, _ = _setup(tmp_path)
    with pytest.raises(SystemExit):          # the survey's own out dir
        main(["enrich-access", "--scrape", str(sp), "--survey-out-dir", str(survey), "--out-dir", str(survey),
              "--offline"])
    with pytest.raises(SystemExit):          # supplements alone make no workbook
        (tmp_path / "a").mkdir()
        (tmp_path / "a" / access.SUPPLEMENT_NAME).write_text("", encoding="utf-8")
        main(["excel", "--results-dir", str(tmp_path / "a"), "--out", str(tmp_path / "w.xlsx")])
    with pytest.raises(SystemExit):          # a dir with neither file
        (tmp_path / "b").mkdir()
        main(["excel", "--results-dir", str(survey), "--results-dir", str(tmp_path / "b"),
              "--out", str(tmp_path / "w.xlsx")])
    assert main(["enrich-access", "--scrape", str(sp), "--survey-out-dir", str(survey),
                 "--out-dir", str(tmp_path / "acc"), "--offline"]) == 0
    rows = (tmp_path / "acc" / access.SUPPLEMENT_NAME).read_text().splitlines()
    assert len(rows) == 5 and all(json.loads(r)["access_error"] == "offline" for r in rows)

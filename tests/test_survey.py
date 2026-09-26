"""Offline tests for the survey subpackage (no network, no ONNX model).

Synthetic images, DICOM files and archives are generated in tmp_path, so the
tests exercise the archive walker, the frame loaders, the DICOM-aligned
aggregation, laterality, weblink cleaning and CMDS schema validity.
"""

from __future__ import annotations

import io
import json
import tarfile
import zipfile

import numpy as np
import pytest
from PIL import Image

from envision_eye_actionable.survey import cmds, dicom_map, weblinks
from envision_eye_actionable.survey.archives import RecordWalker
from envision_eye_actionable.survey.constants import detect_ext, file_kind, wanted_for_download
from envision_eye_actionable.survey.images import content_crop, load_frame, preprocess


def _png_bytes(size=(300, 200), color=(180, 60, 20), mode="RGB", plain=False) -> bytes:
    """A PNG of ``color`` with seeded +-12 noise, so it reads as a photo and
    not as a mask (a flat image is mask-like and gets the MASK label);
    ``plain`` gives the flat image."""
    if plain:
        img = Image.new(mode, size, color if mode == "RGB" else 128)
    else:
        w, h = size
        base = np.array(color if mode == "RGB" else 128, np.int16)
        rng = np.random.default_rng(w * 7 + h * 13 + int(np.sum(base)))
        noise = rng.integers(-12, 13, (h, w, 3) if mode == "RGB" else (h, w))
        img = Image.fromarray(np.clip(base + noise, 0, 255).astype(np.uint8), mode)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _mask_png(size=(64, 48), levels=(0, 255)) -> bytes:
    """A binary (or few-level) label map."""
    w, h = size
    arr = np.full((h, w), levels[0], np.uint8)
    for i, v in enumerate(levels[1:], 1):
        arr[(h * i) // (len(levels) + 1):, (w * i) // (len(levels) + 1):] = v
    buf = io.BytesIO()
    Image.fromarray(arr, "L").save(buf, format="PNG")
    return buf.getvalue()


def _dicom_bytes(frames=3, rows=64, cols=48, **extra) -> bytes:
    pydicom = pytest.importorskip("pydicom")
    from pydicom.dataset import FileDataset, FileMetaDataset
    from pydicom.uid import ExplicitVRLittleEndian, generate_uid

    meta = FileMetaDataset()
    meta.MediaStorageSOPClassUID = "1.2.840.10008.5.1.4.1.1.77.1.5.4"
    meta.MediaStorageSOPInstanceUID = generate_uid()
    meta.TransferSyntaxUID = ExplicitVRLittleEndian
    ds = FileDataset("x.dcm", {}, file_meta=meta, preamble=b"\0" * 128)
    ds.SOPClassUID = meta.MediaStorageSOPClassUID
    ds.Modality = "OPT"
    ds.Manufacturer = "Heidelberg Engineering"
    ds.PatientName = "SHOULD^NOT^LEAK"
    ds.Rows, ds.Columns = rows, cols
    ds.NumberOfFrames = frames
    ds.SamplesPerPixel = 1
    ds.PhotometricInterpretation = "MONOCHROME2"
    ds.BitsAllocated = ds.BitsStored = 8
    ds.HighBit = 7
    ds.PixelRepresentation = 0
    ds.PixelData = (np.arange(frames * rows * cols) % 255).astype(np.uint8).tobytes()
    for k, v in extra.items():
        setattr(ds, k, v)
    buf = io.BytesIO()
    pydicom.dcmwrite(buf, ds, enforce_file_format=True)
    return buf.getvalue()


def test_kinds_and_download_filter():
    assert detect_ext("a/b/scan.nii.gz") == ".nii.gz"
    assert file_kind("x.tar.gz") == "archive"
    assert file_kind("data.z01") == "fragment"
    assert file_kind("img.tif.gz") == "compressed"
    assert wanted_for_download("fundus.zip")
    assert not wanted_for_download("paper.pdf")
    assert not wanted_for_download("reads.fastq.gz")


def test_walker_zip_nested_tar_and_dicom(tmp_path):
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as z:
        z.writestr("deep/OD_001.png", _png_bytes())
    tbuf = io.BytesIO()
    with tarfile.open(fileobj=tbuf, mode="w:gz") as t:
        data = _png_bytes(mode="L")
        info = tarfile.TarInfo("t/pt3_L_scan.png")
        info.size = len(data)
        t.addfile(info, io.BytesIO(data))
    outer = tmp_path / "record.zip"
    with zipfile.ZipFile(outer, "w") as z:
        z.writestr("images/right/a.png", _png_bytes())
        z.writestr("inner.zip", inner.getvalue())
        z.writestr("stack.tar.gz", tbuf.getvalue())
        z.writestr("dicom/IM0001", _dicom_bytes())   # extensionless DICOM
        z.writestr("readme.txt", "hello")

    w = RecordWalker(scratch=tmp_path / "scratch")
    (tmp_path / "scratch").mkdir()
    w.add_file(outer, "record.zip")
    assert len(w.entries) == 4, [e.display for e in w.entries]
    kinds = sorted(e.kind for e in w.entries)
    assert kinds == ["dicom", "raster", "raster", "raster"]

    loaded = {}
    for e, data, path, err in w.read_entries(w.entries):
        assert err is None, err
        loaded[e.display] = load_frame(e.kind, e.ext, data, path)
    dcm = next(v for k, v in loaded.items() if k.endswith("IM0001"))
    assert dcm.image is not None and dcm.facts["frames"] == 3
    assert dcm.facts["dicom"]["Modality"] == "OPT"
    assert "PatientName" not in json.dumps(dcm.facts)
    for v in loaded.values():
        assert v.image is not None and v.image.mode == "RGB"


def test_sixteen_bit_tiff_and_preprocess(tmp_path):
    arr = (np.linspace(0, 4000, 256 * 200).reshape(200, 256)).astype(np.uint16)
    p = tmp_path / "oct.tif"
    Image.fromarray(arr).save(p)
    ld = load_frame("raster", ".tif", None, p)
    assert ld.image is not None and ld.facts["bits"] == 16
    x = preprocess(ld.image)
    assert x.shape == (3, 224, 224) and x.dtype == np.float32


def test_content_crop_matches_training_rule():
    img = Image.new("RGB", (400, 300), (0, 0, 0))
    img.paste((200, 100, 50), (100, 50, 300, 250))
    out = content_crop(img)
    assert out.size == (200, 200)
    blank = Image.new("RGB", (100, 100), (0, 0, 0))
    assert content_crop(blank).size == (100, 100)


@pytest.mark.parametrize("path,side", [
    ("set/OD/img_01.png", "R"),
    ("12_L.png", "L"),
    ("L.png", "U"),            # bare letter without another field
    ("scan_L1_R2.png", "U"),   # glued to digits
    ("pt_left_OS.jpg", "L"),
    ("a_R_b_L.jpg", "U"),      # conflict
    ("zip.zip!/both_eyes/x.png", "B"),
])
def test_laterality(path, side):
    assert dicom_map.laterality(path)[0] == side


def test_class_attributes():
    octa = dicom_map.class_attributes("OCTA", path_text="superficial_slab.png")
    assert octa["Modality"] == "OPTENF" and "128265" in octa["OphthalmicImageType"]
    cfp16 = dicom_map.class_attributes("CFP", bits=16)
    assert cfp16["SOPClassUID"].endswith(".5.2")


def test_manufacturer_hints_guard_weak_optos_words():
    names = dicom_map.manufacturer_hints("Collected at the University of California, Irvine", [])
    assert not any(h["manufacturer"] == "Optos" for h in names)
    hits = dicom_map.manufacturer_hints("Optos California ultra-widefield images", ["a/scan.e2e"])
    makers = [h["manufacturer"] for h in hits]
    assert makers[0] == "Heidelberg Engineering" and "Optos" in makers


def test_float_multiband_tiff(tmp_path):
    tifffile = pytest.importorskip("tifffile")
    arr = np.random.default_rng(0).random((5, 64, 80)).astype(np.float32)
    arr[:, :4, :4] = -3.4e38  # GDAL nodata
    p = tmp_path / "stack.tif"
    tifffile.imwrite(p, arr)
    ld = load_frame("raster", ".tif", None, p)
    assert ld.image is not None and ld.image.size == (80, 64)


def test_weblink_cleaning_and_categories():
    assert weblinks.clean_url("https://github.com/LLNL/tool/releases/The") == "https://github.com/LLNL/tool/releases/"
    assert weblinks.clean_url("https://software.llnl.gov/tool/(updated") == "https://software.llnl.gov/tool/"
    assert weblinks.categorize("https://doi.org/10.5061/dryad.abc")[0] == "dataset_repository"
    assert weblinks.categorize("https://doi.org/10.1016/j.x.2022")[0] == "paper_doi"
    assert weblinks.categorize("https://drive.google.com/file/d/1")[0] == "cloud_drive"
    links = weblinks.collect_links("123", {"external_links": [
        "https://zenodo.org/records/123", "https://www.kaggle.com/datasets/x/y"]}, None, None)
    assert [link["url"] for link in links] == ["https://www.kaggle.com/datasets/x/y"]
    assert links[0]["dataset_likely"]


def test_cmds_documents_validate():
    pytest.importorskip("jsonschema")
    scrape = {"source_id": "10043461", "doi": "10.5281/zenodo.10043461", "title": "T",
              "url": "https://zenodo.org/records/10043461", "access_type": "open", "license": "cc-by-4.0"}
    datacite = {
        "doi": "10.5281/zenodo.10043461", "titles": [{"title": "T"}], "publicationYear": "2023",
        "creators": [{"name": "Doe, Jane", "nameType": "Personal",
                      "nameIdentifiers": [{"nameIdentifier": "0000-0003-2190-4909", "nameIdentifierScheme": "ORCID"}],
                      "affiliation": [{"name": "Uni", "affiliationIdentifier": "https://ror.org/040af2s02",
                                       "affiliationIdentifierScheme": "ROR"}]}],
        "rightsList": [{"rights": "Creative Commons Attribution 4.0 International", "rightsIdentifier": "cc-by-4.0",
                        "rightsUri": "https://creativecommons.org/licenses/by/4.0/legalcode"},
                       {"rightsUri": "info:eu-repo/semantics/openAccess"}],
        "alternateIdentifiers": [{"alternateIdentifier": "oai:zenodo.org:1", "alternateIdentifierType": "oai"}],
        "relatedIdentifiers": [{"relatedIdentifier": "10.5281/zenodo.10043460", "relationType": "IsVersionOf",
                                "relatedIdentifierType": "DOI"}],
        "language": "eng", "publisher": {"name": "Zenodo"},
    }
    legacy = {"metadata": {"access_right": "open", "relations": {"version": [{"index": 0}]}}}
    dd, prov = cmds.build_dataset_description(scrape, legacy, datacite, {"present_classes": ["CFP", "OCT"],
                                                                         "n_files": 2, "total_bytes": 10})
    assert cmds.validate(dd, "dataset_description") == []
    assert prov["datasetConsent"].startswith("placeholder")
    assert dd["language"] == "en" and dd["version"] == "1"
    dsd = cmds.build_structure_description(
        {"CFP": {"count": 80, "est_files": 800, "mean_conf": 0.9},
         "OCT": {"count": 20, "est_files": 200, "mean_conf": 0.8}},
        ["CFP", "OCT"], {"n_images": 1000, "n_classified": 100, "formats": [".jpg"]})
    assert cmds.validate(dsd, "dataset_structure_description") == []
    assert [d["directoryName"] for d in dsd["directoryList"]] == ["retinal_photography", "retinal_oct"]


def test_split_sets_are_colocated(tmp_path):
    from envision_eye_actionable.survey.runner import _colocate_split_sets
    local, dl = tmp_path / "local", tmp_path / "dl"
    local.mkdir()
    dl.mkdir()
    (local / "set.zip").write_bytes(b"x")
    (dl / "set.z01").write_bytes(b"y")
    (local / "other.png").write_bytes(b"z")
    out = _colocate_split_sets([(local / "set.zip", "set.zip"), (dl / "set.z01", "set.z01"),
                                (local / "other.png", "other.png")], dl)
    paths = {k: p for p, k in out}
    assert paths["set.zip"].parent == dl and paths["set.zip"].is_symlink()
    assert paths["other.png"].parent == local


# ---------------------------------------------------------------------------
# review fixes
# ---------------------------------------------------------------------------
def _datacite_min(**extra):
    base = {"doi": "10.5281/zenodo.10043461", "titles": [{"title": "T"}], "publicationYear": "2023",
            "creators": [{"name": "Doe, Jane", "nameType": "Personal"}], "publisher": {"name": "Zenodo"}}
    base.update(extra)
    return base


def _dd(datacite, legacy=None, scrape=None):
    scrape = scrape or {"source_id": "10043461", "title": "T", "access_type": "open"}
    return cmds.build_dataset_description(scrape, legacy or {"metadata": {"access_right": "open"}},
                                          datacite, {"present_classes": ["CFP"]})


def test_datacite_date_types_are_kept():
    pytest.importorskip("jsonschema")
    dd, _ = _dd(_datacite_min(dates=[{"date": "2023-10-26", "dateType": "Issued"},
                                     {"date": "2024-07-11", "dateType": "Updated"},
                                     {"date": "2024", "dateType": "NotAType"}]))
    assert [d["dateType"] for d in dd["date"]] == ["Issued", "Updated", "Other"]
    assert cmds.validate(dd, "dataset_description") == []


def test_enum_values_raises_for_unknown_definition():
    with pytest.raises(KeyError):
        cmds.enum_values("dataset_description", "dateType")      # inline, not a definition
    assert "Issued" in cmds.enum_values("dataset_description", cmds.DATE_TYPE)
    for name in ("identifierType", "titleType", "contributorType", "relationType", "resourceItemType"):
        assert cmds.enum_values("dataset_description", name)


def test_rights_spdx_and_org_schemes():
    pytest.importorskip("jsonschema")
    dc = _datacite_min(
        rightsList=[{"rights": "CC BY 4.0", "rightsIdentifier": "cc-by-4.0", "rightsIdentifierScheme": "spdx"},
                    {"rights": "Custom", "rightsIdentifier": "my-license"}],
        creators=[{"name": "Doe, Jane", "nameType": "Personal",
                   "affiliation": [{"name": "Uni", "affiliationIdentifier": "0000000121865347",
                                    "affiliationIdentifierScheme": "ISNI"}]}])
    dd, _ = _dd(dc)
    ids = [r.get("rightsIdentifier") for r in dd["rights"]]
    assert ids[0] == {"rightsIdentifierValue": "CC-BY-4.0", "rightsIdentifierScheme": "SPDX",
                      "schemeURI": "https://spdx.org/licenses/"}
    assert ids[1] == {"rightsIdentifierValue": "my-license", "rightsIdentifierScheme": "Zenodo license id"}
    aff = dd["creator"][0]["affiliation"][0]["affiliationIdentifier"]
    assert aff["affiliationIdentifierScheme"] == "ISNI" and aff["schemeURI"] == "https://isni.org"
    org = dd["managingOrganization"]["managingOrganizationIdentifier"]
    assert org["managingOrganizationScheme"] == "ISNI" and org["schemeURI"] == "https://isni.org"
    assert cmds.validate(dd, "dataset_description") == []
    # legacy-only license id is mapped too
    dd2, _ = _dd({}, legacy={"metadata": {"access_right": "open", "license": {"id": "cc-zero"}}})
    assert dd2["rights"][0]["rightsIdentifier"]["rightsIdentifierValue"] == "CC0-1.0"


def test_placeholders_make_no_false_claims():
    dd, _ = _dd(_datacite_min(), scrape={"source_id": "1", "title": "T", "access_type": "restricted"})
    assert "license" not in dd["datasetConsent"]["consentsDetails"].lower()
    assert "urlLastChecked" not in dd["accessDetails"]


def test_dicom_standard_only_on_dirs_with_dicom():
    classes = {"CFP": {"count": 10, "est_files": 10, "mean_conf": 0.9},
               "OCT": {"count": 5, "est_files": 5, "mean_conf": 0.9}}
    dsd = cmds.build_structure_description(classes, ["CFP", "OCT"], {
        "n_images": 15, "n_classified": 15, "dicom_classes": ["OCT", "NEG"], "formats": [".dcm", ".jpg"]})
    std = {d["directoryName"]: [s["standardName"] for s in d["relatedStandard"]] for d in dsd["directoryList"]}
    assert not any("DICOM" in s for s in std["retinal_photography"])
    assert any("DICOM" in s for s in std["retinal_oct"])
    dsd2 = cmds.build_structure_description(classes, ["CFP", "OCT"], {
        "n_images": 15, "n_classified": 15, "dicom_classes": ["NEG"], "formats": [".dcm"]})
    assert all(len(d["relatedStandard"]) == 1 for d in dsd2["directoryList"])


def _img(cls, path, dicom=None, **facts):
    f = {"rows": 10, "cols": 10, **facts}
    if dicom is not None:
        f["dicom"] = dicom
    return {"cls": cls, "facts": f, "path": path}


def test_dicom_status_and_header_values():
    # A stray CT DICOM next to fundus JPEGs: values stay from the class map.
    imgs = [_img("CFP", f"a/{i}.jpg") for i in range(5)] + [_img("NEG", "ct/x.dcm", {"Modality": "CT"})]
    agg = dicom_map.aggregate(imgs, [i["path"] for i in imgs], ["CFP", "NEG"], "CFP", "", [])
    assert agg["dicom_mapping_status"] == "mixed" and agg["dicom_Modality"] == "OP"
    # DICOM of the dominant class: header majority wins.
    hdr = {"Modality": "OPT", "SOPClassUID": "1.2.840.10008.5.1.4.1.1.77.1.5.4",
           "ImageType": ["ORIGINAL", "PRIMARY"], "ImageLaterality": "R"}
    imgs = [_img("OCT", f"s/{i}.dcm", hdr) for i in range(3)] + [_img("OCT", "s/left_OS.png")]
    agg = dicom_map.aggregate(imgs, [i["path"] for i in imgs], ["OCT"], "OCT", "", [])
    assert agg["dicom_mapping_status"] == "read_from_header"
    assert agg["dicom_Modality"] == "OPT" and agg["dicom_ImageType"] == "ORIGINAL\\PRIMARY"
    assert agg["laterality_R"] == 3 and agg["laterality_L"] == 1
    assert agg["laterality_source"] == "dicom_header:3; basename:1"


def test_bad_pixel_spacing_does_not_crash_aggregate():
    imgs = [_img("CFP", "a.jpg", spacing_mm=0.01), _img("CFP", "b.jpg", spacing_mm=[0.1, 0.1, 0.2]),
            _img("CFP", "c.jpg", spacing_mm=[0.01, 0.02])]
    agg = dicom_map.aggregate(imgs, ["a.jpg", "b.jpg", "c.jpg"], ["CFP"], "CFP", "", [])
    assert agg["dicom_PixelSpacing_mm"] == "0.01\\0.02 (1)"


def test_single_value_dicom_pixel_spacing(tmp_path):
    pydicom = pytest.importorskip("pydicom")
    ds = pydicom.dcmread(io.BytesIO(_dicom_bytes(frames=1)), force=True)
    ds.PixelSpacing = 0.5
    buf = io.BytesIO()
    ds.save_as(buf)
    ld = load_frame("dicom", ".dcm", buf.getvalue(), None)
    assert ld.image is not None and "spacing_mm" not in ld.facts


def test_oct_anatomy_needs_explicit_phrases():
    retina = "SCT 5665001"
    assert retina in dicom_map.class_attributes(
        "OCT", path_text="oct/b1.png", record_text="scan angle of 30 degrees; light passes the cornea")[
        "AnatomicRegion"]
    assert "31636006" in dicom_map.class_attributes(
        "OCT", path_text="x.png", record_text="Anterior segment OCT of the iris")["AnatomicRegion"]
    assert "31636006" in dicom_map.class_attributes("OCT", path_text="data/cornea/1.png")["AnatomicRegion"]
    assert "81016008" in dicom_map.class_attributes("OCT", path_text="onh_scans/1.png")["AnatomicRegion"]


def _squash_then_eval_tf(img, size=224, stored=256):
    """Training path: content crop, squash to a 256 PNG square, then
    Resize(224) + CenterCrop(224) (torchvision semantics, written out)."""
    from envision_eye_actionable.survey.images import _MEAN, _STD
    img = content_crop(img).resize((stored, stored), Image.BILINEAR)
    w, h = img.size
    nw, nh = (size, int(size * h / w)) if w <= h else (int(size * w / h), size)
    img = img.resize((nw, nh), Image.BILINEAR)
    left, top = int(round((nw - size) / 2.0)), int(round((nh - size) / 2.0))
    img = img.crop((left, top, left + size, top + size))
    a = np.asarray(img, np.float32).transpose(2, 0, 1) / 255.0
    return (a - _MEAN) / _STD


@pytest.mark.parametrize("size", [(640, 480), (300, 1000), (768, 496), (1024, 1024), (4000, 300)])
def test_preprocess_matches_training_squash(size):
    rng = np.random.default_rng(1)
    img = Image.fromarray(rng.integers(0, 255, (size[1], size[0], 3), dtype=np.uint8), "RGB")
    diff = np.abs(preprocess(img) - _squash_then_eval_tf(img)) * 0.229 * 255   # back to 8-bit units
    assert diff.max() <= 1.0


def test_preprocess_keeps_the_sides_of_wide_images():
    # A 768x496 OCT-like frame with bright bands at the far left and right:
    # a short-side resize + center crop would cut both off.
    a = np.full((496, 768, 3), 120, np.uint8)
    a[:, :60] = 250
    a[:, -60:] = 60                    # darker, but above the content-crop threshold
    x = preprocess(Image.fromarray(a, "RGB"))
    col = x[0].mean(axis=0)            # mean over rows, channel 0
    assert col[:5].mean() > 1.5 and col[-5:].mean() < -0.9


def test_preprocess_extreme_aspect_is_cheap():
    img = Image.new("RGB", (100000, 2), (120, 80, 40))
    assert preprocess(img).shape == (3, 224, 224)


def test_content_crop_box_matches_where():
    rng = np.random.default_rng(3)
    for _ in range(5):
        a = np.zeros((120, 160, 3), np.uint8)
        y0, x0 = rng.integers(0, 40, 2)
        a[y0:y0 + 60, x0:x0 + 90] = rng.integers(40, 255, (60, 90, 3))
        img = Image.fromarray(a)
        g = np.asarray(img.convert("L"), np.float32)
        ys, xs = np.where(g > max(12.0, float(g.max()) * 0.08))
        assert content_crop(img).size == (int(xs.max() - xs.min() + 1), int(ys.max() - ys.min() + 1))


def test_volume_uncompressed_cap(tmp_path):
    nrrd = pytest.importorskip("nrrd")
    vol = np.zeros((40, 50, 60), np.uint16)
    p = tmp_path / "v.nrrd"
    nrrd.write(str(p), vol, {"encoding": "gzip"})
    from envision_eye_actionable.survey.images import _load_volume
    ld = _load_volume(p, ".nrrd", 80_000_000, max_bytes=100_000)   # 240 KB uncompressed, tiny on disk
    assert ld.image is None and ld.error == "volume too large to load"
    assert p.stat().st_size < 100_000
    ok = _load_volume(p, ".nrrd", 80_000_000)
    assert ok.image is not None and ok.facts["frames"] == 40


def test_mha_middle_slice(tmp_path):
    sitk = pytest.importorskip("SimpleITK")
    vol = np.zeros((8, 30, 40), np.uint8)       # z, y, x
    vol[4] = (np.arange(40, dtype=np.uint8) * 5)[None, :]   # only the middle slice has content
    im = sitk.GetImageFromArray(vol)
    im.SetSpacing((0.1, 0.2, 0.3))
    p = tmp_path / "v.mha"
    sitk.WriteImage(im, str(p), True)
    ld = load_frame("volume", ".mha", None, p)
    assert ld.image is not None and ld.image.size == (40, 30)
    assert ld.facts["frames"] == 8 and ld.facts["spacing_mm"] == [0.2, 0.1]
    assert np.asarray(ld.image).max() > 0


def test_nested_archives_without_images_are_deleted(tmp_path):
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as z:
        z.writestr("notes.txt", "x")
    inner_img = io.BytesIO()
    with zipfile.ZipFile(inner_img, "w") as z:
        z.writestr("eye.png", _png_bytes())
    outer = tmp_path / "outer.zip"
    with zipfile.ZipFile(outer, "w") as z:
        z.writestr("p1.zip", inner.getvalue())
        z.writestr("p2.zip", inner_img.getvalue())
    scratch = tmp_path / "scratch"
    scratch.mkdir()
    w = RecordWalker(scratch=scratch)
    w.add_file(outer, "outer.zip")
    kept = [p.name for p in (scratch / "nested").rglob("*") if p.is_file()]
    assert kept == ["p2.zip"] and len(w.entries) == 1
    got = list(w.read_entries(w.entries))
    assert got[0][3] is None


def test_read_entries_no_duplicate_errors(tmp_path):
    z = tmp_path / "a.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for i in range(3):
            zf.writestr(f"{i}.png", _png_bytes())
    w = RecordWalker(scratch=tmp_path)
    w.add_file(z, "a.zip")

    def flaky(c, group, max_inmem):
        yield group[0], b"x", None, None
        raise RuntimeError("corrupt stream")
    w._read_zip = flaky
    out = list(w.read_entries(w.entries))
    assert len(out) == 3 and [o[3] for o in out].count(None) == 1


def test_read_7z_respects_disk_floor(tmp_path):
    from envision_eye_actionable.survey.archives import Container, Entry
    w = RecordWalker(scratch=tmp_path, disk_floor_bytes=10 ** 18)   # never room
    c = Container("7z", tmp_path / "x.7z", "x.7z", 0, use_7z=True)
    group = [Entry(0, f"{i}.png", f"x.7z!/{i}.png", ".png", "raster", 10) for i in range(3)]
    out = list(w._read_7z(c, group))
    assert [o[3] for o in out] == ["skipped (disk floor)"] * 3


def test_retry_after_parsing():
    from envision_eye_actionable.survey.zenodo import retry_after_seconds
    assert retry_after_seconds("30", 5) == 30
    assert retry_after_seconds("60.0", 5) == 60
    assert retry_after_seconds(None, 5) == 60           # 429 floor
    assert retry_after_seconds(None, 5, 503) == 5
    assert retry_after_seconds("Wed, 21 Oct 2015 07:28:00 GMT", 5) == 0   # in the past
    assert retry_after_seconds("junk", 5) == 60


class _Resp:
    def __init__(self, status, body=b"", headers=None):
        self.status_code, self.body, self.headers = status, body, headers or {}

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def iter_content(self, chunk_size):
        yield self.body


class _Session:
    def __init__(self, responses):
        self.responses, self.calls = list(responses), 0

    def get(self, *a, **k):
        self.calls += 1
        return self.responses.pop(0) if self.responses else _Resp(200, b"bad")


class _Throttle:
    def __init__(self):
        self.waits, self.backoffs, self.observed = 0, [], []

    def throttle(self):
        self.waits += 1

    def observe(self, r):
        self.observed.append(r.status_code)

    def backoff(self, s):
        self.backoffs.append(s)


def test_stream_download_md5_cap_and_shared_throttle(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import zenodo
    sleeps = []
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: sleeps.append(s))
    th = _Throttle()
    sess = _Session([_Resp(429, headers={"Retry-After": "90"})] + [_Resp(200, b"bad")] * 10)
    ok, err = zenodo.stream_download(sess, "u", tmp_path / "f.bin", 3, "0" * 32, throttle=th)
    assert not ok and err == "md5 mismatch"
    assert sess.calls == 4                 # 429 + first try + 2 restarts
    assert th.waits == sess.calls and th.backoffs == [90.0]
    assert len(sleeps) == 2                # backoff between md5 restarts, none for the 429 (shared)


def test_crashed_record_is_recovered(tmp_path):
    from envision_eye_actionable.survey.results import load_results
    from envision_eye_actionable.survey.runner import Survey, SurveyConfig
    cfg = SurveyConfig(scrape_path=tmp_path / "s.json", model_path=tmp_path / "m.onnx",
                       out_dir=tmp_path / "out", downloads_dir=tmp_path / "dl", scratch_dir=tmp_path / "scratch")
    cfg.out_dir.mkdir()
    s = Survey.__new__(Survey)
    s.cfg = cfg
    s._claim_scratch()
    rs = cfg.scratch_dir / "records"
    (rs / "123" / "dl").mkdir(parents=True)
    (rs / "123" / "dl" / "big.zip.part").write_bytes(b"x")
    (rs / "999").mkdir()
    (rs / "keep_me").mkdir()
    (cfg.scratch_dir / "456").mkdir()          # outside records/: never touched
    s.results_path = cfg.out_dir / "survey_results.jsonl"
    s.events_path = cfg.out_dir / "survey_events.jsonl"
    s.model_info = {}
    s._append_row({"record_id": "1", "status": "ok"})
    s._append_row({"record_id": "123", "status": "started", "started_at": "t"})
    done = load_results(s.results_path)
    crashed = s._recover_crashed(done, {"123": {"source_id": "123", "title": "Big"}})
    s._sweep_scratch()
    assert crashed == ["123"] and done["123"]["status"] == "crashed"
    assert load_results(s.results_path)["123"]["title"] == "Big"
    assert not (rs / "123").exists() and not (rs / "999").exists()
    assert (rs / "keep_me").exists() and (cfg.scratch_dir / "456").exists()


def test_workbook_sheets_and_directory_counts(tmp_path):
    pytest.importorskip("openpyxl")
    from envision_eye_actionable.survey.excel import build_workbook
    from openpyxl import load_workbook
    out = tmp_path / "res"
    (out / "cmds" / "1").mkdir(parents=True)
    (out / "cmds" / "1" / "dataset_description.json").write_text('{"a": 1}', encoding="utf-8")
    rows = [{"record_id": "1", "status": "ok", "cmds_dir": "cmds/1", "cmds_directories": "retinal_photography/cfp (3)",
             "dd_dates": "Issued: 2023", "laterality_source": "basename:3"},
            {"record_id": "2", "status": "no_image_files", "cmds_dir": "cmds/2", "cmds_directories": ""}]
    res = out / "survey_results.jsonl"
    res.write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
    build_workbook(res, out / "w.xlsx")
    wb = load_workbook(out / "w.xlsx")
    assert "CMDS_JSON" in wb.sheetnames
    head = [c.value for c in wb["Records"][1]]
    for col in ("CMDS JSON folder (relative to the results dir)", "dataset_description dates",
                "Laterality source (files per source)", "Fields set from the classifier",
                "Manufacturer hint sources"):
        assert col in head
    sf = {r[1].value: r[4].value for r in wb["Schema_Fields"].iter_rows(min_row=2)
          if r[0].value == "dataset_structure_description"}
    assert sf["directoryList"] == 1 and sf["metadataFileList"] == 2
    assert wb["CMDS_JSON"]["E2"].value == '{"a": 1}'


# ---------------------------------------------------------------------------
# second review round
# ---------------------------------------------------------------------------
def test_scratch_dir_must_be_the_surveys_own(tmp_path):
    from envision_eye_actionable.survey.runner import SCRATCH_MARKER, Survey, SurveyConfig
    data = tmp_path / "actionable"
    (data / "123").mkdir(parents=True)
    (data / "123" / "keep.json").write_text("{}")
    cfg = SurveyConfig(scrape_path=tmp_path / "s.json", model_path=tmp_path / "m.onnx",
                       out_dir=tmp_path / "out", downloads_dir=tmp_path / "dl", scratch_dir=data)
    s = Survey.__new__(Survey)
    s.cfg = cfg
    with pytest.raises(ValueError, match="not created by the survey"):
        s._claim_scratch()
    assert (data / "123" / "keep.json").exists()
    # A new dir is claimed with a marker; a claimed dir is accepted again.
    cfg.scratch_dir = tmp_path / "new_scratch"
    s._claim_scratch()
    assert (cfg.scratch_dir / SCRATCH_MARKER).exists() and (cfg.scratch_dir / "records").is_dir()
    (cfg.scratch_dir / "records" / "5").mkdir()
    s._claim_scratch()


def test_relative_paths_are_resolved_before_the_overlap_guard(tmp_path, monkeypatch):
    from pathlib import Path

    from envision_eye_actionable.survey.runner import Survey, SurveyConfig
    monkeypatch.chdir(tmp_path)
    cfg = SurveyConfig(scrape_path=Path("s.json"), model_path=Path("m.onnx"), out_dir=Path("out"),
                       downloads_dir=Path("data/dl"), scratch_dir=Path("data"))
    s = Survey.__new__(Survey)
    s.cfg = cfg
    with pytest.raises(ValueError, match="overlaps downloads"):
        s._guard_paths()
    assert cfg.scratch_dir.is_absolute()


def test_cli_rejects_missing_explicit_downloads_dir(tmp_path, capsys):
    from envision_eye_actionable.survey.cli import main
    (tmp_path / "m.onnx").write_bytes(b"x")
    with pytest.raises(SystemExit):
        main(["run", "--model", str(tmp_path / "m.onnx"), "--downloads-dir", str(tmp_path / "nope")])
    assert "does not exist" in capsys.readouterr().err


def test_truncated_jsonl_line_is_terminated(tmp_path):
    from envision_eye_actionable.survey.results import load_results
    from envision_eye_actionable.survey.runner import _terminate_last_line
    p = tmp_path / "r.jsonl"
    p.write_text('{"record_id": "1", "status": "ok"}\n{"record_id": "2", "sta', encoding="utf-8")
    _terminate_last_line(p)
    _terminate_last_line(p)                    # idempotent
    with open(p, "a", encoding="utf-8") as fp:
        fp.write(json.dumps({"record_id": "3", "status": "started"}) + "\n")
    rows = load_results(p)
    assert rows["3"]["status"] == "started" and "2" not in rows
    _terminate_last_line(tmp_path / "missing.jsonl")


def test_octa_slab_ignores_loose_words_in_prose():
    def code(**k):
        return dicom_map.class_attributes("OCTA", **k)["OphthalmicImageType"]
    assert "128259" in code(record_text="A deep learning dataset of OCTA images")
    assert "128259" in code(record_text="License: CC BY 4.0; optic disc and choroid")
    assert "128269" in code(record_text="en face images of the deep capillary plexus")
    assert "128273" in code(record_text="choriocapillaris flow deficits")
    # two slabs named -> generic default
    assert "128259" in code(record_text="superficial capillary plexus and deep capillary plexus")
    # paths keep the loose keywords, and win over the text
    assert "128269" in code(path_text="p1/dcp/1.png p1/dcp/2.png", record_text="choriocapillaris")
    assert "128259" in code(path_text="p1/svp/1.png p1/dcp/1.png")


def test_dicom_software_and_header_attributes():
    raw = _dicom_bytes(SoftwareVersions=["6.16", "2.0"], BodyPartExamined="EYE",
                       HorizontalFieldOfView=30, IlluminationWaveLength=870, SeriesDescription="Volume IR")
    ld = load_frame("dicom", ".dcm", raw, None)
    assert ld.facts["software"] == "6.16\\2.0"
    assert float(ld.facts["dicom"]["HorizontalFieldOfView"]) == 30
    hdr = {"Modality": "OPT", "SOPClassUID": "1.2.840.10008.5.1.4.1.1.77.1.5.4",
           "BodyPartExamined": "EYE", "AnatomicRegionSequence": "SCT 81016008 Optic nerve head",
           "OphthalmicImageTypeCodeSequence": "DCM 1 X", "HorizontalFieldOfView": 20.0,
           "IlluminationWaveLength": 870.0, "SeriesDescription": "ONH cube"}
    imgs = [_img("OCT", f"s/{i}.dcm", hdr, software="6.16") for i in range(3)]
    agg = dicom_map.aggregate(imgs, [i["path"] for i in imgs], ["OCT"], "OCT", "", [])
    assert agg["dicom_AnatomicRegion_code"] == "SCT 81016008 Optic nerve head"
    assert agg["dicom_OphthalmicImageType_code"] == "DCM 1 X"
    assert "AnatomicRegion" in agg["dicom_header_attributes_used"]
    assert agg["dicom_SoftwareVersions"] == "6.16 (3)"
    assert agg["dicom_HorizontalFieldOfView"] == "20.0 (3)"
    assert agg["dicom_IlluminationWaveLength"] == "870.0 (3)"
    assert agg["dicom_SeriesDescription"] == "ONH cube (3)"


def test_mask_like_flags_color_label_maps():
    import numpy as np
    from PIL import Image
    from envision_eye_actionable.survey import images
    lab = np.zeros((100, 100, 3), np.uint8)
    lab[20:50, 20:50] = (128, 0, 128)
    lab[60:90, 60:90] = (255, 165, 0)
    f = {}
    images._pixel_flags(Image.fromarray(lab), f)
    assert f["mask_like"] is True and f["rgb_is_gray"] is False
    rng = np.random.default_rng(0)
    photo = rng.integers(0, 256, (100, 100, 3), dtype=np.uint8)
    f = {}
    images._pixel_flags(Image.fromarray(photo), f)
    assert f["mask_like"] is False


def test_samples_per_pixel_drops_alpha():
    # A Photoshop RGBA TIFF reports 4 samples; DICOM RGB has exactly 3.
    imgs = [_img("NEG", "fig.tif", photometric="RGB", bits=8, samples=4),
            _img("NEG", "m.png", photometric="MONOCHROME2", bits=8, samples=2)]
    agg = dicom_map.aggregate(imgs, ["fig.tif", "m.png"], [], "NEG", "", [])
    assert agg["dicom_SamplesPerPixel_set"] == "3 (1); 1 (1)"


def test_pixel_module_columns_and_sixteen_bit_rule():
    imgs = [_img("CFP", "a.jpg", photometric="RGB", bits=8, lossy="01", lossy_method="ISO_10918_1",
                 description="Topcon  TRC-50DX\nfield 1"),
            _img("CFP", "b.tif", photometric="MONOCHROME2", bits=12)]
    agg = dicom_map.aggregate(imgs, ["a.jpg", "b.tif"], ["CFP"], "CFP", "", [])
    assert agg["dicom_SamplesPerPixel_set"] == "3 (1); 1 (1)"
    assert agg["dicom_BitsAllocated_set"] == "8 (1); 16 (1)"
    assert agg["dicom_HighBit_set"] == "7 (1); 11 (1)"
    assert agg["dicom_LossyImageCompressionMethod"] == "ISO_10918_1 (1)"
    assert agg["dicom_ImageComments"] == "Topcon TRC-50DX field 1 (1)"
    assert "Diffuse direct illumination" in agg["dicom_IlluminationType_code"]
    assert agg["dicom_pixel_depth_note"] == ""
    # 32-bit / float data never get the 16-bit SOP class, and are flagged.
    assert dicom_map.class_attributes("CFP", bits=32)["SOPClassUID"].endswith(".5.1")
    assert dicom_map.class_attributes("FAF", bits=12)["SOPClassUID"].endswith(".5.2")
    assert dicom_map.class_attributes("PSC", bits=16)["SOPClassUID"].endswith(".5.2")
    agg = dicom_map.aggregate([_img("FAF", "f.nii", bits=32, float=True)], ["f.nii"], ["FAF"], "FAF", "", [])
    assert agg["dicom_pixel_depth_note"].startswith("1 sampled image")


def test_jpeg_facts_carry_the_compression_method():
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (120, 60, 30)).save(buf, format="JPEG")
    ld = load_frame("raster", ".jpg", buf.getvalue(), None)
    assert ld.facts["lossy"] == "01" and ld.facts["lossy_method"] == "ISO_10918_1"


def test_compressed_non_images_are_not_downloaded():
    assert wanted_for_download("scan.tif.gz") and wanted_for_download("dicoms.tar.bz2")
    assert wanted_for_download("slice.nii.gz") and wanted_for_download("IM0001.gz")
    # arrays are pixel-bearing now (image-shaped datasets are converted)
    assert wanted_for_download("arr.npy.gz") and wanted_for_download("stack.h5.gz")
    for name in ("table.csv.gz", "genes.gtf.gz", "notes.txt.xz", "weights.pt.gz"):
        assert not wanted_for_download(name), name


def test_planar_tiff_uses_image_dimensions(tmp_path):
    tifffile = pytest.importorskip("tifffile")
    from envision_eye_actionable.survey.images import _load_tiff
    arr = np.random.default_rng(0).integers(0, 4000, (3, 40, 50)).astype(np.uint16)
    p = tmp_path / "planar.tif"
    tifffile.imwrite(p, arr, photometric="rgb", planarconfig="separate")
    ld = _load_tiff(None, p, ".tif", max_pixels=10_000)
    assert (ld.facts["rows"], ld.facts["cols"], ld.facts["samples"]) == (40, 50, 3)
    assert ld.image is not None and ld.image.size == (50, 40)
    assert "strided" not in ld.facts["conversion"]
    # over the byte budget (40*50*3*2 = 12000 bytes) or the pixel budget on
    # H*W (2000, not S*H = 120): read strided, never refused
    ld = _load_tiff(None, p, ".tif", max_pixels=10_000, max_bytes=10_000)
    assert ld.image is not None and "strided 1/" in ld.facts["conversion"]
    ld = _load_tiff(None, p, ".tif", max_pixels=1_999)
    assert ld.image is not None and ld.image.size == (25, 20) and "strided 1/2" in ld.facts["conversion"]


def test_datacite_subject_codes_and_keywords_are_merged():
    dc = _datacite_min(subjects=[
        {"subject": "Retina", "subjectScheme": "MeSH", "schemeUri": "https://id.nlm.nih.gov/mesh/",
         "valueUri": "https://id.nlm.nih.gov/mesh/D012160"},
        {"subject": "fundus"}])
    dd, _ = _dd(dc, legacy={"metadata": {"access_right": "open", "keywords": ["retina", "glaucoma"]}})
    subj = {s["subjectValue"]: s for s in dd["subject"]}
    assert subj["Retina"]["subjectIdentifier"] == {
        "classificationCode": "D012160", "subjectScheme": "MeSH",
        "schemeURI": "https://id.nlm.nih.gov/mesh/", "valueURI": "https://id.nlm.nih.gov/mesh/D012160"}
    assert "glaucoma" in subj and "retina" not in subj and "fundus" in subj
    pytest.importorskip("jsonschema")
    assert cmds.validate(dd, "dataset_description") == []


class _LinkResp:
    def __init__(self, status, url, headers=None):
        self.status_code, self.url, self.headers = status, url, headers or {}

    def close(self):
        pass


class _LinkSession:
    def __init__(self, statuses):
        self.statuses, self.calls = list(statuses), []
        self.headers = {}

    def head(self, url, **k):
        self.calls.append(url)
        st = self.statuses.pop(0)
        if isinstance(st, Exception):
            raise st
        return _LinkResp(st, url, {"Retry-After": "1"} if st == 429 else {})

    get = head


def test_link_checker_429_cooldown_and_cache(tmp_path, monkeypatch):
    import requests
    monkeypatch.setattr(weblinks.time, "sleep", lambda s: None)
    lc = weblinks.LinkChecker(tmp_path / "links.json", min_interval=0)
    lc.session = _LinkSession([429, 200])
    assert lc.check("https://example.org/a")["http_status"] == 200      # retried once
    assert "https://example.org/a" in lc.cache
    # 5xx: HEAD then GET fallback; transport error; 429 twice.
    lc.session = _LinkSession([503, 503, requests.ConnectionError("x"), 429, 429])
    assert lc.check("https://example.org/b")["http_status"] == 503
    assert lc.check("https://example.org/c")["error"] == "ConnectionError"
    assert lc.check("https://example.org/d")["http_status"] == 429
    assert not {"https://example.org/b", "https://example.org/c", "https://example.org/d"} & set(lc.cache)
    lc.save()
    lc2 = weblinks.LinkChecker(tmp_path / "links.json", min_interval=0)
    assert set(lc2.cache) == {"https://example.org/a"}
    # A host in a long cooldown is not queried again during this run.
    lc2.session = _LinkSession([429])
    lc2.session.head = lambda url, **k: _LinkResp(429, url, {"Retry-After": "3600"})
    assert "Retry-After 3600" in lc2.check("https://slow.example/x")["error"]
    lc2.session = _LinkSession([])
    assert "not checked" in lc2.check("https://slow.example/y")["error"]
    # Zenodo links use the survey's shared throttle and its 429 backoff.
    th = _Throttle()
    lc3 = weblinks.LinkChecker(tmp_path / "l3.json", min_interval=0, zenodo=th)
    lc3.session = _LinkSession([429, 200])
    lc3.check("https://doi.org/10.5281/zenodo.123")
    assert th.waits == 2 and th.backoffs == [1.0] and th.observed == [429, 200]


def test_metadata_only_cmds_for_records_without_classification(tmp_path):
    pytest.importorskip("jsonschema")
    from envision_eye_actionable.survey.runner import Survey, SurveyConfig
    cfg = SurveyConfig(scrape_path=tmp_path / "s.json", model_path=tmp_path / "m.onnx",
                       out_dir=tmp_path / "out", downloads_dir=tmp_path / "dl", scratch_dir=tmp_path / "sc")
    s = Survey.__new__(Survey)
    s.cfg = cfg
    rec = {"source_id": "10043461", "title": "T", "access_type": "open", "size_mb": 5}
    row = {"record_id": "10043461", "status": "skipped_disk"}
    s._write_cmds_metadata_only(row, rec, {"metadata": {"access_right": "open"}}, _datacite_min(),
                                "skipped_disk: files not read")
    assert row["dd_valid"] and row["dsd_valid"] and row["cmds_directories"] == ""
    assert row["cmds_note"] == "skipped_disk: files not read"
    assert (cfg.out_dir / "cmds" / "10043461" / "dataset_description.json").exists()


def test_encrypted_archives_are_skipped_without_prompting(tmp_path):
    import subprocess

    from envision_eye_actionable.survey import archives
    if not archives.SEVEN_ZIP:
        pytest.skip("no 7z binary")
    src = tmp_path / "img.png"
    src.write_bytes(_png_bytes())
    for name, flags in (("hdr.7z", ["-mhe=on"]), ("members.7z", []), ("enc.zip", ["-tzip"])):
        subprocess.run([archives.SEVEN_ZIP, "a", "-psecret", *flags, str(tmp_path / name), str(src)],
                       check=True, capture_output=True, stdin=subprocess.DEVNULL, timeout=60)
    (tmp_path / "scratch").mkdir()
    w = RecordWalker(scratch=tmp_path / "scratch")
    for name in ("hdr.7z", "members.7z", "enc.zip"):
        w.add_file(tmp_path / name, name)
    assert w.entries == [] and w.kind_counts["encrypted"] == 3
    assert any("encrypted archive skipped" in e for e in w.errors)


# ---------------------------------------------------------------------------
# full-run hardening: remote zip sampling, throttle, skipped_size, columns
# ---------------------------------------------------------------------------
class _HttpResp:
    """Minimal requests.Response stand-in (status, headers, body)."""

    def __init__(self, status, body=b"", headers=None):
        self.status_code, self.body, self.headers = status, body, headers or {}

    @property
    def content(self):
        return self.body

    def json(self):
        return json.loads(self.body)

    def iter_content(self, chunk_size=1 << 20):
        for i in range(0, len(self.body), chunk_size):
            yield self.body[i:i + chunk_size]

    def close(self):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


class _FakeZenodo:
    """Serves one record's files like Zenodo: /content (with Range), the
    zip /container listing (optionally truncated) and /container/<member>."""

    def __init__(self, rid, files: dict, truncated=False, member_status=200, listing_status=200,
                 hide_truncation=False):
        self.rid, self.files = rid, files
        self.truncated, self.member_status, self.listing_status = truncated, member_status, listing_status
        # hide_truncation: cut the listing but say truncated=false (total stays right)
        self.hide_truncation = hide_truncation
        self.calls = []

    def get(self, url, headers=None, stream=False, timeout=None):
        from urllib.parse import unquote

        from envision_eye_actionable.survey.zenodo import API
        headers = headers or {}
        self.calls.append((url, headers.get("Range")))
        prefix = f"{API}/{self.rid}/files/"
        if not url.startswith(prefix):
            return _HttpResp(404)
        key_q, _, tail = url[len(prefix):].partition("/")
        data = self.files.get(unquote(key_q))
        if data is None:
            return _HttpResp(404)
        if tail == "content":
            rng = headers.get("Range")
            if rng:
                a, b = (int(x) for x in rng.split("=")[1].split("-"))
                return _HttpResp(206, data[a:b + 1])
            return _HttpResp(200, data)
        if tail == "container" and self.listing_status != 200:
            return _HttpResp(self.listing_status)
        try:
            zf = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile:
            return _HttpResp(500)                 # Zenodo answers 500 for files it cannot open
        if tail == "container":
            entries = [{"key": i.filename, "size": i.file_size, "compressed_size": i.compress_size, "crc": i.CRC}
                       for i in zf.infolist() if not i.is_dir()]
            total = len(entries)
            if self.hide_truncation:
                return _HttpResp(200, json.dumps({"entries": entries[:2], "truncated": False,
                                                  "total": total}).encode())
            if self.truncated:
                entries = entries[:2]
            return _HttpResp(200, json.dumps({"entries": entries, "truncated": self.truncated,
                                              "total": total if not self.truncated else len(entries)}).encode())
        if tail.startswith("container/"):
            if self.member_status != 200:
                return _HttpResp(self.member_status)
            return _HttpResp(200, zf.read(unquote(tail[len("container/"):])))
        return _HttpResp(404)


def _zip_bytes(members: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def _client(tmp_path, session, **kw):
    from envision_eye_actionable.survey.zenodo import ZenodoClient
    c = ZenodoClient(tmp_path / "cache", None, min_interval=0, max_per_minute=None, **kw)
    c.session = session
    return c


def test_throttle_keeps_a_sliding_minute_under_the_limit(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import zenodo
    clock = [1000.0]
    monkeypatch.setattr(zenodo.time, "monotonic", lambda: clock[0])
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: clock.__setitem__(0, clock[0] + s))
    c = zenodo.ZenodoClient(tmp_path / "cache", None, min_interval=0.1, max_per_minute=110)
    stamps = []
    for _ in range(300):
        c.throttle()
        stamps.append(clock[0])
    assert c.n_requests == 300
    for i in range(len(stamps)):
        in_window = [t for t in stamps if stamps[i] <= t < stamps[i] + 60.0]
        assert len(in_window) <= 110
    assert min(b - a for a, b in zip(stamps, stamps[1:])) >= 0.1 - 1e-9
    # 300 requests at 110/min need at least two full windows
    assert stamps[-1] - stamps[0] >= 120.0
    # a 429 cooldown pauses the next request
    c.backoff(30)
    t0 = clock[0]
    c.throttle()
    assert clock[0] - t0 >= 30


def test_get_retries_with_exponential_backoff(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import zenodo
    sleeps = []
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: sleeps.append(s))

    class Seq:
        def __init__(self, statuses):
            self.statuses, self.n = list(statuses), 0

        def get(self, *a, **k):
            self.n += 1
            st = self.statuses.pop(0)
            return _HttpResp(st, b"{}", {"Retry-After": "0.2"} if st == 429 else {})

    c = _client(tmp_path, Seq([503, 502, 429, 200]))
    r = c.get("https://zenodo.org/api/records/1", base_delay=2.0)
    assert r.status_code == 200 and c.session.n == 4
    assert sleeps[:2] == [2.0, 4.0]          # 5xx: exponential
    assert c._not_before > 0                 # 429: shared cooldown from Retry-After
    c2 = _client(tmp_path, Seq([404]))
    with pytest.raises(zenodo.RemoteError) as ei:
        c2.get("u")
    assert ei.value.status == 404 and c2.session.n == 1
    c3 = _client(tmp_path, Seq([500] * 3))
    with pytest.raises(zenodo.RemoteError, match="gave up"):
        c3.get("u", max_tries=3)


def test_allocate_is_proportional_and_exact():
    from envision_eye_actionable.survey.remote import allocate
    assert allocate({"a": 900, "b": 100}, 300) == {"a": 270, "b": 30}
    assert allocate({"a": 5, "b": 3}, 300) == {"a": 5, "b": 3}
    out = allocate({"a": 1, "b": 1000, "c": 10}, 50)
    assert sum(out.values()) == 50 and out["a"] <= 1 and out["c"] <= 10
    out = allocate({"a": 2, "b": 2, "c": 2}, 5)
    assert sum(out.values()) == 5


def test_remote_listing_container_then_range_fallback(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import remote, zenodo
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: None)     # retry backoff of the refused member
    png = _png_bytes()
    members = {f"set/img_{i:02d}.png": png for i in range(6)}
    members.update({"__MACOSX/set/._img_00.png": b"junk", "set/._img_01.png": b"junk", "README.txt": b"hi",
                    "set/inner.zip": _zip_bytes({"a.png": png})})
    data = _zip_bytes(members)
    rid = "42"
    fake = _FakeZenodo(rid, {"imgs.zip": data})
    c = _client(tmp_path, fake)
    lst = remote.list_zip(c, rid, "imgs.zip", len(data))
    assert lst.source == "container" and not lst.truncated_listing
    assert len(lst.images) == 6 and len(lst.nested) == 1          # junk members dropped
    assert not any("MACOSX" in m.name or "/._" in m.name for m in lst.members)

    # truncated container listing -> central directory read by range requests
    fake_t = _FakeZenodo(rid, {"imgs.zip": data}, truncated=True)
    c_t = _client(tmp_path, fake_t)
    lst_t = remote.list_zip(c_t, rid, "imgs.zip", len(data))
    assert lst_t.source == "range" and lst_t.truncated_listing
    assert sorted(m.name for m in lst_t.images) == sorted(m.name for m in lst.images)
    assert any(rng for _, rng in fake_t.calls)                    # Range headers were sent

    # member fetch: container endpoint, CRC checked
    m = lst.images[0]
    ok, err = remote.fetch_member(c, rid, lst, m, tmp_path / "m" / "x.png", len(data))
    assert ok and (tmp_path / "m" / "x.png").read_bytes() == png
    # container endpoint refuses the member -> range read of the member bytes
    fake_r = _FakeZenodo(rid, {"imgs.zip": data}, member_status=500)
    c_r = _client(tmp_path, fake_r)
    rz = remote.RangeZips(c_r, rid)
    ok, err = remote.fetch_member(c_r, rid, lst, m, tmp_path / "r" / "x.png", len(data), rz)
    rz.close()
    assert ok, err
    assert (tmp_path / "r" / "x.png").read_bytes() == png
    # a member path can never escape the scratch root
    assert remote.safe_member_path(tmp_path, 0, "../../etc/passwd").is_relative_to(tmp_path)


def test_range_file_never_downloads_a_whole_file(tmp_path):
    from envision_eye_actionable.survey import remote
    from envision_eye_actionable.survey.zenodo import RemoteError

    class IgnoresRange:
        def get(self, url, headers=None, stream=False, timeout=None):
            assert stream, "range reads must stream so the body is not read before the status check"
            return _HttpResp(200, b"x" * 100)

    fp = remote.HttpRangeFile(_client(tmp_path, IgnoresRange()), "u", 10_000)
    fp.seek(-22, 2)
    with pytest.raises(RemoteError, match="ignored the Range"):
        fp.read(22)


def test_sample_remote_zips_spreads_the_cap_and_is_seeded(tmp_path):
    import random

    from envision_eye_actionable.survey import remote
    png = _png_bytes(size=(40, 30))
    big = _zip_bytes({f"a/{i:03d}.png": png for i in range(90)})
    small = _zip_bytes({f"b/{i:03d}.jpg": png for i in range(10)})
    rid = "7"
    zips = [{"key": "big.zip", "size": len(big)}, {"key": "small.zip", "size": len(small)}]

    def run(dest):
        c = _client(tmp_path, _FakeZenodo(rid, {"big.zip": big, "small.zip": small}))
        return remote.sample_remote_zips(c, rid, zips, dest, 20, random.Random("42:7"), workers=2)

    r1, r2 = run(tmp_path / "r1"), run(tmp_path / "r2")
    assert r1.n_images_listed == 100 and len(r1.sampled) == 20 and len(r1.fetched) == 20
    by_zip = {}
    for _, disp in r1.fetched:
        by_zip[disp.split("!/")[0]] = by_zip.get(disp.split("!/")[0], 0) + 1
    assert by_zip == {"big.zip": 18, "small.zip": 2}
    assert sorted(d for _, d in r1.fetched) == sorted(d for _, d in r2.fetched)   # seeded
    assert len(r1.unfetched_image_paths) == 80
    assert r1.n_requests == 2 + 20          # two listings + one request per member
    assert all(p.is_relative_to(tmp_path / "r1") and p.exists() for p, _ in r1.fetched)


def test_sample_remote_zips_opens_nested_archives(tmp_path):
    import random

    from envision_eye_actionable.survey import remote
    png = _png_bytes(size=(40, 30))
    inner = [_zip_bytes({f"p{j}/{i}.png": png for i in range(4)}) for j in range(3)]
    outer = _zip_bytes({f"patient_{j}.zip": inner[j] for j in range(3)})
    rid = "8"
    c = _client(tmp_path, _FakeZenodo(rid, {"outer.zip": outer}))
    (tmp_path / "w").mkdir()
    w = RecordWalker(scratch=tmp_path / "w")

    def on_nested(p, d):
        n0 = len(w.entries)
        w.add_file(p, d)
        return len(w.entries) - n0

    res = remote.sample_remote_zips(c, rid, [{"key": "outer.zip", "size": len(outer)}], tmp_path / "r", 6,
                                    random.Random(1), nested_max=5, on_nested=on_nested)
    assert res.n_images_listed == 0 and res.n_nested_listed == 3
    assert res.n_nested_fetched == 2 and res.nested_images == 8     # stops once the cap is reached
    assert len(w.entries) == 8 and all("!/" in e.display for e in w.entries)


class _FakeClf:
    """Every image is IR with p=0.9."""

    def predict(self, batch):
        p = np.full((len(batch), 7), 0.1 / 6, np.float32)
        p[:, 1] = 0.9
        return p


def _survey(tmp_path, legacy: dict, files: dict, rid: str, **cfg_kw):
    from envision_eye_actionable.survey.runner import Survey, SurveyConfig
    meta = tmp_path / "meta"
    meta.mkdir(exist_ok=True)
    (meta / f"{rid}.json").write_text(json.dumps(legacy), encoding="utf-8")
    (tmp_path / "dl").mkdir(exist_ok=True)
    cfg_kw.setdefault("triage_cap", 0)          # one pass, unless a test asks for the two-pass sampling
    cfg = SurveyConfig(scrape_path=tmp_path / "s.json", model_path=tmp_path / "m.onnx", out_dir=tmp_path / "out",
                       downloads_dir=tmp_path / "dl", metadata_dir=meta, scratch_dir=tmp_path / "scratch",
                       check_links=False, disk_floor_gb=0, **cfg_kw)
    cfg.out_dir.mkdir(exist_ok=True)
    s = Survey.__new__(Survey)
    s.cfg = cfg
    s._guard_paths()
    s._claim_scratch()
    s.results_path = cfg.out_dir / "survey_results.jsonl"
    s.events_path = cfg.out_dir / "survey_events.jsonl"
    s.zenodo = _client(tmp_path / "out", _FakeZenodo(rid, files))
    s.zenodo.metadata_dir = meta
    s.links = None
    s.clf = _FakeClf()
    s.model_info = {}
    return s


def _legacy(rid, files: dict):
    import hashlib
    from envision_eye_actionable.survey.zenodo import API
    return {"metadata": {"access_right": "open", "title": "Topcon Triton infrared images"},
            "files": [{"key": k, "size": len(v) if isinstance(v, bytes) else v,
                       "checksum": "md5:" + hashlib.md5(v).hexdigest() if isinstance(v, bytes) else "",
                       "links": {"self": f"{API}/{rid}/files/{k}/content"}} for k, v in files.items()]}


def test_process_record_samples_remote_zips(tmp_path):
    pytest.importorskip("jsonschema")
    rid = "555"
    png = _png_bytes(size=(64, 48), color=(90, 90, 90))
    data = _zip_bytes({f"ir/OD_{i:02d}.png": png for i in range(12)} | {"notes.txt": b"x"})
    files = {"ir.zip": data}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=5)
    row = s.process_record({"source_id": rid, "title": "Topcon Triton infrared images"})
    assert row["status"] == "ok", row.get("error")
    assert row["sampling_mode"] == "remote_zip"
    assert row["n_images_listed"] == 12 and row["n_image_files"] == 12
    assert row["n_remote_sampled"] == 5 and row["n_remote_fetched"] == 5 and row["n_classified"] == 5
    assert row["remote_listing_sources"] == "container" and row["remote_requests"] == 6
    assert row["argmax_IR"] == 5 and row["argmax_frac_IR"] == 1.0 and row["argmax_dominant_class"] == "IR"
    assert row["n_IR"] == 5 and row["frac_IR"] == 1.0
    assert row["_class_detail"]["IR"]["est_files"] == 12        # scaled to the listed population
    assert row["laterality_R"] == 12                            # names of unfetched members count too
    assert "Topcon" in row["ir_device_hint"] and row["ir_non_spectralis_candidate"] is True
    assert row["dd_valid"] and row["dsd_valid"]
    assert not (s.records_scratch / rid).exists()              # scratch removed


def test_process_record_samples_top_level_images_and_skips_large_downloads(tmp_path):
    rid = "556"
    png = _png_bytes(size=(64, 48))
    files = {f"img_{i}.png": png for i in range(8)}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=3)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["sampling_mode"] == "download" and row["n_files_to_download"] == 3
    assert row["n_toplevel_images_not_fetched"] == 5
    assert row["n_images_listed"] == 8 and row["n_classified"] == 3

    rid = "557"
    big = {"scans.rar": 20_000_000_000, "readme.pdf": 1000}
    (tmp_path / "b").mkdir()
    s = _survey(tmp_path / "b", _legacy(rid, big), {}, rid)
    row = s.process_record({"source_id": rid, "title": "Big", "external_links": ["https://example.org/data"]})
    assert row["status"] == "skipped_size" and row["sampling_mode"] == "none"
    assert [f["key"] for f in row["files_list"]] == ["scans.rar", "readme.pdf"]
    assert row["skipped_size_files"][0]["key"] == "scans.rar"
    assert "n_weblinks" in row and row["cmds_note"].startswith("skipped_size")
    assert row["dd_valid"]


def test_ir_device_fields():
    from envision_eye_actionable.survey.dicom_map import ir_device_fields
    base = {"n_classified": 100, "frac_IR": 0.25, "dominant_class": "CFP"}
    assert ir_device_fields({**base, "manufacturer_hint_text": "Topcon [topcon]"})["ir_non_spectralis_candidate"]
    h = ir_device_fields({**base, "manufacturer_hint_text": "Heidelberg Engineering Spectralis [spectralis]"})
    assert not h["ir_non_spectralis_candidate"] and "Heidelberg" in h["ir_device_hint"]
    assert ir_device_fields(base) == {"ir_device_hint": "no manufacturer hint found",
                                      "ir_non_spectralis_candidate": False}
    assert ir_device_fields({"n_classified": 100, "frac_IR": 0.1, "argmax_frac_IR": 0.19,
                             "dominant_class": "CFP"})["ir_device_hint"] == ""
    dom = ir_device_fields({"n_classified": 10, "frac_IR": 0.1, "argmax_dominant_class": "IR",
                            "ir_Manufacturer": "Zeiss", "ir_Manufacturer_source": "exif"})
    assert dom["ir_non_spectralis_candidate"] and "Zeiss (IR images, exif)" in dom["ir_device_hint"]
    # the record-level maker of other images is not IR evidence
    other = ir_device_fields({**base, "dicom_Manufacturer": "Carl Zeiss", "dicom_Manufacturer_source": "dicom_header"})
    assert other == {"ir_device_hint": "no manufacturer hint found", "ir_non_spectralis_candidate": False}
    # text hints are labelled as record-level evidence
    assert "(record-level text)" in ir_device_fields({**base, "manufacturer_hint_text": "Topcon [topcon]",
                                                      "manufacturer_hint_source": "text"})["ir_device_hint"]
    # 97 NEG + 3 IR: IR wins the dominant *eye* class, but that is no IR trigger
    mostly_neg = {"n_classified": 100, "frac_IR": 0.03, "argmax_frac_IR": 0.03, "dominant_class": "NEG",
                  "argmax_dominant_class": "NEG", "dominant_eye_class": "IR",
                  "manufacturer_hint_text": "Topcon [topcon]", "manufacturer_hint_source": "text"}
    assert ir_device_fields(mostly_neg) == {"ir_device_hint": "", "ir_non_spectralis_candidate": False}


def test_ir_maker_comes_from_ir_images_only():
    from envision_eye_actionable.survey.dicom_map import aggregate
    zeiss_oct = [_img("OCT", f"oct/{i}.dcm", dicom={"Manufacturer": "Carl Zeiss Meditec", "Modality": "OPT"})
                 for i in range(3)]
    ir = [_img("IR", f"ir/{i}.png") for i in range(1)]
    agg = aggregate(zeiss_oct + ir, [im["path"] for im in zeiss_oct + ir], ["OCT", "IR"], "OCT", "", [])
    assert agg["dicom_Manufacturer"] == "Carl Zeiss Meditec"
    assert agg["ir_Manufacturer"] == "" and agg["ir_Manufacturer_source"] == ""
    ir_topcon = [_img("IR", f"ir/{i}.jpg", make="TOPCON", model="Triton") for i in range(2)]
    agg = aggregate(zeiss_oct + ir_topcon, [], ["OCT", "IR"], "OCT", "", [])
    assert agg["ir_Manufacturer"] == "TOPCON" and agg["ir_Manufacturer_source"] == "exif"
    assert agg["ir_ManufacturerModelName"] == "Triton"


def test_dominant_ties_are_the_same_in_run_and_backfill():
    from envision_eye_actionable.survey.constants import CLASSES, dominant_of
    from envision_eye_actionable.survey.excel import backfill
    assert dominant_of({"IR": 2, "CFP": 2}) == "CFP"          # first in CLASSES, whatever the dict order
    assert dominant_of({"OCT": 0.5, "UNCERTAIN": 0.5}) == "OCT"
    assert dominant_of({}) is None
    row = backfill({"n_classified": 4, "argmax_IR": 2, "argmax_CFP": 2})
    assert row["argmax_dominant_class"] == dominant_of({"IR": 2, "CFP": 2}, CLASSES) == "CFP"


def test_placeholders_are_the_least_assertive_schema_values():
    pytest.importorskip("jsonschema")
    dd, prov = _dd(_datacite_min(), scrape={"source_id": "1", "title": "T", "access_type": "open"})
    de, co = dd["datasetDeIdentLevel"], dd["datasetConsent"]
    # de-identification (project decision): public Zenodo image release, so
    # direct identifiers are taken as removed; the methods are not reported
    assert de["deIdentType"] == "DeIdentificationApplied"
    assert de["deIdentDirect"] is True
    assert [k for k, v in de.items() if isinstance(v, bool) and not v] == [
        "deIdentHIPAA", "deIdentDates", "deIdentNonarr", "deIdentKAnon"]
    details = de["deIdentDetails"]
    assert "publicly released on zenodo as image files" in details.lower()
    assert "not reported by the depositor" in details
    assert "not reported, not confirmed absent" in details
    assert prov["datasetDeIdentLevel"].startswith("derived (public release on Zenodo")
    # consent stays the least assertive placeholder
    assert co["consentType"] == "ConsentSpecifiedNotElsewhereCategorised"
    assert not any(v for k, v in co.items() if isinstance(v, bool))
    assert co["consentsDetails"].startswith("Not reported by source")
    assert prov["datasetConsent"] == "placeholder (not reported by source)"
    assert cmds.validate(dd, "dataset_description") == []


@pytest.mark.parametrize("access,summary", [("closed", {}), ("restricted", {}),
                                            ("open", {"n_images": 0, "files_ext_counts": {".pdf": 1}})])
def test_deident_text_follows_the_access_right_and_the_images(access, summary):
    pytest.importorskip("jsonschema")
    scrape = {"source_id": "1", "title": "T", "access_type": access}
    dd, prov = cmds.build_dataset_description(scrape, {"metadata": {"access_right": access}}, _datacite_min(),
                                              {"present_classes": [], **summary})
    de = dd["datasetDeIdentLevel"]
    # the enum and the booleans are the project default everywhere
    assert {k: v for k, v in de.items() if k != "deIdentDetails"} == {
        k: v for k, v in cmds.DEIDENT_PUBLIC_RELEASE.items() if k != "deIdentDetails"}
    # the text no longer claims a public image release
    assert "publicly released on zenodo as image files" not in de["deIdentDetails"].lower()
    assert f"access_right: {access}" in de["deIdentDetails"] and "no image files read" in de["deIdentDetails"]
    assert "project default" in de["deIdentDetails"]
    assert prov["datasetDeIdentLevel"].startswith(f"derived (Zenodo deposit, access_right {access}")
    if access in ("restricted", "closed"):
        assert "by the depositor" in de["deIdentDetails"]
        assert "de-identification is not reported" in de["deIdentDetails"]
        assert "de-identification not reported" in prov["datasetDeIdentLevel"]
    else:
        assert "de-identification is not reported" not in de["deIdentDetails"]
    if access == "closed":
        assert "not available from Zenodo" in dd["accessDetails"]["description"]
    assert cmds.validate(dd, "dataset_description") == []
    assert cmds.CMDS_NAME == "Clinical Multimodal Data Structure (CMDS) v0.1.1"


def test_workbook_has_argmax_sampling_and_ir_columns(tmp_path):
    pytest.importorskip("openpyxl")
    from envision_eye_actionable.survey.excel import build_workbook
    from openpyxl import load_workbook
    out = tmp_path / "res"
    out.mkdir()
    new = {"record_id": "1", "status": "ok", "sampling_mode": "remote_zip", "n_classified": 10, "n_IR": 6,
           "frac_IR": 0.6, "argmax_IR": 8, "argmax_frac_IR": 0.8, "dominant_class": "IR",
           "ir_device_hint": "Topcon [topcon] (text)", "ir_non_spectralis_candidate": True,
           "n_images_listed": 400, "n_image_files": 400, "n_remote_sampled": 10}
    old = {"record_id": "2", "status": "ok", "n_classified": 4, "n_files_local": 1, "argmax_IR": 3,
           "argmax_CFP": 1, "frac_IR": 0.5, "dominant_class": "IR", "manufacturer_hint_text": "Canon CR-2 [canon]",
           "n_image_files": 4}
    skipped = {"record_id": "3", "status": "skipped_size", "sampling_mode": "none",
               "skipped_size_files": [{"key": "a.rar", "size": 2e10, "url": "https://zenodo.org/x"}]}
    res = out / "survey_results.jsonl"
    res.write_text("\n".join(json.dumps(r) for r in (new, old, skipped)) + "\n", encoding="utf-8")
    build_workbook(res, out / "w.xlsx")
    ws = load_workbook(out / "w.xlsx")["Records"]
    head = [c.value for c in ws[1]]
    for col in ("Sampling mode (local, remote_zip, download)", "argmax n IR", "argmax frac IR", "frac IR",
                "Dominant class (argmax, no threshold)", "IR device hint (IR dominant or >= 20% of non-mask images)",
                "IR non-Spectralis candidate", "Images seen in listings", "Remote members sampled",
                "Files over --max-download-gb (key, size, link)"):
        assert col in head, col
    assert head.index("argmax n IR") < head.index("frac CFP") < head.index("argmax frac IR")
    rows = {r[0].value: {h: c.value for h, c in zip(head, r)} for r in ws.iter_rows(min_row=2)}
    assert rows["1"]["argmax frac IR"] == 0.8 and rows["1"]["IR non-Spectralis candidate"] is True
    # rows written before these columns existed are backfilled
    assert rows["2"]["argmax frac IR"] == 0.75 and rows["2"]["Sampling mode (local, remote_zip, download)"] == "local"
    assert rows["2"]["IR non-Spectralis candidate"] is True and "Canon" in rows["2"]["IR device hint (IR dominant or >= 20% of non-mask images)"]
    assert "a.rar (20.0 GB) https://zenodo.org/x" in rows["3"]["Files over --max-download-gb (key, size, link)"]


# ---------------------------------------------------------------------------
# third review round: memory bounds, resume safety, weighting, provenance
# ---------------------------------------------------------------------------
class _BrokenBody(_HttpResp):
    @property
    def content(self):
        import requests
        raise requests.exceptions.ChunkedEncodingError("connection broken mid-body")


def test_range_read_retries_a_broken_body_and_bounds_its_cache(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import remote
    monkeypatch.setattr(remote.time, "sleep", lambda s: None)
    data = bytes(range(256)) * 4096                      # 1 MB

    class Flaky:
        def __init__(self):
            self.n = 0

        def get(self, url, headers=None, stream=False, timeout=None):
            self.n += 1
            a, b = (int(x) for x in headers["Range"].split("=")[1].split("-"))
            if self.n == 1:
                return _BrokenBody(206)
            return _HttpResp(206, data[a:b + 1])

    s = Flaky()
    fp = remote.HttpRangeFile(_client(tmp_path, s), "u", len(data), block=64 << 10, max_cache_bytes=200 << 10)
    fp.seek(1000)
    assert fp.read(10) == data[1000:1010] and s.n == 2           # retried once
    # a read bigger than the cache bound is returned but never kept
    fp.seek(0)
    assert fp.read(len(data)) == data
    assert fp._cache_bytes <= 200 << 10 and all(len(b) <= 200 << 10 for b in fp._cache.values())
    for i in range(10):                                           # many blocks: bytes stay bounded
        fp.seek(i * 70_000)
        fp.read(100)
    assert fp._cache_bytes <= 200 << 10
    # prefetch never asks for more than PREFETCH_MAX_BYTES in one request
    monkeypatch.setattr(remote, "PREFETCH_MAX_BYTES", 128 << 10)
    seen = []
    orig = fp._fetch
    fp._fetch = lambda start, length: seen.append(length) or orig(start, length)
    fp.close()
    fp.prefetch(0, len(data))
    assert seen and max(seen) <= 128 << 10


def test_range_member_fetch_streams_stored_and_deflated_members(tmp_path):
    from envision_eye_actionable.survey import remote
    rnd = np.random.default_rng(0).integers(0, 256, 3 << 20, dtype=np.uint8).tobytes()   # 3 MB, incompressible
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(zipfile.ZipInfo("a/stored.tif"), rnd, compress_type=zipfile.ZIP_STORED)
        zf.writestr("a/deflated.tif", b"\x00" * (2 << 20), compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("a/empty.tif", b"", compress_type=zipfile.ZIP_DEFLATED)
    data = buf.getvalue()
    rid = "60"
    fake = _FakeZenodo(rid, {"z.zip": data}, member_status=404)     # container refuses every member
    c = _client(tmp_path, fake)
    lst = remote.list_zip(c, rid, "z.zip", len(data))
    rz = remote.RangeZips(c, rid)
    try:
        for m in lst.members:
            dest = tmp_path / "out" / m.name
            ok, err = remote.fetch_member(c, rid, lst, m, dest, len(data), rz)
            assert ok, err
        assert (tmp_path / "out/a/stored.tif").read_bytes() == rnd
        assert (tmp_path / "out/a/deflated.tif").read_bytes() == b"\x00" * (2 << 20)
        assert (tmp_path / "out/a/empty.tif").read_bytes() == b""
        fp, _ = rz.get("z.zip", len(data))
        assert all(len(b) < 1 << 20 for b in fp._cache.values())     # member data never went into the cache
    finally:
        rz.close()
    # members over the cap are refused before any request
    n_calls = len(fake.calls)
    big = remote.Member("a/stored.tif", 3 << 20)
    ok, err = remote.fetch_member(c, rid, lst, big, tmp_path / "x", len(data), None, max_bytes=1 << 20)
    assert not ok and "cap" in err and len(fake.calls) == n_calls


def test_nested_archives_over_the_cap_are_not_fetched(tmp_path, monkeypatch):
    import random

    from envision_eye_actionable.survey import remote
    png = _png_bytes(size=(40, 30))
    small = _zip_bytes({"p/0.png": png})
    big = _zip_bytes({f"q/{i}.png": _png_bytes(size=(40, 30), color=(i, 0, 0)) for i in range(20)})
    outer = _zip_bytes({"small.zip": small, "big.zip": big})
    small_size = zipfile.ZipFile(io.BytesIO(outer)).getinfo("small.zip").file_size
    fake = _FakeZenodo("9", {"outer.zip": outer})
    c = _client(tmp_path, fake)
    got = []
    res = remote.sample_remote_zips(c, "9", [{"key": "outer.zip", "size": len(outer)}], tmp_path / "r", 50,
                                    random.Random(1), nested_max=5, on_nested=lambda p, d: got.append(d) or 1,
                                    nested_max_bytes=small_size + 10)
    assert res.n_nested_listed == 2 and res.n_nested_fetched == 1 and res.n_nested_over_cap == 1
    assert got == ["outer.zip!/small.zip"] and not res.failures
    assert not any("big.zip" in url for url, _ in fake.calls)


def test_failed_container_listing_skips_container_member_fetches(tmp_path, monkeypatch):
    import random

    from envision_eye_actionable.survey import remote, zenodo
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: None)
    png = _png_bytes(size=(40, 30))
    data = _zip_bytes({f"i/{k}.png": png for k in range(6)})
    fake = _FakeZenodo("11", {"x.zip": data}, listing_status=500)
    c = _client(tmp_path, fake)
    res = remote.sample_remote_zips(c, "11", [{"key": "x.zip", "size": len(data)}], tmp_path / "r", 4,
                                    random.Random(1))
    lst = res.listings[0]
    assert lst.container_failed and lst.source == "range" and not lst.failed
    assert len(res.fetched) == 4 and not res.failures
    assert not any("/container/" in url for url, _ in fake.calls)       # no per-member container retries
    assert sum(1 for url, _ in fake.calls if url.endswith("/container")) == 3   # the listing's own tries


def test_listing_failures_are_retryable_statuses(tmp_path, monkeypatch):
    pytest.importorskip("jsonschema")
    from envision_eye_actionable.survey import zenodo
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: None)
    rid = "600"
    files = {"bad.zip": b"not a zip at all" * 10}
    s = _survey(tmp_path, _legacy(rid, files), files, rid)
    s.zenodo.session.listing_status = 503
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "remote_listing_failed", row.get("error")
    assert row["remote_listing_failures"] == 1

    rid = "601"
    png = _png_bytes(size=(64, 48))
    good = _zip_bytes({f"ir/OD_{i}.png": png for i in range(3)})
    files = {"good.zip": good, "bad.zip": b"not a zip at all" * 10}
    (tmp_path / "b").mkdir()
    s = _survey(tmp_path / "b", _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["remote_listing_failures"] == 1
    assert row["status"] == "ok_partial_download"


class _BrightDarkClf:
    """CFP for bright images, IR for dark ones (by the normalised mean)."""

    def predict(self, batch):
        p = np.full((len(batch), 7), 0.01, np.float32)
        bright = batch[:, 0].mean(axis=(1, 2)) > 0
        p[bright, 0] = 0.94
        p[~bright, 1] = 0.94
        return p


def test_mixed_local_and_remote_records_are_weighted_by_stratum(tmp_path):
    pytest.importorskip("jsonschema")
    rid = "610"
    bright = _png_bytes(size=(64, 48), color=(220, 220, 220))
    dark = _png_bytes(size=(64, 48), color=(40, 40, 40))
    local = _zip_bytes({f"cfp/{i}.png": bright for i in range(10)})
    remote_zip = _zip_bytes({f"ir/{i:03d}.png": dark for i in range(100)})
    files = {"local.zip": local, "remote.zip": remote_zip}
    s = _survey(tmp_path, _legacy(rid, files), {"remote.zip": remote_zip}, rid, remote_cap=10)
    (s.cfg.downloads_dir / rid).mkdir(parents=True)
    (s.cfg.downloads_dir / rid / "local.zip").write_bytes(local)
    s.clf = _BrightDarkClf()
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["sampling_mode"] == "local+remote_zip", row.get("error")
    assert row["n_classified"] == 20 and row["n_classified_by_source"] == {"files": 10, "remote": 10}
    assert row["n_CFP"] == 10 and row["n_IR"] == 10                     # raw sample counts
    assert row["class_weighting"] == "stratified"
    assert row["frac_IR"] == round(100 / 110, 4) and row["frac_CFP"] == round(10 / 110, 4)
    assert row["argmax_frac_IR"] == round(100 / 110, 4)
    assert row["dominant_class"] == "IR" and row["dominant_eye_class"] == "IR"
    assert row["_class_detail"]["IR"]["est_files"] == 100 and row["_class_detail"]["CFP"]["est_files"] == 10


def test_download_budget_is_filled_smallest_first(tmp_path):
    from envision_eye_actionable.survey.runner import fit_budget
    files = [{"key": "huge.rar", "size": 20_000}, {"key": "a.tif", "size": 10},
             {"key": "set.part1.rar", "size": 400}, {"key": "set.part2.rar", "size": 200},
             {"key": "b.tif", "size": 20}]
    kept, over = fit_budget(files, 700)
    assert [f["key"] for f in kept] == ["a.tif", "set.part1.rar", "set.part2.rar", "b.tif"]
    assert [f["key"] for f in over] == ["huge.rar"]
    kept, over = fit_budget(files, 500)             # the split set (600) never goes half
    assert [f["key"] for f in kept] == ["a.tif", "b.tif"]
    assert {f["key"] for f in over} == {"huge.rar", "set.part1.rar", "set.part2.rar"}

    rid = "620"
    png = _png_bytes(size=(64, 48))
    legacy = _legacy(rid, {"scans.rar": 20_000_000_000, "img_0.png": png, "img_1.png": png})
    s = _survey(tmp_path, legacy, {"img_0.png": png, "img_1.png": png}, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "ok", row.get("error")
    assert row["n_files_skipped_size"] == 1 and row["skipped_size_files"][0]["key"] == "scans.rar"
    assert row["n_files_to_download"] == 2 and row["n_classified"] == 2 and "size_note" in row


def test_throttle_does_not_hold_its_lock_while_sleeping(tmp_path):
    import threading
    import time as _time

    from envision_eye_actionable.survey.zenodo import ZenodoClient
    c = ZenodoClient(tmp_path / "cache", None, min_interval=0, max_per_minute=None)
    c._not_before = _time.monotonic() + 0.4
    done = []
    t0 = _time.monotonic()
    t = threading.Thread(target=lambda: (c.throttle(), done.append(_time.monotonic())))
    t.start()
    _time.sleep(0.05)
    c.backoff(0.8)                                   # a 429 elsewhere while the thread sleeps
    assert _time.monotonic() - t0 < 0.3              # backoff did not wait for the sleeper
    t.join(5)
    assert done and done[0] - t0 >= 0.8              # the sleeper honoured the new cooldown


def test_model_sidecar_hash_must_match(tmp_path):
    import hashlib

    from envision_eye_actionable.survey.classifier import read_sidecar
    m = tmp_path / "m.onnx"
    m.write_bytes(b"model bytes")
    side = tmp_path / "m.json"
    good = {"classes": ["CFP", "IR", "PSC", "FAF", "OCT", "OCTA", "NEG"], "parity_max_abs_diff": 1e-6,
            "onnx_sha256": hashlib.sha256(b"model bytes").hexdigest(), "checkpoint": "c.pt"}
    side.write_text(json.dumps(good), encoding="utf-8")
    assert read_sidecar(m)["checkpoint"] == "c.pt"
    m.write_bytes(b"a re-exported model")             # swapped model, stale sidecar
    with pytest.raises(ValueError, match="does not match its sidecar"):
        read_sidecar(m)
    side.unlink()
    assert read_sidecar(m) == {}                      # no sidecar: warning only


def test_export_refuses_without_onnxruntime(tmp_path, monkeypatch):
    import sys

    from envision_eye_actionable.survey import export_onnx
    monkeypatch.setitem(sys.modules, "onnxruntime", None)      # import fails
    out = tmp_path / "m.onnx"
    with pytest.raises(RuntimeError, match="needs onnxruntime"):
        export_onnx.export(tmp_path / "c.pt", "regnety_004", out)
    assert not out.exists() and not out.with_suffix(".json").exists()


def test_cli_requires_a_model(tmp_path, monkeypatch, capsys):
    from envision_eye_actionable.survey import cli
    monkeypatch.delenv("ENVISION_SURVEY_MODEL", raising=False)
    monkeypatch.setattr(cli, "_default_model", lambda: None)
    with pytest.raises(SystemExit):
        cli.main(["run", "--scrape", str(tmp_path / "s.json")])
    assert "ENVISION_SURVEY_MODEL" in capsys.readouterr().err


# ---------------------------------------------------------------------------
# fourth review round: remote strata, scratch names, nested budget, listing
# checks, IR evidence, review flags, parity through the survey path
# ---------------------------------------------------------------------------
def test_short_container_listing_is_not_trusted(tmp_path):
    from envision_eye_actionable.survey import remote
    short = remote.container_listing_short
    assert short({"entries": [{}] * 5, "truncated": True})
    assert short({"entries": [{}] * 5, "truncated": False, "total": 9})          # total above the page
    assert short({"entries": [{}] * 990, "directories": [{}] * 10, "truncated": False, "total": 1000})
    assert not short({"entries": [{}] * 5, "truncated": False, "total": 5})
    assert not short({"entries": [{}] * 5, "directories": [{}] * 2, "truncated": False, "total": 7})
    assert not short({"entries": [{}] * 5})                                        # no total: trusted
    png = _png_bytes(size=(20, 20))
    data = _zip_bytes({f"s/{i}.png": png for i in range(6)})
    fake = _FakeZenodo("70", {"z.zip": data}, hide_truncation=True)
    lst = remote.list_zip(_client(tmp_path, fake), "70", "z.zip", len(data))
    assert lst.source == "range" and lst.truncated_listing and len(lst.images) == 6
    assert "truncated (total 6)" in remote.RemoteSample(listings=[lst]).detail()


def test_scratch_names_are_flat_unique_and_keep_the_suffix(tmp_path):
    import random

    from envision_eye_actionable.survey import remote
    long_dir = "图" * 120                           # 360 bytes in UTF-8: over the 255-byte name limit
    names = [f"{long_dir}/a.png", "x.png", "x.png/y.png", "a:b.png", "a_b.png", "../x.png", "d/scan",
             "s/v.tar.gz", "s/part.7z.001"]
    paths = [remote.safe_member_path(tmp_path, 0, n, i) for i, n in enumerate(names)]
    assert len(set(paths)) == len(paths)
    assert all(p.parent == tmp_path / "0" and len(p.name.encode()) < 60 for p in paths)
    assert [p.name.split("_", 1)[1][12:] for p in paths] == [
        ".png", ".png", ".png", ".png", ".png", ".png", "", ".tar.gz", ".7z.001"]
    assert file_kind(paths[-1].name) == "archive" and file_kind(paths[-2].name) == "archive"
    # every such member is fetched; none aborts the record
    png = _png_bytes(size=(20, 20))
    data = _zip_bytes({n: png for n in names[:6]})
    c = _client(tmp_path, _FakeZenodo("71", {"z.zip": data}))
    res = remote.sample_remote_zips(c, "71", [{"key": "z.zip", "size": len(data)}], tmp_path / "r", 50,
                                    random.Random(1))
    assert not res.failures and len(res.fetched) == 6
    assert len({p for p, _ in res.fetched}) == 6 and all(p.read_bytes() == png for p, _ in res.fetched)
    assert sorted(d.split("!/", 1)[1] for _, d in res.fetched) == sorted(names[:6])
    # a scratch path problem is a per-member failure, not an exception
    (tmp_path / "file").write_bytes(b"")
    lst = res.listings[0]
    ok, err = remote.fetch_member(c, "71", lst, lst.images[0], tmp_path / "file" / "sub" / "m.png", len(data))
    assert not ok and err.startswith("scratch dir")


def test_nested_fetching_has_its_own_budget_and_stops_on_empty_archives(tmp_path):
    import random

    from envision_eye_actionable.survey import remote
    empty_inner = [_zip_bytes({f"p{j}/notes.txt": b"x" * 50}) for j in range(8)]
    outer = _zip_bytes({f"patient_{j}.zip": empty_inner[j] for j in range(8)})
    zips = [{"key": "outer.zip", "size": len(outer)}]
    c = _client(tmp_path, _FakeZenodo("72", {"outer.zip": outer}))
    res = remote.sample_remote_zips(c, "72", zips, tmp_path / "r", 50, random.Random(1), nested_max=20,
                                    on_nested=lambda p, d: 0, nested_empty_streak=3)
    assert res.n_nested_fetched == 3 and "in a row" in res.nested_stop and res.nested_bytes > 0
    one = zipfile.ZipFile(io.BytesIO(outer)).getinfo("patient_0.zip").file_size
    c = _client(tmp_path, _FakeZenodo("72", {"outer.zip": outer}))
    res = remote.sample_remote_zips(c, "72", zips, tmp_path / "r2", 50, random.Random(1), nested_max=20,
                                    on_nested=lambda p, d: 1, nested_budget_bytes=int(one * 2.5))
    assert res.n_nested_fetched == 2 and res.n_nested_over_budget == 6
    assert res.nested_stop == "nested byte budget used up" and res.nested_bytes <= one * 2.5


class _BrightNegDarkOctClf:
    """NEG for bright images, OCT for dark ones."""

    def predict(self, batch):
        p = np.full((len(batch), 7), 0.01, np.float32)
        bright = batch[:, 0].mean(axis=(1, 2)) > 0
        p[bright, 6] = 0.94
        p[~bright, 4] = 0.94
        return p


class _MaskOctaClf(_BrightNegDarkOctClf):
    """As _BrightNegDarkOctClf, but noise-free inputs (a binary mask comes
    out flat once content-cropped to its foreground) are OCTA, as missed
    masks are with the real model."""

    def predict(self, batch):
        p = super().predict(batch)
        hi = batch[:, 0].std(axis=(1, 2)) < 0.01
        p[hi] = 0.01
        p[hi, 5] = 0.94
        return p


def test_remote_direct_and_nested_images_are_separate_strata(tmp_path):
    pytest.importorskip("jsonschema")
    rid = "630"
    bright = _png_bytes(size=(64, 48), color=(220, 220, 220))
    dark = _png_bytes(size=(64, 48), color=(40, 40, 40))
    inner = [_zip_bytes({f"pt{j}/{i}.png": dark for i in range(3)}) for j in range(10)]
    outer = _zip_bytes({f"figures/f{i}.png": bright for i in range(5)}
                       | {f"patients/pt{j}.zip": inner[j] for j in range(10)})
    files = {"data.zip": outer}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=50, remote_nested_max=2)
    s.clf = _BrightNegDarkOctClf()
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["sampling_mode"] == "remote_zip", row.get("error")
    assert row["n_remote_nested_fetched"] == 2 and row["n_remote_nested_images"] == 6
    assert row["n_image_files"] == 35                                  # 5 direct + 3 x 10 nested (estimated)
    assert row["n_classified_by_source"] == {"remote": 5, "remote_nested": 6}
    assert row["class_weighting"] == "stratified"
    assert row["n_NEG"] == 5 and row["n_OCT"] == 6                     # raw sample counts
    assert row["frac_OCT"] == round(30 / 35, 4) and row["frac_NEG"] == round(5 / 35, 4)
    assert row["argmax_frac_OCT"] == round(30 / 35, 4)
    assert row["_class_detail"]["OCT"]["est_files"] == 30
    assert row["remote_nested_bytes"] > 0


def test_ir_evidence_includes_below_threshold_ir_images():
    from envision_eye_actionable.survey.dicom_map import aggregate
    unsure_ir = [{**_img("UNCERTAIN", f"ir/{i}.jpg", make="TOPCON", model="Triton"), "top": "IR"}
                 for i in range(3)]
    cfp = [{**_img("CFP", f"cfp/{i}.jpg", make="Canon"), "top": "CFP"} for i in range(5)]
    agg = aggregate(unsure_ir + cfp, [], ["CFP"], "CFP", "", [])
    assert agg["ir_Manufacturer"] == "TOPCON" and agg["ir_Manufacturer_source"] == "exif"
    assert agg["ir_ManufacturerModelName"] == "Triton"


def test_strong_ir_maker_decides_the_non_spectralis_flag():
    from envision_eye_actionable.survey.dicom_map import ir_device_fields
    base = {"n_classified": 100, "frac_IR": 0.3, "dominant_class": "IR",
            "manufacturer_hint_text": "Heidelberg Engineering Spectralis [spectralis]",
            "manufacturer_hint_source": "text"}
    topcon = ir_device_fields({**base, "ir_Manufacturer": "TOPCON", "ir_Manufacturer_source": "dicom_header"})
    assert topcon["ir_non_spectralis_candidate"] is True
    assert "TOPCON (IR images, dicom_header)" in topcon["ir_device_hint"]
    assert "Spectralis" in topcon["ir_device_hint"]                  # the weak hint is still listed
    heid = ir_device_fields({**base, "manufacturer_hint_text": "Topcon [topcon]",
                             "ir_Manufacturer": "Heidelberg Engineering", "ir_Manufacturer_source": "exif"})
    assert heid["ir_non_spectralis_candidate"] is False
    assert ir_device_fields(base)["ir_non_spectralis_candidate"] is False        # weak hints only


def test_modality_review_flags_known_false_positive_patterns():
    pytest.importorskip("jsonschema")
    from envision_eye_actionable.survey.dicom_map import modality_review
    microscopy = {"present_eye_classes": "OCTA", "dominant_eye_class": "OCTA", "setfit_label": "NEGATIVE",
                  "setfit_prob_eye_imaging": 0.011, "mask_like_frac": 0.897,
                  "_class_detail": {"OCTA": {"count": 250, "est_files": 13000, "mean_conf": 0.71}}}
    r = modality_review(microscopy)
    assert r["modality_low_trust"] is True
    assert [f.split(" ")[0] for f in r["modality_review_flags"].split("; ")] == [
        "setfit_not_eye", "mask_like", "low_confidence"]
    fundus = {"present_eye_classes": "CFP", "dominant_eye_class": "CFP", "setfit_label": "EYE_IMAGING",
              "setfit_prob_eye_imaging": 0.98, "mask_like_frac": 0.0,
              "_class_detail": {"CFP": {"count": 800, "est_files": 800, "mean_conf": 0.82}}}
    assert modality_review(fundus) == {"modality_review_flags": "", "modality_low_trust": False}
    assert modality_review({"present_eye_classes": "", "setfit_prob_eye_imaging": 0.0})["modality_low_trust"] is False
    dsd = cmds.build_structure_description(microscopy["_class_detail"], ["OCTA"], {
        "n_images": 15000, "n_classified": 300, "review_note": r["modality_review_flags"]})
    assert cmds.validate(dsd, "dataset_structure_description") == []
    octa = dsd["directoryList"][0]
    assert "REVIEW" in octa["directoryDescription"]
    assert "setfit_not_eye" in octa["directoryList"][0]["directoryDescription"]


def test_workbook_reports_review_flags_and_parity(tmp_path):
    pytest.importorskip("openpyxl")
    from envision_eye_actionable.survey.excel import build_workbook
    from openpyxl import load_workbook
    out = tmp_path / "res"
    out.mkdir()
    old = {"record_id": "1", "status": "ok", "n_classified": 300, "present_eye_classes": "OCTA",
           "dominant_eye_class": "OCTA", "setfit_prob_eye_imaging": 0.01, "setfit_label": "NEGATIVE",
           "mask_like_frac": 0.9, "_class_detail": {"OCTA": {"count": 250, "est_files": 1, "mean_conf": 0.7}}}
    res = out / "survey_results.jsonl"
    res.write_text(json.dumps(old) + "\n", encoding="utf-8")
    meta = {"parity_max_abs_diff": 1.2e-6,
            "parity_real_images": {"n": 21, "max_logit_diff": 2.4e-6, "argmax_agree": 21},
            "parity_survey_preprocess": {"n": 21, "max_logit_diff": 5.5e-6, "argmax_agree": 21}}
    build_workbook(res, out / "w.xlsx", meta)
    wb = load_workbook(out / "w.xlsx")
    readme = {r[0].value: r[1].value for r in wb["README"].iter_rows()}
    assert "survey decoding and preprocessing: n 21" in readme["Model ONNX parity"]
    assert "OCTA" in readme["Eye modality review"]
    ws = wb["Records"]
    head = [c.value for c in ws[1]]
    row = {h: c.value for h, c in zip(head, next(ws.iter_rows(min_row=2)))}
    assert row["Eye modality low trust (review)"] is True                 # backfilled for an older row
    assert "setfit_not_eye" in row["Eye modality review flags"]


def test_export_parity_can_use_the_survey_preprocessing_standalone():
    import importlib.util
    from pathlib import Path

    from envision_eye_actionable.survey import export_onnx
    assert hasattr(export_onnx._survey_images_module(), "preprocess")
    # run as a standalone file (no package): images.py is loaded from next to it
    spec = importlib.util.spec_from_file_location("standalone_export_onnx", Path(export_onnx.__file__))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    images = mod._survey_images_module()
    assert images is not None and hasattr(images, "load_frame")
    arr = images.preprocess(Image.new("RGB", (300, 200), (120, 30, 10)))
    assert arr.shape == (3, 224, 224)


# ---------------------------------------------------------------------------
# fourth round: de-identification, byte bounds, MASK label, eye status,
# partitions, several processes
# ---------------------------------------------------------------------------
def test_mask_detector_is_conservative():
    from envision_eye_actionable.survey.images import is_mask_like
    rng = np.random.default_rng(3)
    # binary and 3-level label maps, a 5-color label map
    assert is_mask_like(Image.open(io.BytesIO(_mask_png())))
    assert is_mask_like(Image.open(io.BytesIO(_mask_png(levels=(0, 100, 200)))))
    lab = np.zeros((90, 90, 3), np.uint8)
    for k, c in enumerate([(255, 0, 0), (0, 255, 0), (0, 0, 255), (255, 255, 0), (0, 255, 255)]):
        lab[k * 15:(k + 1) * 15] = c
    assert is_mask_like(Image.fromarray(lab))
    # a saved gray mask with chroma residue (+-1 in one channel) is still a mask
    m = np.asarray(Image.open(io.BytesIO(_mask_png())).convert("RGB")).astype(np.int16)
    m[..., 0] += rng.integers(0, 2, m.shape[:2])
    assert is_mask_like(Image.fromarray(np.clip(m, 0, 255).astype(np.uint8)))
    # dark, low-contrast grayscale (an IR or FAF frame): 8 levels is not a mask
    dark = (rng.integers(0, 8, (200, 200)) + 10).astype(np.uint8)
    assert not is_mask_like(Image.fromarray(dark, "L"))
    # a mostly black OCT-like frame with a textured band
    oct_like = np.zeros((200, 300), np.uint8)
    oct_like[80:120] = rng.integers(40, 200, (40, 300))
    assert not is_mask_like(Image.fromarray(oct_like, "L"))
    # a JPEG-saved mask gains many values at its edges, but its flat regions
    # still dominate: flagged by the dominance test
    buf = io.BytesIO()
    Image.open(io.BytesIO(_mask_png(size=(256, 256)))).save(buf, format="JPEG", quality=70)
    assert is_mask_like(Image.open(io.BytesIO(buf.getvalue())))
    # a JPEG of a textured gray image is not
    buf = io.BytesIO()
    Image.open(io.BytesIO(_png_bytes(size=(256, 256), mode="L"))).save(buf, format="JPEG", quality=70)
    assert not is_mask_like(Image.open(io.BytesIO(buf.getvalue())))


def test_masks_get_the_mask_label_not_a_modality(tmp_path):
    pytest.importorskip("jsonschema")
    import gzip
    rid = "700"
    photo = _png_bytes(size=(64, 48), color=(90, 90, 90))
    members = {f"img/{i:02d}.png": photo for i in range(6)} | {f"seg/{i:02d}.png": _mask_png() for i in range(4)}
    files = {"data.zip": _zip_bytes(members)}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=50)
    row = s.process_record({"source_id": rid, "title": "Topcon infrared images"})   # _FakeClf: all IR 0.9
    assert row["status"] == "ok", row.get("error")
    assert row["n_classified"] == 10 and row["n_classified_non_mask"] == 6
    assert row["n_MASK"] == 4 and row["n_IR"] == 6 and row["argmax_MASK"] == 4 and row["argmax_IR"] == 6
    assert row["frac_MASK"] == 0.4 and row["argmax_frac_MASK"] == 0.4 and row["frac_IR"] == 0.6
    assert row["dominant_class"] == "IR" and row["argmax_dominant_class"] == "IR"
    assert row["eye_image_fraction"] == 1.0 and row["present_eye_classes"] == "IR"
    assert row["mask_like_frac"] == 0.4
    preds = [json.loads(x) for x in gzip.open(s.cfg.out_dir / "predictions" / f"{rid}.jsonl.gz", "rt")]
    masks = [p for p in preds if p.get("mask_like")]
    assert len(masks) == 4 and all(p["label"] == "MASK" and p["top"] == "IR" and p["model_label"] == "IR"
                                   for p in masks)
    assert all(p["label"] == "IR" and "mask_like" not in p for p in preds if not p.get("mask_like"))

    # a record of masks only: dominant MASK, no eye images (mask-dominated)
    rid = "701"
    files = {"seg.zip": _zip_bytes({f"seg/{i}.png": _mask_png() for i in range(5)})}
    (tmp_path / "b").mkdir()
    s = _survey(tmp_path / "b", _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "mask_dominated", row.get("error")
    assert row["mask_dominated"] is True and row["review_eye_classes"] == ""
    assert row["dominant_class"] == "MASK" and row["argmax_dominant_class"] == "MASK"
    assert row["n_classified_non_mask"] == 0 and row["eye_image_fraction"] is None
    assert row["present_eye_classes"] == "" and row["mean_confidence"] is None
    assert row["dd_valid"] and row["dsd_valid"] and "n_weblinks" in row


@pytest.mark.parametrize("n_eye,status", [(2, "no_eye_images"), (3, "ok")])
def test_eye_status_needs_the_min_eye_fraction_of_non_mask_images(tmp_path, n_eye, status):
    pytest.importorskip("jsonschema")
    rid = "710"
    bright = _png_bytes(size=(64, 48), color=(220, 220, 220))     # NEG
    dark = _png_bytes(size=(64, 48), color=(40, 40, 40))          # OCT
    members = ({f"neg/{i:02d}.png": bright for i in range(50 - n_eye)}
               | {f"oct/{i:02d}.png": dark for i in range(n_eye)}
               | {f"seg/{i:02d}.png": _mask_png() for i in range(30)})   # masks do not dilute the fraction
    files = {"d.zip": _zip_bytes(members)}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=100)
    s.clf = _MaskOctaClf()
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == status, row.get("error")
    assert row["eye_image_fraction"] == round(n_eye / 50, 4)
    assert row["dominant_class"] == "NEG" and row["n_MASK"] == 30
    assert row["present_eye_classes"] == ("OCT" if status == "ok" else "")
    assert row["min_eye_fraction"] == 0.05


class _UncertainIrClf:
    """Argmax IR for every image, but under the 0.6 threshold (UNCERTAIN)."""

    def predict(self, batch):
        p = np.full((len(batch), 7), 0.5 / 6, np.float32)
        p[:, 1] = 0.5
        return p


def test_no_eye_record_gets_no_ir_device_fields(tmp_path):
    pytest.importorskip("jsonschema")
    from envision_eye_actionable.survey.dicom_map import ir_device_fields
    rid = "715"
    png = _png_bytes(size=(64, 48), color=(90, 90, 90))
    files = {"d.zip": _zip_bytes({f"img/{i:02d}.png": png for i in range(10)})}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=50)
    s.clf = _UncertainIrClf()
    row = s.process_record({"source_id": rid, "title": "Topcon Triton infrared images"})
    assert row["status"] == "no_eye_images", row.get("error")
    assert row["argmax_frac_IR_non_mask"] == 1.0 and row["argmax_dominant_class"] == "IR"
    assert row["eye_image_fraction"] == 0.0 and row["present_eye_classes"] == ""
    assert row["ir_device_hint"] == "" and row["ir_non_spectralis_candidate"] is False
    # the same shares with an eye class present do trigger (the gate is the eye decision)
    fired = ir_device_fields({**row, "present_eye_classes": "IR", "manufacturer_hint_text": "Topcon [topcon]"})
    assert fired["ir_non_spectralis_candidate"] is True


def test_remote_members_over_the_cap_are_replaced_and_the_record_budget_stops_fetching(tmp_path):
    import random

    from envision_eye_actionable.survey import remote
    small = {f"s/{i:02d}.png": _png_bytes(size=(40, 30), color=(i * 9, 50, 50)) for i in range(20)}
    big = {f"b/{i}.png": _png_bytes(size=(400, 300), color=(i * 9, 80, 80)) for i in range(5)}
    data = _zip_bytes(small | big)
    zips = [{"key": "x.zip", "size": len(data)}]
    # the cap applies to the bytes a fetch moves (compressed size; deflate
    # adds a few bytes to an incompressible PNG)
    cap_bytes = max(len(v) for v in small.values()) + 64
    assert min(len(v) for v in big.values()) > cap_bytes
    c = _client(tmp_path, _FakeZenodo("80", {"x.zip": data}))
    res = remote.sample_remote_zips(c, "80", zips, tmp_path / "r", 10, random.Random(1),
                                    max_member_bytes=cap_bytes)
    assert res.n_members_skipped_oversize == 5
    assert len(res.fetched) == 10 and all(d.split("!/")[1].startswith("s/") for _, d in res.fetched)
    assert res.bytes_fetched == sum(len(small[d.split("!/")[1]]) for _, d in res.fetched)
    assert res.n_members_over_budget == 0 and res.budget_stop == ""

    budget = int(cap_bytes * 3.5)
    c = _client(tmp_path, _FakeZenodo("80", {"x.zip": data}))
    res = remote.sample_remote_zips(c, "80", zips, tmp_path / "r2", 10, random.Random(1),
                                    max_member_bytes=cap_bytes, record_budget_bytes=budget)
    assert 3 <= len(res.fetched) < 10 and res.bytes_fetched <= budget
    assert res.n_members_over_budget == 10 - len(res.fetched) and "budget" in res.budget_stop
    assert len(res.sampled) == len(res.fetched)
    assert len(res.unfetched_image_paths) == 25 - len(res.fetched)      # dropped members stay listed
    assert res.n_images_listed == 25
    assert res.n_requests == 1 + len(res.fetched)                     # nothing requested past the budget


class _CorruptContainerZenodo(_FakeZenodo):
    """Container member bodies arrive whole but fail the CRC check (last byte
    flipped), so every member also costs a range read."""

    def get(self, url, headers=None, stream=False, timeout=None):
        r = super().get(url, headers=headers, stream=stream, timeout=timeout)
        if "/container/" in url and r.status_code == 200 and r.body:
            r.body = r.body[:-1] + bytes([r.body[-1] ^ 0xFF])
        return r


def test_record_budget_counts_bytes_of_failed_transfers(tmp_path):
    import random

    from envision_eye_actionable.survey import remote
    members = {f"s/{i:02d}.png": _png_bytes(size=(40, 30), color=(i * 9, 50, 50)) for i in range(10)}
    data = _zip_bytes(members)
    zips = [{"key": "x.zip", "size": len(data)}]
    biggest = max(len(v) for v in members.values())
    budget = int(biggest * 6)
    c = _client(tmp_path, _CorruptContainerZenodo("81", {"x.zip": data}))
    res = remote.sample_remote_zips(c, "81", zips, tmp_path / "r", 10, random.Random(1), workers=1,
                                    record_budget_bytes=budget)
    # every member's container body was received and thrown away: all of it counts
    assert res.bytes_received > res.bytes_fetched > 0
    assert res.bytes_received <= budget + biggest                   # at most one chunk past the budget
    assert 1 <= len(res.fetched) < 6                                # planned sizes alone would allow 6
    assert res.n_members_over_budget == 10 - len(res.fetched) and "budget" in res.budget_stop
    assert len(res.unfetched_image_paths) == 10 - len(res.fetched)  # stopped members stay listed
    assert not res.failures                                          # a budget stop is not a failure

    # without a budget the same fetches all succeed, and the count still covers the failed attempts
    c = _client(tmp_path, _CorruptContainerZenodo("81", {"x.zip": data}))
    res = remote.sample_remote_zips(c, "81", zips, tmp_path / "r2", 10, random.Random(1), workers=1)
    assert len(res.fetched) == 10 and res.bytes_received > 1.5 * res.bytes_fetched


def test_process_record_reports_remote_byte_columns(tmp_path):
    pytest.importorskip("jsonschema")
    rid = "720"
    photo = _png_bytes(size=(64, 48), color=(90, 90, 90))
    files = {"ir.zip": _zip_bytes({f"ir/OD_{i:02d}.png": photo for i in range(12)})}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=5,
                remote_record_budget_gb=len(photo) * 3.5 / 1e9, remote_max_member_mb=1.0)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "ok", row.get("error")
    assert row["n_members_skipped_oversize"] == 0
    assert row["n_remote_fetched"] == 3 and row["n_remote_over_budget"] == 2
    assert row["remote_bytes_fetched"] == 3 * len(photo) and row["remote_budget_stop"]
    assert row["n_image_files"] == 12
    assert row["remote_record_budget_gb"] == len(photo) * 3.5 / 1e9 and row["remote_max_member_mb"] == 1.0

    rid = "721"                                            # no remote zip: zero, not missing
    png = _png_bytes(size=(64, 48))
    files = {"a.png": png}
    (tmp_path / "b").mkdir()
    s = _survey(tmp_path / "b", _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["n_members_skipped_oversize"] == 0 and row["remote_bytes_fetched"] == 0


def test_partition_writes_id_lists_by_expected_mode(tmp_path):
    from envision_eye_actionable.survey.runner import expected_mode, partition_records, read_ids_file
    meta, cache, dl = tmp_path / "meta", tmp_path / "cache", tmp_path / "dl"
    for d in (meta, cache, dl):
        d.mkdir()
    recs = {
        "1": {"a.zip": b"zipbytes"},                       # on disk -> local
        "2": {"b.zip": 1000},                              # remote zip
        "3": {"c.rar": 1000, "r.pdf": 10},                 # download
        "4": {"r.pdf": 10},                                # nothing an image survey reads
        "5": {"d.zip": 1000, "e.rar": 1000},               # remote zip wins over download
        "6": {},                                           # no files
    }
    for rid, files in recs.items():
        (meta if rid != "5" else cache).joinpath(f"{rid}.json").write_text(json.dumps(_legacy(rid, files)))
    (dl / "1").mkdir()
    (dl / "1" / "a.zip").write_bytes(b"zipbytes")
    scrape = tmp_path / "s.json"
    scrape.write_text(json.dumps([{"source_id": r} for r in [*recs, "7"]]))   # 7: no metadata file
    out = tmp_path / "parts"
    summary = partition_records(scrape, dl, [meta, cache], out)
    ids = {m: read_ids_file(out / f"ids_{m}.txt") for m in ("local", "remote_zip", "download", "none", "unknown")}
    assert ids == {"local": ["1"], "remote_zip": ["2", "5"], "download": ["3"], "none": ["4", "6"],
                   "unknown": ["7"]}
    assert summary["counts"] == {"local": 1, "remote_zip": 2, "download": 1, "none": 2, "unknown": 1}
    assert json.loads((out / "partition_summary.json").read_text())["total"] == 7
    # --no-remote-zip: zips are downloads
    assert expected_mode([{"key": "b.zip", "size": 5}], dl / "2", remote_zip=False) == "download"


def test_ids_file_and_rpm_options(tmp_path, monkeypatch, capsys):
    from envision_eye_actionable.survey import cli
    from envision_eye_actionable.survey.runner import read_ids_file
    f = tmp_path / "ids.txt"
    f.write_text("# remote partition\n12 34\n56,12\n\n78  # trailing comment\n")
    assert read_ids_file(f) == ["12", "34", "56", "78"]
    model = tmp_path / "m.onnx"
    model.write_bytes(b"x")
    seen = {}
    monkeypatch.setattr("envision_eye_actionable.survey.runner.run_survey", lambda cfg: seen.update(cfg=cfg) or {})
    assert cli.main(["run", "--model", str(model), "--ids", "99", "12", "--ids-file", str(f), "--rpm", "90"]) == 0
    assert seen["cfg"].ids == ["99", "12", "34", "56", "78"] and seen["cfg"].zenodo_per_minute == 90
    assert seen["cfg"].min_eye_fraction == 0.05 and seen["cfg"].remote_record_budget_gb == 2.0
    assert seen["cfg"].remote_max_member_mb == 200.0
    assert cli.main(["run", "--model", str(model), "--zenodo-rpm", "20"]) == 0      # old spelling kept
    assert seen["cfg"].zenodo_per_minute == 20 and seen["cfg"].ids == []
    empty = tmp_path / "empty.txt"
    empty.write_text("# nothing\n")
    with pytest.raises(SystemExit):                 # an empty list must never mean "every record"
        cli.main(["run", "--model", str(model), "--ids-file", str(empty)])
    assert "no record ids" in capsys.readouterr().err


def test_one_process_per_out_and_scratch_dir(tmp_path):
    from envision_eye_actionable.survey import runner
    if not runner.LOCKING_AVAILABLE:
        pytest.skip("no file locking on this platform")
    held = runner.acquire_lock(tmp_path)
    with pytest.raises(ValueError, match="another survey process"):
        runner.acquire_lock(tmp_path)
    held.close()
    runner.acquire_lock(tmp_path).close()           # free again once released


def test_workbook_merges_several_results_dirs(tmp_path):
    pytest.importorskip("openpyxl")
    from envision_eye_actionable.survey import cli
    from openpyxl import load_workbook
    a, b = tmp_path / "survey_remote", tmp_path / "survey_download"
    for d, rows in ((a, [{"record_id": "1", "status": "error", "cmds_dir": "cmds/1"},
                         {"record_id": "2", "status": "ok", "cmds_dir": "cmds/2", "n_MASK": 3, "frac_MASK": 0.3,
                          "argmax_MASK": 3, "remote_bytes_fetched": 12345, "n_members_skipped_oversize": 2}]),
                    (b, [{"record_id": "1", "status": "ok", "cmds_dir": "cmds/1"},
                         {"record_id": "3", "status": "no_files", "cmds_dir": "cmds/3"}])):
        for r in rows:
            (d / "cmds" / r["record_id"]).mkdir(parents=True)
            (d / "cmds" / r["record_id"] / "dataset_description.json").write_text(
                json.dumps({"from": d.name, "id": r["record_id"]}), encoding="utf-8")
        (d / "survey_results.jsonl").write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
    out = tmp_path / "w.xlsx"
    assert cli.main(["excel", "--results-dir", str(a), "--results-dir", str(b), "--out", str(out)]) == 0
    wb = load_workbook(out)
    ws = wb["Records"]
    head = [c.value for c in ws[1]]
    rows = {r[0].value: {h: c.value for h, c in zip(head, r)} for r in ws.iter_rows(min_row=2)}
    assert set(rows) == {"1", "2", "3"}
    assert rows["1"]["Survey status"] == "ok"                                      # the last dir wins
    assert rows["1"]["Results dir (when several runs were merged)"] == "survey_download"
    assert rows["2"]["n MASK"] == 3 and rows["2"]["frac MASK"] == 0.3 and rows["2"]["argmax n MASK"] == 3
    assert rows["2"]["Remote bytes fetched (members + nested archives, failed transfers included)"] == 12345
    assert rows["2"]["Remote members never sampled (over --remote-max-member-mb)"] == 2
    cj = {r[0].value: json.loads(r[4].value) for r in wb["CMDS_JSON"].iter_rows(min_row=2)}
    assert cj["1"]["from"] == "survey_download" and cj["2"]["from"] == "survey_remote"   # each row's own dir
    readme = {r[0].value: r[1].value for r in wb["README"].iter_rows()}
    assert readme["Results merged from"] == "survey_remote; survey_download"
    assert "MASK" in readme["Classes"] and "min-eye-fraction" in readme["Eye images"]


def test_oversize_remote_members_stay_out_of_the_population(tmp_path):
    """Members over --remote-max-member-mb are outside the sampling frame, so
    the class mix of the small members is not projected onto them."""
    pytest.importorskip("jsonschema")
    rid = "730"
    small = {f"s/OD_{i:02d}.png": _png_bytes(size=(40, 30), color=(90, 90, 90)) for i in range(6)}
    rng = np.random.default_rng(0)
    big = {}
    for i in range(4):
        buf = io.BytesIO()
        Image.fromarray(rng.integers(0, 255, (120, 160, 3), dtype=np.uint8)).save(buf, "PNG")
        big[f"b/OD_{i}.png"] = buf.getvalue()
    cap = max(len(v) for v in small.values()) + 64     # compressed size (see above)
    assert min(len(v) for v in big.values()) > cap
    files = {"x.zip": _zip_bytes(small | big)}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=50, remote_max_member_mb=cap / 1e6)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["n_classified"] == 6, row.get("error")
    assert row["n_members_skipped_oversize"] == 4 and row["n_images_unsampleable"] == 4
    assert row["n_images_listed"] == 10          # every listed image is still seen
    assert row["n_image_files"] == 6             # but only the sampleable ones are the population
    assert "4 larger image member(s)" in row["fraction_scope_note"]
    assert row["n_members_zero_size"] == 0


def test_ir_trigger_and_mask_flag_use_non_mask_and_weighted_shares():
    from envision_eye_actionable.survey.dicom_map import aggregate, ir_device_fields, modality_review
    # 50% masks, IR 30% of the non-mask images (15% of all): substantial IR
    half_masks = {"n_classified": 100, "frac_IR": 0.15, "argmax_frac_IR": 0.15, "frac_IR_non_mask": 0.3,
                  "argmax_frac_IR_non_mask": 0.3, "dominant_class": "NEG", "argmax_dominant_class": "NEG",
                  "manufacturer_hint_text": "Topcon [topcon]", "manufacturer_hint_source": "text"}
    assert ir_device_fields(half_masks)["ir_non_spectralis_candidate"] is True
    low = {**half_masks, "frac_IR_non_mask": 0.1, "argmax_frac_IR_non_mask": 0.1, "frac_IR": 0.25}
    assert ir_device_fields(low)["ir_device_hint"] == ""       # the non-mask share decides when present
    # the review flag follows frac_MASK (weighted), mask_like_frac only for old rows
    base = {"present_eye_classes": "OCTA", "dominant_eye_class": "OCTA", "setfit_prob_eye_imaging": 0.9}
    assert "mask_like" in modality_review({**base, "frac_MASK": 0.5, "mask_like_frac": 0.1})["modality_review_flags"]
    assert modality_review({**base, "frac_MASK": 0.1, "mask_like_frac": 0.5})["modality_low_trust"] is False
    assert modality_review({**base, "mask_like_frac": 0.5})["modality_low_trust"] is True
    # mask_like_frac counts only images the pixel check ran on
    imgs = [_img("MASK", "m.png", mask_like=True), _img("OCTA", "o.png", mask_like=False),
            {"cls": "UNCERTAIN", "facts": {"format": "TIFF", "rows": 10, "cols": 10}, "path": ""}]
    assert aggregate(imgs, [], ["OCTA"], "OCTA", "", [])["mask_like_frac"] == 0.5


def test_classify_reports_ir_share_of_non_mask_images(tmp_path):
    rid = "731"
    files = {"a.png": _png_bytes(size=(64, 48))}
    s = _survey(tmp_path, _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert "frac_IR_non_mask" in row and "argmax_frac_IR_non_mask" in row


def test_zenodo_429_cooldown_is_shared_between_processes(tmp_path):
    from envision_eye_actionable.survey import zenodo
    shared = tmp_path / "state"
    a = _client(tmp_path / "a", None, shared_state_dir=shared)
    b = _client(tmp_path / "b", None, shared_state_dir=shared)
    solo = _client(tmp_path / "c", None)
    a.backoff(30)
    assert (shared / zenodo.SHARED_COOLDOWN_NAME).is_file()
    with b._lock:
        b._read_shared()
    assert 25 < b._not_before - zenodo.time.monotonic() <= 30.5
    a.backoff(5)                                   # a shorter cooldown never shortens the shared one
    assert float((shared / zenodo.SHARED_COOLDOWN_NAME).read_text()) - zenodo.time.time() > 25
    with solo._lock:
        solo._read_shared()
    assert solo._not_before == 0.0                 # no shared dir: not affected


def test_merge_keeps_finished_rows_and_names_same_leaf_dirs_apart(tmp_path):
    from envision_eye_actionable.survey.excel import merge_results, results_dir_labels
    a, b = tmp_path / "survey_smoke" / "results", tmp_path / "survey_smoke2" / "results"
    for d, rows in ((a, [{"record_id": "1", "status": "ok"}, {"record_id": "2", "status": "error"}]),
                    (b, [{"record_id": "1", "status": "started"}, {"record_id": "2", "status": "crashed"},
                         {"record_id": "3", "status": "started"}])):
        d.mkdir(parents=True)
        (d / "survey_results.jsonl").write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
    paths = [a / "survey_results.jsonl", b / "survey_results.jsonl"]
    labels = results_dir_labels(paths)
    assert labels == {str(paths[0]): "survey_smoke/results", str(paths[1]): "survey_smoke2/results"}
    assert all(not v.startswith("/") and ":" not in v for v in labels.values())
    m = merge_results(paths)
    assert m["1"]["status"] == "ok" and m["1"]["results_dir"] == "survey_smoke/results"   # started never wins
    assert m["2"]["status"] == "error"                                                    # nor crashed
    assert m["3"]["status"] == "started" and m["3"]["results_dir"] == "survey_smoke2/results"
    c = tmp_path / "other"
    assert results_dir_labels([a / "x.jsonl", c / "x.jsonl"]) == {str(a / "x.jsonl"): "results",
                                                                  str(c / "x.jsonl"): "other"}


# ---------------------------------------------------------------------------
# sixth round: combined request budget, disk reservations, locking fallback
# ---------------------------------------------------------------------------
def test_shared_request_budget_caps_processes_together(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import zenodo
    clock = [1_000_000.0]
    monkeypatch.setattr(zenodo.time, "monotonic", lambda: clock[0])
    monkeypatch.setattr(zenodo.time, "time", lambda: clock[0])
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: clock.__setitem__(0, clock[0] + s))
    shared = tmp_path / "state"
    # two processes, each at the default per-process limit, sharing a combined budget of 10
    a = zenodo.ZenodoClient(tmp_path / "a", None, min_interval=0, max_per_minute=110,
                            shared_state_dir=shared, shared_per_minute=10)
    b = zenodo.ZenodoClient(tmp_path / "b", None, min_interval=0, max_per_minute=110,
                            shared_state_dir=shared, shared_per_minute=10)
    stamps = []
    for i in range(30):
        (a if i % 2 else b).throttle()
        stamps.append(clock[0])
    for t in stamps:
        assert len([u for u in stamps if t <= u < t + 60.0]) <= 10
    assert stamps[-1] - stamps[0] >= 120.0              # 30 requests at 10/min: at least two full windows
    assert a.n_requests == b.n_requests == 15
    # without a shared dir each process only has its own limit
    solo = zenodo.ZenodoClient(tmp_path / "c", None, min_interval=0, max_per_minute=110)
    t0 = clock[0]
    for _ in range(30):
        solo.throttle()
    assert clock[0] == t0


def test_disk_reservations_of_other_processes_count_and_stale_ones_go(tmp_path):
    import os
    import shutil

    from envision_eye_actionable.survey import locks
    state, disk = tmp_path / "state", tmp_path / "scratch"
    disk.mkdir()
    mine = locks.DiskReservations(state, disk)
    assert mine.others() == 0
    # another live process: a reservation file whose lock is held
    other_pid = os.getpid() + 100000
    lock_fp = open(state / f"disk_reserve.{other_pid}.lock", "a+")
    if locks.LOCKING_AVAILABLE:
        assert locks.try_lock(lock_fp)
    (state / f"disk_reserve.{other_pid}.json").write_text(json.dumps({"pid": other_pid, "dev": mine.dev,
                                                                       "bytes": 5_000_000}))
    (state / "disk_reserve.9.json").write_text(json.dumps({"pid": 9, "dev": mine.dev + 1, "bytes": 7}))
    mine.reserve(123)                                    # its own reservation is not subtracted
    assert mine.others() == 5_000_000                    # the other disk's reservation does not count
    assert mine.free() <= shutil.disk_usage(disk).free - 5_000_000
    if locks.LOCKING_AVAILABLE:
        # the other process ends: its lock is free, its reservation is removed
        locks.unlock(lock_fp)
        lock_fp.close()
        assert mine.others() == 0
        assert not (state / f"disk_reserve.{other_pid}.json").exists()
    else:
        lock_fp.close()
    mine.close()
    assert not list(state.glob(f"disk_reserve.{os.getpid()}.*"))


def test_download_check_subtracts_other_reservations(tmp_path):
    pytest.importorskip("jsonschema")
    import shutil as _sh

    from envision_eye_actionable.survey import locks
    rid = "730"
    files = {"scan.tif": _png_bytes(size=(64, 48))}
    s = _survey(tmp_path, _legacy(rid, files), files, rid)
    s.cfg.shared_state_dir = tmp_path / "state"
    s.disk = locks.DiskReservations(s.cfg.shared_state_dir, s.cfg.scratch_dir)
    free = _sh.disk_usage(s.cfg.scratch_dir).free
    s.cfg.disk_floor_gb = max(0.0, (free - 10_000_000) / 1e9)    # 10 MB above the floor
    s.disk.others = lambda: 50_000_000                           # another process has 50 MB in flight
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "skipped_disk", row.get("error")
    s.disk.others = lambda: 0
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] != "skipped_disk", row.get("error")
    s.disk.close()


def test_stream_download_stops_at_the_disk_floor(tmp_path):
    from envision_eye_actionable.survey import zenodo

    class _Sess:
        def get(self, url, headers=None, stream=False, timeout=None):
            return _HttpResp(200, b"x" * (3 << 20))

    calls = []

    def room():
        calls.append(1)
        return len(calls) < 2                               # the disk fills up after the first check

    dest = tmp_path / "f.bin"
    ok, err = zenodo.stream_download(_Sess(), "http://x/f", dest, 3 << 20, room_check=room, room_every=1 << 20)
    assert not ok and err == "disk floor reached"
    assert not dest.exists() and not dest.with_name("f.bin.part").exists()
    ok, err = zenodo.stream_download(_Sess(), "http://x/f", dest, 3 << 20, room_check=lambda: True)
    assert ok and dest.stat().st_size == 3 << 20


def test_lock_fallback_warns_and_skips_the_scratch_sweep(tmp_path, monkeypatch, capsys):
    from envision_eye_actionable.survey import runner
    monkeypatch.setattr(runner, "LOCKING_AVAILABLE", False)
    assert runner.acquire_lock(tmp_path) is None
    assert "no file locking" in capsys.readouterr().err
    files = {"a.png": _png_bytes(size=(8, 8))}
    s = _survey(tmp_path, _legacy("1", files), files, "1")
    stale = s.records_scratch / "123"
    stale.mkdir(parents=True)
    s._sweep_scratch()
    assert stale.is_dir()                                   # another process may own it
    monkeypatch.setattr(runner, "LOCKING_AVAILABLE", True)
    s._sweep_scratch()
    assert not stale.exists()


# ---------------------------------------------------------------------------
# seventh round: dominance mask test, mask-dominated records, full CMDS
# folder in the workbook
# ---------------------------------------------------------------------------
def _antialiased_binary_mask() -> Image.Image:
    """A binary disk and vessel line drawn at 1024 and resized with
    interpolation: 0 and 255 plus a hundred-odd edge levels."""
    from PIL import ImageDraw
    big = Image.new("L", (1024, 1024), 0)
    d = ImageDraw.Draw(big)
    d.ellipse((200, 300, 700, 800), fill=255)
    d.line((0, 0, 1024, 1024), fill=255, width=30)
    return big.resize((256, 256), Image.BILINEAR)


def _antialiased_label_map() -> Image.Image:
    from PIL import ImageDraw
    lab = Image.new("RGB", (1024, 1024), (0, 0, 0))
    d = ImageDraw.Draw(lab)
    for k, c in enumerate([(255, 0, 0), (0, 255, 0), (0, 0, 255), (255, 255, 0), (0, 255, 255)]):
        d.ellipse((60 + k * 170, 100 + k * 120, 360 + k * 170, 400 + k * 120), fill=c)
    return lab.resize((256, 256), Image.BILINEAR)


def _posterized(noise_frac: float) -> Image.Image:
    """Four flat gray bands with ``noise_frac`` of the pixels random."""
    rng = np.random.default_rng(1)
    a = np.zeros((256, 256), np.uint8)
    a[:, :77], a[:, 77:128], a[:, 128:192], a[:, 192:] = 30, 90, 150, 210
    noisy = rng.random(a.shape) < noise_frac
    a[noisy] = rng.integers(0, 256, int(noisy.sum()))
    return Image.fromarray(a, "L")


def test_mask_dominance_catches_antialiased_masks_and_label_maps():
    from envision_eye_actionable.survey.images import MASK_SAMPLE, is_mask_like
    aa = _antialiased_binary_mask()
    small = np.asarray(aa.resize((MASK_SAMPLE, MASK_SAMPLE), Image.NEAREST))
    assert len(np.unique(small)) > 4                  # the exact-count rule alone misses it
    assert is_mask_like(aa)
    lab = _antialiased_label_map()
    small = np.asarray(lab.resize((MASK_SAMPLE, MASK_SAMPLE), Image.NEAREST)).reshape(-1, 3)
    assert len(np.unique(small, axis=0)) > 8
    assert is_mask_like(lab)
    # the 95% coverage bar: four bands with 2% noise are dominated, with 10% not
    assert is_mask_like(_posterized(0.02))
    assert not is_mask_like(_posterized(0.10))


def test_mask_dominance_ignores_a_small_object_on_black():
    """A small textured object on a black canvas: the background alone
    covers over 99% of the sample, but the object's pixels are spread over
    many values, so it is not a mask."""
    from envision_eye_actionable.survey.images import is_mask_like
    rng = np.random.default_rng(1)
    canvas = np.zeros((300, 300), np.uint8)
    canvas[140:160, 140:160] = rng.integers(50, 71, (20, 20))
    assert (canvas == 0).mean() > 0.99
    assert not is_mask_like(Image.fromarray(canvas, "L"))
    # the same canvas with the object replaced by a flat blob with a soft
    # (many-valued) edge is a mask: the edge is a small part of the object
    blob = np.zeros((300, 300), np.uint8)
    blob[130:170, 130:170] = 200
    blob[127:130, 130:170] = np.arange(40) * 4 + 20
    blob[170:173, 130:170] = np.arange(40) * 4 + 22
    small = np.asarray(Image.fromarray(blob, "L").resize((128, 128), Image.NEAREST))
    assert len(np.unique(small)) > 4
    assert is_mask_like(Image.fromarray(blob, "L"))


def _mask_record(tmp_path, rid: str, n_mask: int, n_photo: int):
    photo = _png_bytes(size=(64, 48), color=(90, 90, 90))
    members = ({f"img/{i:03d}.png": photo for i in range(n_photo)}
               | {f"seg/{i:03d}.png": _mask_png() for i in range(n_mask)})
    files = {"data.zip": _zip_bytes(members)}
    tmp_path.mkdir(exist_ok=True)
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=500)
    return s, s.process_record({"source_id": rid, "title": "Topcon infrared images"})   # _FakeClf: IR 0.9


def test_mask_dominated_record_is_not_ok_from_a_small_remainder(tmp_path):
    pytest.importorskip("jsonschema")
    s, row = _mask_record(tmp_path / "a", "740", n_mask=30, n_photo=8)
    assert row["n_MASK"] == 30 and row["n_classified_non_mask"] == 8, row.get("error")
    assert row["status"] == "mask_dominated"
    assert row["mask_dominated"] is True
    assert row["present_eye_classes"] == "" and row["review_eye_classes"] == "IR"
    assert row["dominant_eye_class"] is None and row["_class_detail"] == {}
    assert row["cmds_directories"] == "" and "n_weblinks" in row
    dsd = json.loads((s.cfg.out_dir / row["cmds_dir"] / "dataset_structure_description.json").read_text())
    assert dsd["directoryList"] == []
    assert row["dsd_valid"] and row["dd_valid"]


@pytest.mark.parametrize("n_mask,n_photo,status", [
    (60, 50, "ok"),     # masks over half, but 50 non-mask images: the normal rule
    (20, 30, "ok"),     # under 50 non-mask images, but masks under half
    (30, 9, "mask_dominated"),    # masks over half, 9 confident IR images: under the exemption
    (30, 30, "ok"),     # exactly half masks, 30 confident IR images: exempt (photos next to masks)
    (30, 10, "ok"),     # 10 confident IR images: exempt
])
def test_mask_dominated_needs_both_the_mask_share_and_a_small_remainder(tmp_path, n_mask, n_photo, status):
    pytest.importorskip("jsonschema")
    _, row = _mask_record(tmp_path, "741", n_mask=n_mask, n_photo=n_photo)
    assert row["n_MASK"] == n_mask and row["n_classified_non_mask"] == n_photo, row.get("error")
    assert row["status"] == status
    if status == "ok":
        assert row["mask_dominated"] is False and row["present_eye_classes"] == "IR"
        assert row["review_eye_classes"] == "" and "retinal_photography/ir" in row["cmds_directories"]


def test_workbook_shows_mask_dominated_records_and_full_cmds_folders(tmp_path):
    pytest.importorskip("jsonschema")
    pytest.importorskip("openpyxl")
    from envision_eye_actionable.survey.excel import build_workbook
    from openpyxl import load_workbook
    s, row = _mask_record(tmp_path / "a", "742", n_mask=30, n_photo=8)
    s2, row2 = _mask_record(tmp_path / "b", "743", n_mask=5, n_photo=10)
    assert row["status"] == "mask_dominated" and row2["status"] == "ok"
    for srv, r in ((s, row), (s2, row2)):
        srv.results_path.write_text(json.dumps(r, default=str) + "\n", encoding="utf-8")
    out = tmp_path / "w.xlsx"
    build_workbook([s.results_path, s2.results_path], out)
    wb = load_workbook(out)
    ws = wb["Records"]
    head = [c.value for c in ws[1]]
    rows = {r[0].value: {h: c.value for h, c in zip(head, r)} for r in ws.iter_rows(min_row=2)}
    review = "Eye modalities seen in a mask-dominated record (review only, not present)"
    assert rows["742"]["Survey status"] == "mask_dominated" and rows["742"][review] == "IR"
    assert rows["742"]["Eye modalities present"] in (None, "")
    assert rows["743"]["Eye modalities present"] == "IR" and rows["743"][review] in (None, "")
    rc = [r[0].value for r in wb["Record_Classes"].iter_rows(min_row=2)]
    assert rc == ["743"]                                  # no modality rows for the mask-dominated record
    readme = {r[0].value: r[1].value for r in wb["README"].iter_rows()}
    assert readme["Records mask-dominated (eye classes review only)"] == 1
    assert "mask_dominated: 1" in readme["Records by status"]
    assert "mask_dominated" in readme["Eye images"]
    # the CMDS folder in full: each results dir plus cmds/<id>, both sheets
    full = "CMDS JSON folder (results dir + cmds/<id>)"
    for rid, srv in (("742", s), ("743", s2)):
        want = (srv.cfg.out_dir.resolve() / "cmds" / rid).as_posix()
        assert rows[rid][full] == want
        assert rows[rid]["CMDS JSON folder (relative to the results dir)"] == f"cmds/{rid}"
    cj_head = [c.value for c in wb["CMDS_JSON"][1]]
    assert cj_head[1] == full
    cj = {r[0].value: r[1].value for r in wb["CMDS_JSON"].iter_rows(min_row=2)}
    assert cj["742"] == (s.cfg.out_dir.resolve() / "cmds" / "742").as_posix()
    assert cj["742"] != cj["743"].replace("743", "742")   # different results dirs stay distinct


# ---------------------------------------------------------------------------
# Archive probe (probe.py)
# ---------------------------------------------------------------------------
def _tar_bytes(members: dict, mode: str = "w") -> bytes:
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode=mode) as tf:
        for name, data in members.items():
            info = tarfile.TarInfo(name)
            info.size = len(data)
            tf.addfile(info, io.BytesIO(data))
    return buf.getvalue()


def _rar4_bytes(members: dict, main_flags: int = 0) -> bytes:
    """A minimal stored RAR 4 archive (header CRCs included)."""
    import struct
    import zlib

    def block(htype, flags, body, add=b""):
        hdr = struct.pack("<BHH", htype, flags, 7 + len(body)) + body
        return struct.pack("<H", zlib.crc32(hdr) & 0xFFFF) + hdr + add

    out = b"Rar!\x1a\x07\x00" + block(0x73, main_flags, b"\x00" * 6)
    for name, data in members.items():
        nb = name.encode()
        body = struct.pack("<IIBIIBBHI", len(data), len(data), 3, zlib.crc32(data), 0x00210000, 20, 0x30,
                           len(nb), 0x81A4 << 16) + nb
        out += block(0x74, 0x8000, body, data)
    return out + block(0x7B, 0x4000, b"")


def _probe(tmp_path, key, data, **kw):
    from envision_eye_actionable.survey import probe
    rid = "900"
    fake = _FakeZenodo(rid, {key: data})
    c = _client(tmp_path, fake)
    kw.setdefault("min_bytes", 0)
    return probe.probe_archive(c, rid, key, len(data), **kw), fake


def _noise(n: int, seed: int = 0) -> bytes:
    return np.random.default_rng(seed).integers(0, 256, n, dtype=np.uint8).tobytes()


def test_probe_member_signal_follows_the_walker():
    from envision_eye_actionable.survey.probe import member_signal
    dicm = b"\0" * 128 + b"DICM" + b"\0" * 10
    assert member_signal("a/b/scan.tif", 10, False) == "image"
    assert member_signal("vol.nii.gz", 10, False) == "image"
    assert member_signal("inner.zip", 10, False) == "nested"
    assert member_signal("x.tif.gz", 10, False) == "nested"
    assert member_signal("table.csv.gz", 10, False) == ""
    assert member_signal("__MACOSX/._a.png", 10, False) == ""
    assert member_signal("dir/._a.png", 10, True) == ""
    assert member_signal("code/main.py", 10, True) == ""
    # extensionless: only where the walker sniffs them (zip, tar), in its size range
    assert member_signal("IM0001", 5000, False) == ""
    assert member_signal("IM0001", 5000, True) == "noext"
    assert member_signal("IM0001", 7, True) == ""                 # under the walker's sniff size
    assert member_signal("IM0001", 80 << 20, True) == ""
    assert member_signal("pkg/README", 5000, True) == ""
    assert member_signal("IM0001", 5000, True, head=dicm) == "image"
    assert member_signal("IM0001", 5000, True, head=b"x" * 200) == ""


def test_probe_tar_hops_over_large_members(tmp_path):
    big = _noise(3 << 20)
    data = _tar_bytes({"raw/a.bin": big, "raw/b.bin": big, "docs/readme.txt": b"hi", "img/OD_1.png": _png_bytes()})
    res, fake = _probe(tmp_path, "d.tar", data)
    assert res.outcome == "images" and res.hit == "img/OD_1.png" and res.method == "tar_headers"
    assert res.n_members == 4 and res.n_images == 1
    assert res.bytes_used < len(data) // 4 and res.n_requests <= 5
    assert all(rng for _, rng in fake.calls)                     # range reads only

    neg = _tar_bytes({"raw/a.bin": big, "raw/b.bin": big, "docs/readme.txt": b"hi", "t.csv": b"1,2"})
    res, _ = _probe(tmp_path, "n.tar", neg)
    assert res.outcome == "no_images" and res.complete and res.n_members == 4
    assert res.kind_counts["other"] == 4 and "docs/readme.txt" in res.names
    assert res.bytes_used < len(neg) // 4
    # one part of a split set is never proven empty
    res, _ = _probe(tmp_path, "n.tar", neg, split=True)
    assert res.outcome == "unknown" and not res.complete


def test_probe_tar_limits_make_it_unknown(tmp_path):
    big = _noise(1 << 20)
    data = _tar_bytes({f"raw/{i}.bin": big for i in range(12)} | {"z.png": _png_bytes()})
    res, _ = _probe(tmp_path, "d.tar", data, max_requests=3)
    assert res.outcome == "unknown" and not res.complete and "request limit" in res.note
    res, _ = _probe(tmp_path, "d.tar", data)                      # within the default limits
    assert res.outcome == "images"


def test_probe_tar_sniffs_extensionless_members(tmp_path):
    dicm = b"\0" * 128 + b"DICM" + b"\0" * 400
    res, _ = _probe(tmp_path, "d.tar", _tar_bytes({"series/IM0001": dicm}))
    assert res.outcome == "images" and res.n_images == 1 and res.n_noext == 0
    res, _ = _probe(tmp_path, "d.tar", _tar_bytes({"bin/tool": b"\x7fELF" + b"\0" * 400, "README": b"x" * 300}))
    assert res.outcome == "no_images" and res.complete
    # an extensionless PNG is an image too (the walker sniffs PNG, JPEG, TIFF, BMP)
    res, _ = _probe(tmp_path, "p.tar", _tar_bytes({"scans/IMG0001": _png_bytes()}))
    assert res.outcome == "images" and res.n_images == 1


def test_probe_compressed_tar_streams_the_head(tmp_path):
    pos = _tar_bytes({"a.txt": b"x" * 1000, "b/OD.tif": b"II*\0" + b"\0" * 100}, mode="w:gz")
    res, fake = _probe(tmp_path, "d.tar.gz", pos)
    assert res.outcome == "images" and res.method == "tar_stream" and res.n_requests == 1
    neg = _tar_bytes({"a.txt": b"x" * 1000, "b.csv": b"1,2"}, mode="w:bz2")
    res, _ = _probe(tmp_path, "d.tar.bz2", neg)
    assert res.outcome == "no_images" and res.complete and res.bytes_used == len(neg)
    # past the byte limit without an image: unknown, never no_images
    big = _tar_bytes({f"n{i}.bin": _noise(400_000, i) for i in range(6)} | {"z.png": b"x"}, mode="w:gz")
    res, _ = _probe(tmp_path, "d.tgz", big, stream_bytes=1 << 20)
    assert res.outcome == "unknown" and not res.complete and "byte limit" in res.note
    assert res.bytes_used <= 1 << 20


def test_probe_7z_reads_the_end_header(tmp_path):
    py7zr = pytest.importorskip("py7zr")
    for name, members, want in (("p.7z", {"a/readme.txt": b"x" * 100, "a/img/OS.png": _png_bytes()}, "images"),
                                ("n.7z", {"a/readme.txt": b"x" * 100, "b/data.csv": b"1,2\n" * 50}, "no_images")):
        src = tmp_path / name
        with py7zr.SevenZipFile(src, "w") as z:
            for k, v in members.items():
                z.writestr(v, k)
            z.writestr(_noise(3 << 20), "a/zz_payload.bin")
        data = src.read_bytes()
        res, _ = _probe(tmp_path, name, data)
        assert res.outcome == want, res.note
        assert res.method == "7z_header" and res.complete
        assert res.n_members == 3
        assert res.bytes_used < len(data) // 2 and res.n_requests <= 4


def test_probe_rar_walks_block_headers(tmp_path):
    pytest.importorskip("rarfile")
    big = _noise(2 << 20)
    data = _rar4_bytes({"raw/a.bin": big, "raw/b.bin": big, "imgs/OD.jpg": b"\xff\xd8" + b"x" * 50})
    res, _ = _probe(tmp_path, "d.rar", data)
    assert res.outcome == "images" and res.method == "rar_headers" and res.hit == "imgs/OD.jpg"
    assert res.bytes_used < len(data) // 2
    neg = _rar4_bytes({"raw/a.bin": big, "raw/b.bin": big, "IM0001": b"\0" * 500})
    res, _ = _probe(tmp_path, "n.rar", neg)
    assert res.outcome == "no_images" and res.complete and res.n_members == 3   # rar: no noext sniff
    # one part of a split set never proves the set empty
    res, _ = _probe(tmp_path, "n.part1.rar", neg, split=True)
    assert res.outcome == "unknown" and not res.complete
    # the main header's volume flag says the same without the name
    vol = _rar4_bytes({"raw/a.bin": big, "IM0001": b"\0" * 500}, main_flags=0x0001)
    res, _ = _probe(tmp_path, "v.rar", vol)
    assert res.outcome == "unknown" and "multi-volume" in res.note


def test_probe_single_compressed_files(tmp_path):
    import gzip
    res, fake = _probe(tmp_path, "scan.tif.gz", gzip.compress(b"II*\0" + b"\0" * 200))
    assert res.outcome == "images" and res.n_requests == 0 and not fake.calls
    dicm = b"\0" * 128 + b"DICM" + _noise(5000)
    res, _ = _probe(tmp_path, "IM0001.gz", gzip.compress(dicm))
    assert res.outcome == "images" and res.method == "inner_sniff" and res.n_requests == 1
    res, _ = _probe(tmp_path, "blob.gz", gzip.compress(_noise(5000)))
    assert res.outcome == "no_images" and res.complete
    res, _ = _probe(tmp_path, "inner.zip.gz", gzip.compress(b"PK"))
    assert res.outcome == "unknown"


def test_probe_record_skips_small_and_split_and_honours_record_limits(tmp_path):
    from envision_eye_actionable.survey import probe
    rid = "901"
    neg = _tar_bytes({"a.txt": _noise(300_000)})
    files = {"a.tar": neg, "b.tar": neg, "c.tar": neg, "s.part1.rar": b"Rar!" * 400, "s.part2.rar": b"Rar!",
             "tiny.tar": _tar_bytes({"a.txt": b"x"})}
    c = _client(tmp_path, _FakeZenodo(rid, files))
    fl = [{"key": k, "size": len(v)} for k, v in files.items()]
    out = probe.probe_record_archives(c, rid, fl, list(files), min_bytes=20_000, record_requests=2)
    by = {r.key: r for r in out.results}
    assert by["a.tar"].outcome == "no_images"
    assert by["tiny.tar"].outcome == "not_probed" and "small" in by["tiny.tar"].note
    assert out.stop == "record probe limit reached"
    assert any(r.note == "record probe limit reached" for r in out.results)
    assert "s.part2.rar" not in by                               # fragments follow their first part
    assert by["s.part1.rar"].outcome in ("unknown", "not_probed")


def test_process_record_skips_archives_probed_negative(tmp_path, monkeypatch):
    pytest.importorskip("jsonschema")
    from envision_eye_actionable.survey import probe
    monkeypatch.setattr(probe, "MIN_PROBE_BYTES", 0)
    rid = "902"
    neg = _tar_bytes({"code/run.py": b"print(1)", "tables/t.csv": b"1,2", "Topcon_export.txt": b"x"},
                     mode="w:gz")
    files = {"code.tar.gz": neg, "readme.pdf": b"%PDF"}
    s = _survey(tmp_path, _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "no_images_in_archives" and row["sampling_mode"] == "none"
    assert row["n_archives_probed"] == 1 and row["n_archive_probe_no_images"] == 1
    assert row["archive_probe_bytes_avoided"] == len(neg) and row["n_files_to_download"] == 0
    assert row["archive_probes"][0]["outcome"] == "no_images"
    assert row["file_kind_counts"]["other"] == 3 and row["n_archives_listed"] == 1
    assert "Topcon" in row["manufacturer_hint_text"]              # member names are catalogued
    assert all(rng for url, rng in s.zenodo.session.calls if url.endswith("/content"))   # no whole download
    assert "n_weblinks" in row and row["dd_valid"]
    ev = [json.loads(line) for line in s.events_path.read_text(encoding="utf-8").splitlines()]
    assert any(e["event"] == "archive_probe" and e["archives"][0]["outcome"] == "no_images" for e in ev)

    # a positive archive next to a negative one: only the positive one is downloaded
    rid = "903"
    pos = _tar_bytes({"ir/OD_1.png": _png_bytes(size=(64, 48)), "ir/OD_2.png": _png_bytes(size=(64, 48))})
    files = {"code.tar.gz": neg, "ir.tar": pos}
    (tmp_path / "b").mkdir()
    s = _survey(tmp_path / "b", _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "ok" and row["n_classified"] == 2 and row["sampling_mode"] == "download"
    assert row["n_archive_probe_images"] == 1 and row["n_archive_probe_no_images"] == 1
    whole = [url for url, rng in s.zenodo.session.calls if not rng and url.endswith("/content")]
    assert len(whole) == 1 and "ir.tar" in whole[0]
    assert row["n_archives_listed"] == 2

    # --no-archive-probe: downloaded whole as before, no_image_files
    rid = "904"
    files = {"code.tar.gz": neg}
    (tmp_path / "c").mkdir()
    s = _survey(tmp_path / "c", _legacy(rid, files), files, rid, archive_probe=False)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "no_image_files" and row["sampling_mode"] == "download"
    assert row["archive_probe_enabled"] is False and "archive_probes" not in row


def test_workbook_lists_archive_probes(tmp_path, monkeypatch):
    pytest.importorskip("jsonschema")
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    from envision_eye_actionable.survey import probe
    from envision_eye_actionable.survey.excel import build_workbook
    monkeypatch.setattr(probe, "MIN_PROBE_BYTES", 0)
    rid = "905"
    files = {"code.tar.gz": _tar_bytes({"a.csv": b"1"}, mode="w:gz")}
    s = _survey(tmp_path, _legacy(rid, files), files, rid)
    s._append_row(s.process_record({"source_id": rid, "title": "T"}))
    build_workbook(s.results_path, tmp_path / "w.xlsx")
    wb = load_workbook(tmp_path / "w.xlsx")
    head = [c.value for c in wb["Archive_Probes"][1]]
    got = dict(zip(head, [c.value for c in wb["Archive_Probes"][2]]))
    assert got["Archive"] == "code.tar.gz" and got["Record status"] == "no_images_in_archives"
    assert got["Outcome (images: downloaded; no_images: not downloaded; unknown: downloaded as before)"] == "no_images"
    rec_head = [c.value for c in wb["Records"][1]]
    assert "Download bytes avoided by the archive probe" in rec_head


# ---------------------------------------------------------------------------
# Format conversion (images.py), audit fixes
# ---------------------------------------------------------------------------
def test_tiff_over_budget_is_read_strided_not_refused(tmp_path):
    """The 15116835 bug: a TIFF over the pixel cap was refused (too_large)
    instead of being read at a reduced resolution."""
    tifffile = pytest.importorskip("tifffile")
    from envision_eye_actionable.survey.images import _strided_page
    yy, xx = np.mgrid[0:200, 0:300]
    arr = (yy * 300 + xx).astype(np.uint16)
    for name, kw in (("strips.tif", {"rowsperstrip": 7}), ("tiles.tif", {"tile": (64, 64)})):
        p = tmp_path / name
        tifffile.imwrite(p, arr, **kw)
        with tifffile.TiffFile(p) as tf:
            got = _strided_page(tf.pages[0], 3)
        assert np.array_equal(got, arr[::3, ::3]), name
        ld = load_frame("raster", ".tif", None, p, max_pixels=10_000)
        assert ld.image is not None and ld.error is None
        assert "strided 1/3" in ld.facts["conversion"] and ld.image.size == (100, 67)
        assert (ld.facts["rows"], ld.facts["cols"]) == (200, 300)
    # float32 GeoTIFF, planar separate, with a -3.4e38 nodata border
    band = np.linspace(0, 1, 200 * 300, dtype=np.float32).reshape(200, 300)
    band2 = band * 2
    band[:10] = band2[:10] = -3.4e38
    p = tmp_path / "geo.tif"
    tifffile.imwrite(p, np.stack([band, band2]), planarconfig="separate", photometric="minisblack")
    with tifffile.TiffFile(p) as tf:
        got = _strided_page(tf.pages[0], 4)
    assert got.shape == (50, 75, 2) and np.array_equal(got[..., 1], band2[::4, ::4])
    ld = load_frame("raster", ".tif", None, p, max_pixels=5_000)
    assert ld.image is not None and "strided 1/4" in ld.facts["conversion"] and ld.facts.get("float")


def test_single_strip_tiff_over_the_budget_is_viewed_or_refused_never_decoded_whole(tmp_path):
    """ImageJ-style TIFFs hold the whole plane in one strip: decoding that
    strip is decoding the plane. Uncompressed: strided through a view of
    the file (or of the bytes in memory); compressed: too_large."""
    tifffile = pytest.importorskip("tifffile")
    from envision_eye_actionable.survey import images
    yy, xx = np.mgrid[0:200, 0:300]
    arr = (yy * 300 + xx).astype(np.uint16)
    rgb = np.stack([arr, arr // 2, arr // 3], axis=-1).astype(np.uint16)
    decoded = []
    orig = images._strided_page
    images._strided_page = lambda page, step: decoded.append(step) or orig(page, step)
    try:
        for name, data, kw in (("one.tif", arr, {}), ("big_endian.tif", arr, {"byteorder": ">"}),
                               ("rgb.tif", rgb, {"photometric": "rgb"}),
                               ("planar.tif", np.moveaxis(rgb, -1, 0), {"photometric": "rgb",
                                                                        "planarconfig": "separate"})):
            p = tmp_path / name
            tifffile.imwrite(p, data, rowsperstrip=200, **kw)
            with tifffile.TiffFile(p) as tf:
                assert images._tiff_segment_bytes(tf.pages[0]) >= 200 * 300 * 2
            want = arr[::3, ::3] if data is arr else rgb[::3, ::3]
            for src in (None, p.read_bytes()):
                ld = images._load_tiff(src, p, ".tif", max_pixels=10_000, max_bytes=50_000)
                assert ld.image is not None and ld.error is None, (name, ld.error)
                assert "uncompressed view" in ld.facts["conversion"] and "strided 1/3" in ld.facts["conversion"]
                assert ld.image.size == (100, 67)
            with tifffile.TiffFile(p) as tf:
                got = images._contiguous_strided(tf, tf.pages[0], 3, None, p)
            assert got.dtype.isnative and np.array_equal(got, want), name
        assert decoded == []                      # no strip was decoded
        # the same plane compressed in one strip: refused, not decoded
        p = tmp_path / "zlib.tif"
        tifffile.imwrite(p, arr, rowsperstrip=200, compression="zlib")
        ld = images._load_tiff(None, p, ".tif", max_pixels=10_000, max_bytes=50_000)
        assert ld.image is None and ld.error.startswith("too_large") and decoded == []
        # many small strips under the budget: decoded strip by strip as before
        p = tmp_path / "strips.tif"
        tifffile.imwrite(p, arr, rowsperstrip=7, compression="zlib")
        ld = images._load_tiff(None, p, ".tif", max_pixels=10_000, max_bytes=50_000)
        assert ld.image is not None and decoded == [3]
    finally:
        images._strided_page = orig


def test_array_to_rgb8_windows_in_place_within_a_memory_bound():
    """A plane at the decode budget must not take 4x its size again in the
    consumer: one float32 working copy plus bool masks."""
    import tracemalloc
    from envision_eye_actionable.survey.images import array_to_rgb8
    rng = np.random.default_rng(1)
    a = rng.normal(100, 20, (1500, 2000)).astype(np.float32)
    a[:5] = -3.4e38                                  # nodata border
    a[5, :10] = np.nan
    ref = a.astype(np.float64)
    ok = np.isfinite(ref) & (np.abs(ref) < 1e30)
    lo, hi = np.percentile(ref[ok][:: max(1, ok.sum() // 1_000_000)], [0.5, 99.5])
    want = np.clip((np.where(ok, ref, lo) - lo) / (hi - lo) * 255, 0, 255).astype(np.uint8)
    tracemalloc.start()
    try:
        img = array_to_rgb8(a)
        _cur, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    got = np.asarray(img)[..., 0]
    assert np.abs(got.astype(int) - want.astype(int)).max() <= 2
    assert (got[:5] == 0).all()                      # nodata goes to the low end
    assert peak < 2.6 * a.nbytes, peak / a.nbytes
    u16 = rng.integers(0, 4096, (1500, 2000)).astype(np.uint16)
    tracemalloc.start()
    try:
        array_to_rgb8(u16)
        _cur, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert peak < 5.0 * u16.nbytes, peak / u16.nbytes    # float32 copy (2x) + masks + sample + uint8 out


def test_tiff_pyramid_level_and_middle_plane(tmp_path):
    tifffile = pytest.importorskip("tifffile")
    a0 = np.random.default_rng(0).integers(0, 255, (1500, 2400), dtype=np.uint8)
    p = tmp_path / "pyr.tif"
    with tifffile.TiffWriter(p) as tw:
        tw.write(a0, subifds=2, tile=(256, 256))
        for k in (2, 4):
            tw.write(a0[::k, ::k], subfiletype=1, tile=(256, 256))
    ld = load_frame("raster", ".tif", None, p)
    # smallest level whose long side is at least 1024: 2400 -> 1200 (level 1)
    assert ld.image.size == (1200, 750) and "pyramid level 1/3" in ld.facts["conversion"]
    assert (ld.facts["full_rows"], ld.facts["full_cols"]) == (1500, 2400)
    # a Z stack of same-shape pages gives its middle plane, not page 0
    rng = np.random.default_rng(1)
    stack = np.stack([np.clip(i * 50 + rng.integers(-5, 6, (40, 50)), 0, 255) for i in range(5)]).astype(np.uint8)
    p = tmp_path / "stack.tif"
    tifffile.imwrite(p, stack, imagej=True, metadata={"axes": "ZYX"})
    ld = load_frame("raster", ".tif", None, p)
    assert "Z=2/5" in ld.facts["conversion"] and ld.facts["frame_index"] == 2
    assert abs(np.asarray(ld.image)[..., 0].astype(float).mean() - np.asarray(stack[2], float).mean()) < 30
    assert ld.facts["frames"] == 5


def test_native_mask_test_on_deep_integer_images():
    from envision_eye_actionable.survey.images import array_to_rgb8, is_mask_like, native_mask_like
    rng = np.random.default_rng(0)
    real = (1000 + rng.integers(-30, 31, (120, 160))).astype(np.uint16)
    real.ravel()[rng.choice(real.size, real.size // 50, replace=False)] = 65535    # 2% saturated
    # the 8-bit test sees a few gray levels; the native values do not
    assert is_mask_like(array_to_rgb8(real))
    assert native_mask_like(real) is False
    labels = (np.arange(120 * 160).reshape(120, 160) // 4000 % 4 * 1000).astype(np.uint16)
    assert native_mask_like(labels) is True
    assert native_mask_like(np.zeros((10, 10), np.uint8)) is None          # 8-bit: the 8-bit test applies
    assert native_mask_like(np.zeros((10, 10, 3), np.uint16)) is None


def test_deep_tiff_facts_use_the_native_mask_test(tmp_path):
    tifffile = pytest.importorskip("tifffile")
    rng = np.random.default_rng(0)
    real = (1000 + rng.integers(-30, 31, (120, 160))).astype(np.uint16)
    real.ravel()[rng.choice(real.size, real.size // 50, replace=False)] = 65535
    p = tmp_path / "deep.tif"
    tifffile.imwrite(p, real)
    ld = load_frame("raster", ".tif", None, p)
    assert ld.facts["mask_like"] is False and ld.facts["mask_test"] == "native"
    buf = io.BytesIO()
    Image.fromarray(real).save(buf, "PNG")                      # 16-bit PNG through PIL
    ld = load_frame("raster", ".png", buf.getvalue(), None)
    assert ld.facts["mask_like"] is False and ld.facts["mask_test"] == "native"


def test_transparency_is_composited_over_white(tmp_path):
    from envision_eye_actionable.survey.images import array_to_rgb8
    im = Image.new("RGBA", (60, 40), (0, 0, 0, 0))
    im.paste((200, 30, 30, 255), (20, 10, 40, 30))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    ld = load_frame("raster", ".png", buf.getvalue(), None)
    assert ld.image.getpixel((0, 0)) == (255, 255, 255) and ld.image.getpixel((30, 20)) == (200, 30, 30)
    arr = np.zeros((40, 60, 4), np.uint8)
    arr[10:30, 20:40] = (200, 30, 30, 255)
    out = array_to_rgb8(arr)
    assert out.getpixel((0, 0)) == (255, 255, 255) and out.getpixel((30, 20)) == (200, 30, 30)


def test_blank_frames_are_flagged():
    ld = load_frame("raster", ".png", _png_bytes(size=(64, 48), plain=True), None)
    assert ld.facts["blank"] is True
    ld = load_frame("raster", ".png", _png_bytes(size=(64, 48)), None)
    assert ld.facts["blank"] is False


def test_svg_embedded_raster_and_sandboxed_render(tmp_path):
    pytest.importorskip("resvg_py")
    import base64
    photo = _png_bytes(size=(80, 60))
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="200" '
           f'height="100"><image width="80" height="60" xlink:href="data:image/png;base64,'
           f'{base64.b64encode(photo).decode()}"/></svg>').encode()
    ld = load_frame("vector", ".svg", svg, None)
    assert ld.image.size == (80, 60) and ld.facts["conversion"] == "SVG embedded PNG image"
    # a vector figure is rendered; a reference to a local file is removed
    red = tmp_path / "red.png"
    Image.new("RGB", (50, 50), (255, 0, 0)).save(red)
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="100" '
           f'height="50"><rect width="100" height="50" fill="#0000ff"/><image width="100" height="50" '
           f'preserveAspectRatio="none" xlink:href="{red.as_posix()}"/><image width="100" height="50" '
           f'href="https://example.org/x.png"/></svg>').encode()
    ld = load_frame("vector", ".svg", svg, None)
    assert ld.image.size == (512, 256) and "rendered" in ld.facts["conversion"]
    for xy in ((20, 20), (256, 128), (490, 240)):
        assert ld.image.getpixel(xy) == (0, 0, 255)                # the rect, not the local red file


def test_compressed_dicom_transfer_syntaxes_decode():
    pydicom = pytest.importorskip("pydicom")
    from pydicom.uid import JPEG2000Lossless, RLELossless
    for uid in (RLELossless, JPEG2000Lossless):
        ds = pydicom.dcmread(io.BytesIO(_dicom_bytes(frames=3, rows=64, cols=48)))
        try:
            ds.compress(uid)
        except Exception as e:  # noqa: BLE001 - encoder plugin missing
            pytest.skip(f"no encoder for {uid}: {e}")
        buf = io.BytesIO()
        ds.save_as(buf, enforce_file_format=True)
        ld = load_frame("dicom", ".dcm", buf.getvalue(), None)
        assert ld.image is not None, (uid, ld.error)
        assert "frame 1/3" in ld.facts["conversion"] and str(uid.name) in ld.facts["source_format"]


def test_arrays_pick_an_image_shaped_dataset(tmp_path):
    from envision_eye_actionable.survey.images import plane_plan
    vol = np.stack([np.full((128, 96), i * 1000, np.uint16) + np.arange(96, dtype=np.uint16) for i in range(10)])
    np.save(tmp_path / "v.npy", vol)
    ld = load_frame("array", ".npy", None, tmp_path / "v.npy")
    assert ld.image is not None and ld.image.size == (96, 128) and "npy (10, 128, 96)" in ld.facts["conversion"]
    np.savez(tmp_path / "z.npz", labels=np.arange(50), images=np.random.default_rng(0).random((4, 64, 80)))
    ld = load_frame("array", ".npz", None, tmp_path / "z.npz")
    assert ld.image is not None and "npz images" in ld.facts["conversion"]
    # a name hint beats a larger plane: oct/bscan over raw/table
    h5py = pytest.importorskip("h5py")
    with h5py.File(tmp_path / "d.h5", "w") as f:
        f["raw/table"] = np.random.default_rng(1).random((300, 400))
        f["oct/bscan"] = np.random.default_rng(2).random((5, 128, 100)).astype(np.float32)
    ld = load_frame("array", ".h5", None, tmp_path / "d.h5")
    assert "oct/bscan" in ld.facts["conversion"] and ld.image.size == (100, 128)
    # MAT v7.3 (HDF5 after a 512-byte header) and MAT v5
    with h5py.File(tmp_path / "m73.mat", "w", userblock_size=512) as f:
        f["img"] = np.random.default_rng(3).random((70, 90))
    ld = load_frame("array", ".mat", None, tmp_path / "m73.mat")
    assert ld.image is not None and ld.facts["source_format"] == "MAT v7.3 (HDF5)"
    scipy_io = pytest.importorskip("scipy.io")
    scipy_io.savemat(tmp_path / "m5.mat", {"img": np.random.default_rng(4).integers(0, 255, (100, 80)).astype(np.uint8),
                                           "vec": np.arange(500)})
    ld = load_frame("array", ".mat", None, tmp_path / "m5.mat")
    assert ld.image is not None and "mat img" in ld.facts["conversion"] and ld.facts["source_format"] == "MAT v5"
    # signals, tables and feature matrices are not images
    np.save(tmp_path / "sig.npy", np.random.default_rng(5).random((4096, 64)))
    ld = load_frame("array", ".npy", None, tmp_path / "sig.npy")
    assert ld.image is None and "not image shaped" in ld.error
    assert plane_plan((4096, 64)) is None and plane_plan((40, 40)) is None and plane_plan((64, 64)) is not None
    assert plane_plan((3, 200, 300))[1] == (200, 300)


def test_network_weights_are_not_images(tmp_path):
    """A Keras model .h5 (seen live: a record's model file classified as
    OCTA) and layer parameter tensors elsewhere are never rendered."""
    from envision_eye_actionable.survey.runner import unread_reason
    h5py = pytest.importorskip("h5py")
    rng = np.random.default_rng(0)
    with h5py.File(tmp_path / "keras.h5", "w") as f:
        f.attrs["keras_version"] = "2.4.0"
        f["layers/dense/vars/0"] = rng.random((512, 256)).astype(np.float32)   # no parameter-like name
    ld = load_frame("array", ".h5", None, tmp_path / "keras.h5")
    assert ld.image is None and "model weights" in ld.error
    assert unread_reason(ld.error) == "not_image_shaped"
    # a group model_weights alone marks the file too
    with h5py.File(tmp_path / "weights_only.h5", "w") as f:
        f["model_weights/conv/conv/kernel:0"] = rng.random((3, 3, 64, 128)).astype(np.float32)
        f["model_weights/dense/w"] = rng.random((300, 200)).astype(np.float32)
    ld = load_frame("array", ".h5", None, tmp_path / "weights_only.h5")
    assert ld.image is None and "model weights" in ld.error
    # parameter tensors in a plain file are skipped; an image next to them is still read
    with h5py.File(tmp_path / "mixed.h5", "w") as f:
        f["encoder/conv1/weight"] = rng.random((256, 256)).astype(np.float32)
        f["encoder/bn1/running_mean"] = rng.random((128, 128)).astype(np.float32)
        f["sample"] = rng.random((70, 90)).astype(np.float32)
    ld = load_frame("array", ".h5", None, tmp_path / "mixed.h5")
    assert ld.image is not None and "hdf5 sample" in ld.facts["conversion"], ld.facts
    with h5py.File(tmp_path / "params.h5", "w") as f:
        f["layer1/weight"] = rng.random((256, 256)).astype(np.float32)
        f["layer1/bias"] = rng.random((128, 128)).astype(np.float32)
    ld = load_frame("array", ".h5", None, tmp_path / "params.h5")
    assert ld.image is None and "no image-shaped dataset among 2" in ld.error
    np.savez(tmp_path / "w.npz", fc_weight=rng.random((200, 200)), fc_bias=rng.random((100, 100)))
    ld = load_frame("array", ".npz", None, tmp_path / "w.npz")
    assert ld.image is None, ld.facts
    # image names that merely contain a parameter word are still images
    with h5py.File(tmp_path / "imgs.h5", "w") as f:
        f["weighted_images"] = rng.random((80, 80)).astype(np.float32)
    ld = load_frame("array", ".h5", None, tmp_path / "imgs.h5")
    assert ld.image is not None


def test_arrays_over_the_budget_are_strided(tmp_path):
    arr = np.random.default_rng(0).random((600, 800)).astype(np.float32)
    np.save(tmp_path / "big.npy", arr)
    ld = load_frame("array", ".npy", None, tmp_path / "big.npy", max_pixels=100_000)
    assert ld.image is not None and "strided 1/3" in ld.facts["conversion"] and ld.image.size == (267, 200)


def test_video_frames_are_averaged(tmp_path):
    imageio_ffmpeg = pytest.importorskip("imageio_ffmpeg")
    vp = tmp_path / "clip.mp4"
    w = imageio_ffmpeg.write_frames(str(vp), (64, 48), fps=10, macro_block_size=1)
    w.send(None)
    rng = np.random.default_rng(0)
    for i in range(30):
        w.send(np.clip(i * 8 + rng.integers(-5, 6, (48, 64, 3)), 0, 255).astype(np.uint8).tobytes())
    w.close()
    ld = load_frame("video", ".mp4", None, vp)
    assert ld.image is not None and len(ld.frames) == 2, ld.error
    assert ld.facts["duration_s"] > 2 and "25%" in ld.facts["conversion"]
    means = [np.asarray(im, float).mean() for im in (ld.frames[0], ld.image, ld.frames[1])]
    assert means[0] < means[1] < means[2]                          # 25%, 50%, 75% of the clip


def _mhd_pair(tmp_path, name="vol"):
    sitk = pytest.importorskip("SimpleITK")
    arr = np.stack([np.full((64, 80), i * 20, np.uint8) + np.arange(80, dtype=np.uint8) for i in range(5)])
    img = sitk.GetImageFromArray(arr)
    p = tmp_path / f"{name}.mhd"
    sitk.WriteImage(img, str(p), useCompression=False)
    return p, tmp_path / f"{name}.raw"


def test_volume_header_pairs_are_read_together(tmp_path):
    hdr, raw = _mhd_pair(tmp_path)
    assert raw.exists()
    ld = load_frame("volume_pair", ".mhd", None, hdr)
    assert ld.image is not None and ld.image.size == (80, 64) and ld.facts["source_format"] == "MHD"
    # inside a zip and a tar: header and data member are extracted together
    zp = tmp_path / "pairs.zip"
    with zipfile.ZipFile(zp, "w") as zf:
        zf.write(hdr, "scans/vol.mhd")
        zf.write(raw, "scans/vol.raw")
        zf.writestr("scans/lone.mhd", hdr.read_bytes())
    tp = tmp_path / "pairs.tar"
    with tarfile.open(tp, "w") as tf:
        tf.add(raw, "scans/vol.raw")                     # data member first
        tf.add(hdr, "scans/vol.mhd")
    for arch in (zp, tp):
        w = RecordWalker(scratch=tmp_path / f"s_{arch.suffix[1:]}")
        w.add_file(arch, arch.name)
        pair = [e for e in w.entries if e.member.endswith("vol.mhd")][0]
        assert pair.kind == "volume_pair" and pair.companion == "scans/vol.raw"
        assert w.kind_counts["volume_data"] == 1
        got = {}
        for e, data, path, err in w.read_entries(w.entries):
            got[e.member] = load_frame(e.kind, e.ext, data, path) if not err else err
        assert got["scans/vol.mhd"].image is not None, (arch, got["scans/vol.mhd"])
        if "scans/lone.mhd" in got:
            lone = got["scans/lone.mhd"]
            assert isinstance(lone, str) or lone.image is None


def test_analyze_pair_and_mrc(tmp_path):
    nib = pytest.importorskip("nibabel")
    arr = np.random.default_rng(0).integers(0, 1000, (70, 60, 4)).astype(np.int16)
    nib.save(nib.AnalyzeImage(arr, np.eye(4)), str(tmp_path / "a.hdr"))
    assert (tmp_path / "a.img").exists()
    ld = load_frame("volume_pair", ".hdr", None, tmp_path / "a.hdr")
    assert ld.image is not None and ld.facts["source_format"] == "ANALYZE"
    (tmp_path / "b.hdr").write_bytes((tmp_path / "a.hdr").read_bytes())      # header without data
    ld = load_frame("volume_pair", ".hdr", None, tmp_path / "b.hdr")
    assert ld.image is None and "companion" in ld.error
    mrcfile = pytest.importorskip("mrcfile")
    with mrcfile.new(str(tmp_path / "m.mrc")) as m:
        m.set_data(np.random.default_rng(1).random((6, 70, 90)).astype(np.float32))
    ld = load_frame("microscopy", ".mrc", None, tmp_path / "m.mrc")
    assert ld.image is not None and ld.image.size == (90, 70)


def _heidelberg_vol(path, size_x=64, n=3, size_z=48, slo=32, bhdr=456):
    import struct
    hdr = bytearray(2048)
    hdr[:11] = b"HSF-OCT-103"
    struct.pack_into("<3i", hdr, 12, size_x, n, size_z)
    struct.pack_into("<2i", hdr, 48, slo, slo)
    struct.pack_into("<i", hdr, 100, bhdr)
    rng = np.random.default_rng(0)
    body = rng.integers(0, 255, slo * slo, dtype=np.uint8).tobytes()
    for i in range(n):
        b = (rng.random((size_z, size_x)) * (i + 1) * 0.1).astype("<f4")
        b[0, 0] = 3.4e38                                    # invalid pixel marker
        body += b"\0" * bhdr + b.tobytes()
    path.write_bytes(bytes(hdr) + body)


def test_vendor_vol_gives_a_bscan_and_a_fundus_view(tmp_path):
    p = tmp_path / "scan.vol"
    _heidelberg_vol(p)
    ld = load_frame("vendor_oct", ".vol", None, p)
    assert ld.image is not None and ld.image.size == (64, 48), ld.error
    assert ld.facts["make"] == "Heidelberg Engineering" and ld.facts["frames"] == 3
    assert (ld.facts["rows"], ld.facts["cols"]) == (48, 64)
    assert "middle B-scan 1/3" in ld.facts["conversion"]
    assert [v[0] for v in ld.views] == ["fundus"] and ld.views[0][1].size == (32, 32)
    (tmp_path / "x.sdb").write_bytes(b"\0" * 100)
    ld = load_frame("vendor_oct", ".sdb", None, tmp_path / "x.sdb")
    assert ld.image is None and "catalogued" in ld.error


def test_noext_members_are_sniffed_for_images_and_bare_dicom(tmp_path):
    pydicom = pytest.importorskip("pydicom")
    from envision_eye_actionable.survey.constants import sniff_kind
    # a DICOM data set without the 128-byte preamble (ACR-NEMA style)
    ds = pydicom.dcmread(io.BytesIO(_dicom_bytes(frames=1)))
    buf = io.BytesIO()
    ds.preamble = None
    ds.file_meta.TransferSyntaxUID = pydicom.uid.ImplicitVRLittleEndian
    pydicom.dcmwrite(buf, ds, enforce_file_format=False, implicit_vr=True, little_endian=True)
    bare = buf.getvalue()
    assert bare[:2] in (b"\x08\x00", b"\x02\x00") and sniff_kind(bare[:132]) == ("dicom", ".dcm")
    assert sniff_kind(_png_bytes()[:132]) == ("raster", ".png")
    assert sniff_kind(b"hello world, this is a readme" * 5) is None
    zp = tmp_path / "n.zip"
    with zipfile.ZipFile(zp, "w") as zf:
        zf.writestr("s/IMG0001", _png_bytes())
        zf.writestr("s/IM0002", bare)
        zf.writestr("s/README", b"text " * 50)
    w = RecordWalker(scratch=tmp_path / "s")
    w.add_file(zp, "n.zip")
    kinds = {e.member: (e.kind, e.ext) for e in w.entries}
    assert kinds == {"s/IMG0001": ("raster", ".png"), "s/IM0002": ("dicom", ".dcm")}
    assert w.member_ext_counts["(none)"] == 3 and sum(w.member_ext_counts.values()) == 3


def test_zip_members_at_wrapped_offsets_are_found(tmp_path):
    """4005629: a zip over 4 GiB without ZIP64 records; zipfile shifts every
    member offset by 4 GiB and members in the first 4 GiB fail."""
    from envision_eye_actionable.survey.archives import open_zip_member, wrapped_offsets
    assert wrapped_offsets(5 << 32, 7 << 32) == [4 << 32, 3 << 32, 2 << 32, 1 << 32, 0, 6 << 32]
    data = _zip_bytes({"a/OD.png": _png_bytes(), "b/OS.png": _png_bytes(color=(10, 200, 10))})
    zf = zipfile.ZipFile(io.BytesIO(data))
    info = zf.getinfo("b/OS.png")
    info.header_offset += 1 << 32                            # what the wrapped central directory gives
    with pytest.raises(Exception):
        zf.open(info).read()
    with open_zip_member(zf, info) as fp:
        assert fp.read() == _png_bytes(color=(10, 200, 10))


def test_remote_member_at_wrapped_offset_and_deflate64(tmp_path, monkeypatch):
    import random

    from envision_eye_actionable.survey import remote, zenodo
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: None)       # the 500 retries back off
    data = _zip_bytes({"a/OD.png": _png_bytes(), "b/OS.png": _png_bytes(color=(10, 200, 10))})
    fake = _FakeZenodo("81", {"x.zip": data}, listing_status=500)      # container API fails: range path
    c = _client(tmp_path, fake)
    lst = remote.list_zip(c, "81", "x.zip", len(data))
    assert lst.container_failed and lst.source == "range"
    rz = remote.RangeZips(c, "81")
    _, zf = rz.get("x.zip", len(data))
    for i in zf.infolist():
        i.header_offset += 1 << 32
    m = next(x for x in lst.members if x.name == "b/OS.png")
    ok, err = remote.fetch_member(c, "81", lst, m, tmp_path / "o.png", len(data), rz)
    assert ok, err
    assert (tmp_path / "o.png").read_bytes() == _png_bytes(color=(10, 200, 10))
    # Deflate64 members (4969418), streamed through inflate64
    from envision_eye_actionable.survey.archives import SEVEN_ZIP
    if not SEVEN_ZIP:
        pytest.skip("7z binary needed to write a Deflate64 zip")
    import subprocess
    src = tmp_path / "d64"
    src.mkdir()
    payload = (_png_bytes(size=(200, 150)) * 20)
    (src / "OD.png").write_bytes(payload)
    zp = tmp_path / "d64.zip"
    subprocess.run([SEVEN_ZIP, "a", "-tzip", "-mm=Deflate64", str(zp), str(src / "OD.png")], check=True,
                   capture_output=True)
    d64 = zp.read_bytes()
    assert zipfile.ZipFile(zp).infolist()[0].compress_type == 9
    fake = _FakeZenodo("82", {"d.zip": d64}, listing_status=500)
    c = _client(tmp_path / "c2", fake)
    res = remote.sample_remote_zips(c, "82", [{"key": "d.zip", "size": len(d64)}], tmp_path / "r", 5,
                                    random.Random(0))
    assert len(res.fetched) == 1 and not res.failures, res.failures
    assert res.fetched[0][0].read_bytes() == payload


def test_remote_member_cap_is_on_the_transfer_and_compressible_members_use_range(tmp_path, monkeypatch):
    """4943680: 794 MB TIFF stacks in a 188 MB zip were never sampled (the
    cap was on the uncompressed size)."""
    import random

    from envision_eye_actionable.survey import remote
    big_flat = b"II*\x00" + bytes(3 << 20)                  # 3 MB, compresses to a few KB
    data = _zip_bytes({"stack/LegStack.tif": big_flat})
    fake = _FakeZenodo("83", {"x.zip": data})
    c = _client(tmp_path, fake)
    monkeypatch.setattr(remote, "COMPRESSED_PREFER_BYTES", 1 << 20)
    res = remote.sample_remote_zips(c, "83", [{"key": "x.zip", "size": len(data)}], tmp_path / "r", 5,
                                    random.Random(0), max_member_bytes=1 << 20)
    assert res.n_members_skipped_oversize == 0 and len(res.fetched) == 1, res.failures
    assert res.fetched[0][0].stat().st_size == len(big_flat)
    assert not any("/container/" in url for url, _ in fake.calls)       # moved compressed, by range
    assert res.bytes_received < 1 << 20


def test_zenodo_token_goes_to_zenodo_only_and_is_never_shown(tmp_path, monkeypatch):
    import requests

    from envision_eye_actionable.survey.zenodo import ZenodoClient, load_token
    secret = "tok-" + "x" * 20
    tf = tmp_path / "token"
    tf.write_text(secret + "\n", encoding="utf-8")
    monkeypatch.delenv("ZENODO_TOKEN", raising=False)
    assert load_token(None) is None and load_token(tf) == secret
    monkeypatch.setenv("ZENODO_TOKEN", "from-env")
    assert load_token(tf) == "from-env"
    c = ZenodoClient(tmp_path / "cache", None, token=secret)
    zen = c.session.prepare_request(requests.Request("GET", "https://zenodo.org/api/records/1"))
    other = c.session.prepare_request(requests.Request("GET", "https://example.org/data.zip"))
    assert zen.headers["Authorization"] == f"Bearer {secret}"
    assert "Authorization" not in other.headers
    assert secret not in repr(c.session.auth) and secret not in str(vars(c).get("session").auth)
    assert ZenodoClient(tmp_path / "cache2", None).session.auth is None


def test_rate_limit_remaining_header_pauses_requests(tmp_path):
    import time as _time
    c = _client(tmp_path, _FakeZenodo("1", {}))
    c.observe(_HttpResp(200, headers={"X-RateLimit-Remaining": "50", "X-RateLimit-Reset": str(_time.time() + 30)}))
    assert c._not_before <= _time.monotonic()
    c.observe(_HttpResp(200, headers={"X-RateLimit-Remaining": "2", "X-RateLimit-Reset": str(_time.time() + 30)}))
    assert 20 < c._not_before - _time.monotonic() <= 31


class _NegClf:
    """Every image NEG with p=0.95 (a record of graphics)."""

    def predict(self, batch):
        p = np.full((len(batch), 7), 0.05 / 6, np.float32)
        p[:, 6] = 0.95
        return p


def test_mask_like_graphics_with_argmax_neg_keep_the_model_label(tmp_path):
    """38 of the 65 mask_dominated records were plots, word clouds and logos:
    flagged by the pixel check, called NEG by the model."""
    pytest.importorskip("jsonschema")
    rid = "750"
    files = {"figs.zip": _zip_bytes({f"fig/{i}.png": _mask_png(levels=(0, 90, 200)) for i in range(6)})}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=50)
    s.clf = _NegClf()
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "no_eye_images", row.get("error")
    assert row["n_MASK"] == 0 and row["n_NEG"] == 6 and row["n_mask_like_veto"] == 6
    assert row["mask_dominated"] is False
    import gzip
    preds = [json.loads(x) for x in gzip.open(s.cfg.out_dir / "predictions" / f"{rid}.jsonl.gz", "rt")]
    assert all(p["mask_like_veto"] and p["label"] == "NEG" for p in preds)


def test_mask_exemption_needs_a_confident_non_octa_class():
    from envision_eye_actionable.survey.runner import mask_exempt_class
    cfp = [{"cls": "CFP", "conf": 0.82}] * 10
    assert mask_exempt_class(cfp) == "CFP"
    assert mask_exempt_class(cfp[:9]) is None                                      # too few
    assert mask_exempt_class([{"cls": "CFP", "conf": 0.7}] * 20) is None           # not confident
    assert mask_exempt_class([{"cls": "OCTA", "conf": 0.95}] * 50) is None         # missed masks look OCTA
    assert mask_exempt_class(cfp + [{"cls": "IR", "conf": 0.9}] * 12) == "IR"      # the larger class


def test_statuses_split_no_image_files_and_images_unreadable(tmp_path):
    pytest.importorskip("jsonschema")
    rid = "760"
    files = {"broken.png": b"\x89PNG\r\n\x1a\n" + b"\0" * 200, "notes.csv": b"a,b"}
    s = _survey(tmp_path, _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "images_unreadable", row.get("error")
    assert row["n_pixel_files_unread_by_reason"] == {"decode_error": 1}
    assert row["member_ext_counts"] == {}                      # no archive members
    import gzip
    pred = json.loads(gzip.open(s.cfg.out_dir / "predictions" / f"{rid}.jsonl.gz", "rt").readline())
    assert pred["reason"] == "decode_error" and pred["path"] == "broken.png"
    # blank frames only: images_unreadable, reason blank
    rid = "761"
    files = {"b.zip": _zip_bytes({f"x/{i}.png": _png_bytes(size=(64, 48), plain=True) for i in range(3)})}
    (tmp_path / "b").mkdir()
    s = _survey(tmp_path / "b", _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "images_unreadable" and row["n_blank"] == 3 and row["n_classified"] == 0
    assert row["n_pixel_files_unread_by_reason"] == {"blank": 3}
    # nothing pixel-bearing listed
    rid = "762"
    files = {"notes.csv": b"a,b", "paper.pdf": b"%PDF"}
    (tmp_path / "c").mkdir()
    s = _survey(tmp_path / "c", _legacy(rid, files), files, rid)
    assert s.process_record({"source_id": rid, "title": "T"})["status"] == "no_image_files"


def test_records_classify_converted_formats(tmp_path):
    """A record mixing a VOL file (B-scan plus fundus view), a volume pair and
    an array: every one converted, summarised per record."""
    pytest.importorskip("jsonschema")
    rid = "770"
    _heidelberg_vol(tmp_path / "scan.vol")
    (tmp_path / "v").mkdir()
    hdr, raw = _mhd_pair(tmp_path / "v")
    buf = io.BytesIO()
    np.save(buf, np.random.default_rng(0).random((3, 90, 70)))
    files = {"scan.vol": (tmp_path / "scan.vol").read_bytes(), "vol.mhd": hdr.read_bytes(),
             "vol.raw": raw.read_bytes(), "stack.npy": buf.getvalue()}
    (tmp_path / "run").mkdir()
    s = _survey(tmp_path / "run", _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "ok", (row.get("error"), row.get("load_error_top"))
    assert row["n_classified"] == 4                           # vol, vol#fundus, mhd pair, npy
    assert set(row["formats_classified"]) == {"Heidelberg VOL", "MHD", "NPY"}
    import gzip
    preds = [json.loads(x) for x in gzip.open(s.cfg.out_dir / "predictions" / f"{rid}.jsonl.gz", "rt")]
    assert {p["path"] for p in preds} == {"scan.vol", "scan.vol#fundus", "vol.mhd", "stack.npy"}
    assert all(p.get("conversion") for p in preds)


def test_series_of_large_top_level_images_is_sampled():
    import random

    from envision_eye_actionable.survey.runner import sample_series
    big = [{"key": f"landcover_{i:03d}.tif", "size": 900 << 20} for i in range(8)]
    small = [{"key": f"thumb_{i:03d}.tif", "size": 1 << 20} for i in range(8)]
    other = [{"key": "data.tar", "size": 900 << 20}]
    keep, rest = sample_series(big + small + other, 3, 64 << 20, random.Random(0))
    assert len(rest) == 5 and all(f["key"].startswith("landcover_") for f in rest)
    assert sum(1 for f in keep if f["key"].startswith("landcover_")) == 3
    assert all(f in keep for f in small + other)
    assert sample_series(big, 0, 64 << 20, random.Random(0)) == (big, [])


def test_split_sets_are_joined_and_zstd_tars_read(tmp_path):
    pytest.importorskip("jsonschema")
    zstandard = pytest.importorskip("zstandard")
    rid = "780"
    tar = _tar_bytes({"img/OD.png": _png_bytes(), "img/OS.png": _png_bytes(color=(20, 90, 200))})
    half = len(tar) // 2
    files = {"scans.tar.partaa": tar[:half], "scans.tar.partab": tar[half:]}
    s = _survey(tmp_path, _legacy(rid, files), files, rid)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["n_classified"] == 2 and row["status"] == "ok", (row.get("error"), row.get("walk_error_detail"))
    assert "2 parts joined" in row["split_join_detail"]
    # .tar.zst: walked and probed
    zst = zstandard.ZstdCompressor().compress(tar)
    w = RecordWalker(scratch=tmp_path / "zs")
    (tmp_path / "x.tar.zst").write_bytes(zst)
    w.add_file(tmp_path / "x.tar.zst", "x.tar.zst")
    assert len(w.entries) == 2
    got = [load_frame(e.kind, e.ext, d, p) for e, d, p, err in w.read_entries(w.entries)]
    assert all(g.image is not None for g in got)
    res, _ = _probe(tmp_path, "x.tar.zst", zst)
    assert res.outcome == "images" and res.method == "tar_stream"
    from envision_eye_actionable.survey.constants import file_kind, is_fragment, split_first, split_stem
    assert file_kind("x.tar.001") == "archive" and is_fragment("x.tar.002") and split_first("x.zip.partaa")
    assert is_fragment("x.zip.partab") and split_stem("x.zip.partab") == "x.zip" and file_kind("x.part01") == "archive"


def test_remote_volume_pairs_fetch_the_data_member(tmp_path):
    import random

    from envision_eye_actionable.survey import remote
    (tmp_path / "v").mkdir()
    hdr, raw = _mhd_pair(tmp_path / "v")
    data = _zip_bytes({"ct/vol.mhd": hdr.read_bytes(), "ct/vol.raw": raw.read_bytes(), "ct/notes.txt": b"x"})
    c = _client(tmp_path, _FakeZenodo("84", {"x.zip": data}))
    res = remote.sample_remote_zips(c, "84", [{"key": "x.zip", "size": len(data)}], tmp_path / "r", 5,
                                    random.Random(0))
    assert len(res.fetched) == 1 and not res.failures, res.failures
    path, disp = res.fetched[0]
    assert disp == "x.zip!/ct/vol.mhd" and (path.parent / "vol.raw").exists()
    ld = load_frame("volume_pair", ".mhd", None, path)
    assert ld.image is not None, ld.error
    assert res.member_ext_counts == {".mhd": 1, ".raw": 1, ".txt": 1}


def _thorlabs_oct(path, intensity=True):
    """A Thorlabs .oct as the instrument writes it: a zip with Header.xml and
    data\*.data members named with backslashes."""
    rng = np.random.default_rng(0)
    video = rng.integers(0, 255, (48, 64, 4), dtype=np.uint8)
    if not intensity:
        video[..., 3] = 255                                      # interleaved BGRA (else: planes)
    files = {"data\\VideoImage.data": video.tobytes()}
    # SizeZ is the row length and SizeX the row count (as in a real file)
    df = ['<DataFile Type="Colored" SizeZ="64" SizeX="48" BytesPerPixel="4">data\\VideoImage.data</DataFile>']
    if intensity:
        vol = rng.random((5, 80, 60)).astype("<f4")             # SizeY x SizeX x SizeZ, z fastest
        vol[2] += 3.0                                            # the middle B-scan is brighter
        files["data\\Intensity.data"] = vol.tobytes()
        df.append('<DataFile Type="Real" SizeZ="60" SizeX="80" SizeY="5" BytesPerPixel="4">'
                  'data\\Intensity.data</DataFile>')
    else:
        z = np.arange(256)
        phase = rng.uniform(0, 2 * np.pi, (300, 1))             # a moving sample: the mean spectrum is flat
        spec = 1000 + 200 * np.cos(2 * np.pi * z[None, :] * 20 / 256 + phase) + rng.integers(0, 5, (300, 256))
        for i in range(3):
            files[f"data\\Spectral{i}.data"] = spec.astype("<u2").tobytes()
            df.append(f'<DataFile Type="Raw" SizeZ="256" SizeX="300" BytesPerPixel="2">data\\Spectral{i}.data'
                      f'</DataFile>')
    files["Header.xml"] = ('<?xml version="1.0"?><Ocity><DataFiles>' + "".join(df)
                           + "</DataFiles></Ocity>").encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as zf:
        for k, v in files.items():
            zf.writestr(k, v)


def test_thorlabs_oct_intensity_and_raw_spectra(tmp_path):
    p = tmp_path / "a.oct"
    _thorlabs_oct(p)
    ld = load_frame("vendor_oct", ".oct", None, p)
    assert ld.image is not None and ld.image.size == (80, 60), ld.error
    assert ld.facts["make"] == "Thorlabs" and "intensity B-scan 2/5" in ld.facts["conversion"]
    assert ld.views and ld.views[0][1].size == (64, 48)
    p = tmp_path / "b.oct"
    _thorlabs_oct(p, intensity=False)
    ld = load_frame("vendor_oct", ".oct", None, p)
    assert ld.image is not None and "raw spectra FFT (300 A-scans of spectral file 1/3)" in ld.facts["conversion"]
    assert ld.views[0][1].size == (64, 48) and not ld.views[0][2]["rgb_is_gray"]    # BGRA read as color
    col = np.asarray(ld.image)[..., 0].astype(float).mean(axis=1)
    assert abs(int(col.argmax()) - 19) <= 1                      # the reflector at depth bin 20
    (tmp_path / "fire.fds").write_text("&HEAD CHID='x' /\n" * 20)
    ld = load_frame("vendor_oct", ".fds", None, tmp_path / "fire.fds")
    assert ld.image is None and "wrong_format" in ld.error


def _rar5_bytes(members: dict, archive_flags: int = 0, encrypted: bool = False) -> bytes:
    """A minimal stored RAR 5 archive (header CRCs included)."""
    import zlib

    def vint(n):
        out = b""
        while True:
            b = n & 0x7F
            n >>= 7
            out += bytes([b | (0x80 if n else 0)])
            if not n:
                return out

    def block(htype, hflags, body, data=b""):
        head = vint(htype) + vint(hflags) + (vint(len(data)) if hflags & 0x2 else b"") + body
        return (zlib.crc32(vint(len(head)) + head) & 0xFFFFFFFF).to_bytes(4, "little") + vint(len(head)) + head + data

    out = b"Rar!\x1a\x07\x01\x00"
    if encrypted:
        out += block(4, 0, vint(0) + vint(0) + b"\0" * 16)
    out += block(1, 0, vint(archive_flags))
    for name, data in members.items():
        nb = name.encode()
        out += block(2, 0x2, vint(0) + vint(len(data)) + vint(0x20) + vint(0) + vint(1) + vint(len(nb)) + nb, data)
    return out + block(5, 0, vint(0))


def test_probe_rar_walker_stops_early_and_reads_rar5(tmp_path):
    big = _noise(2 << 20)
    # an image first: the walk stops there, whatever follows
    data = _rar4_bytes({"imgs/OD.jpg": b"\xff\xd8" + b"x" * 50} | {f"raw/{i}.bin": big for i in range(8)})
    res, _ = _probe(tmp_path, "d.rar", data)
    assert res.outcome == "images" and res.n_members == 1 and res.n_requests <= 2 and res.method == "rar_headers"
    # many small members, no image: the read block keeps growing over short hops
    # (hops of 200 KB: past the 64 KB base block, so without the growth every
    # header would cost a request of its own and the 50-request limit would stop the walk)
    small = {f"tab/{i:04d}.csv": _noise(200_000, i) for i in range(100)}
    data = _rar4_bytes(small)
    res, _ = _probe(tmp_path, "t.rar", data)
    assert res.outcome == "no_images" and res.n_members == 100 and res.n_requests <= 20, (res.n_requests, res.note)
    # RAR 5
    res, _ = _probe(tmp_path, "p5.rar", _rar5_bytes({"a/readme.txt": b"x" * 100, "a/scan.png": _png_bytes()}))
    assert res.outcome == "images" and res.hit == "a/scan.png"
    res, _ = _probe(tmp_path, "n5.rar", _rar5_bytes({"a/readme.txt": b"x" * 100, "b/t.csv": big}))
    assert res.outcome == "no_images" and res.complete and res.n_members == 2 and res.bytes_used < len(big) // 2
    res, _ = _probe(tmp_path, "v5.rar", _rar5_bytes({"b/t.csv": big}, archive_flags=0x1))
    assert res.outcome == "unknown" and "multi-volume" in res.note
    res, _ = _probe(tmp_path, "e5.rar", _rar5_bytes({"b/t.csv": big}, encrypted=True))
    assert res.outcome == "unknown" and "encrypted" in res.note


def test_unsampled_files_follow_a_sample_that_was_all_non_pixel(tmp_path):
    """5084941: 2047 .npy spike-count arrays; the 60 sampled are signals, not
    images, so the record has no image files rather than unreadable ones."""
    pytest.importorskip("jsonschema")
    rid = "790"
    members = {}
    for i in range(12):
        buf = io.BytesIO()
        np.save(buf, np.arange(3599, dtype=float))
        members[f"spikes/cell_{i}.npy"] = buf.getvalue()
    files = {"d.zip": _zip_bytes(members)}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=5)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "no_image_files", row.get("n_pixel_files_unread_by_reason")
    assert row["n_pixel_files_unread_by_reason"] == {"not_sampled_like_sample": 7, "not_image_shaped": 5}
    # one sampled file that is pixel-bearing but unreadable keeps the rest pixel-bearing
    rid = "791"
    members["spikes/cell_0.npy"] = bytes([0x93]) + b"NUMPY broken"
    files = {"d.zip": _zip_bytes(members)}
    (tmp_path / "b").mkdir()
    s = _survey(tmp_path / "b", _legacy(rid, files), files, rid, remote_cap=12)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "images_unreadable" and row["n_pixel_files_unread_by_reason"]["decode_error"] == 1

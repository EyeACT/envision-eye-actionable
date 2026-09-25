"""Offline tests of the full-pull architecture: token handling, two-pass
sampling, the spool, the fetch / process / monitor roles and the pipeline
supervisor, and the streamed workbook (no network, no ONNX model)."""

from __future__ import annotations

import json
import logging
import random
import threading
import time
from pathlib import Path

import numpy as np
import pytest
import requests

from test_survey import (_BrightDarkClf, _BrightNegDarkOctClf, _FakeClf, _FakeZenodo, _HttpResp, _client,
                         _legacy, _png_bytes, _survey, _zip_bytes)

SECRET = "tokSECRET" + "q" * 30


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
class _Multi:
    """Several _FakeZenodo records behind one session; ``leak`` makes the
    member fetches of that record fail with the token in the error text."""

    def __init__(self, fakes: dict, leak: str | None = None):
        self.fakes, self.leak, self.calls = fakes, leak, []

    def get(self, url, headers=None, stream=False, timeout=None):
        self.calls.append(url)
        for rid, fake in self.fakes.items():
            if f"/records/{rid}/" in url:
                if self.leak == rid and "/container/" in url:
                    raise requests.ConnectionError(f"proxy said: Authorization: Bearer {SECRET}")
                return fake.get(url, headers=headers, stream=stream, timeout=timeout)
        return _HttpResp(404)


class _OnnxStub:
    """Stands in for runner.OnnxClassifier (bright CFP, dark IR)."""

    meta = {"arch": "stub"}

    def __init__(self, *a, **k):
        self._clf = _BrightDarkClf()

    def predict(self, batch):
        return self._clf.predict(batch)


class _NegOnnxStub(_OnnxStub):
    def __init__(self, *a, **k):
        self._clf = _BrightNegDarkOctClf()


def _cfg(tmp_path, meta: Path, **kw):
    from envision_eye_actionable.survey.runner import SurveyConfig
    (tmp_path / "dl").mkdir(exist_ok=True)
    base = dict(scrape_path=tmp_path / "scrape.json", model_path=tmp_path / "m.onnx", out_dir=tmp_path / "out",
                downloads_dir=tmp_path / "dl", metadata_dir=meta, scratch_dir=tmp_path / "scratch",
                metadata_cache_dir=tmp_path / "mcache", spool_dir=tmp_path / "spool", state_dir=tmp_path / "state",
                check_links=False, disk_floor_gb=0, poll_s=0.05, zenodo_interval=0, zenodo_per_minute=None,
                zenodo_shared_per_minute=None, fetch_workers=2, download_workers=2)
    base.update(kw)
    return SurveyConfig(**base)


def _write_records(tmp_path, records: dict) -> Path:
    """records: {rid: {filename: bytes}} -> metadata dir + scrape file."""
    meta = tmp_path / "meta"
    meta.mkdir(exist_ok=True)
    scrape = []
    for rid, files in records.items():
        (meta / f"{rid}.json").write_text(json.dumps(_legacy(rid, files)), encoding="utf-8")
        scrape.append({"source_id": rid, "title": f"record {rid}", "label": "NEGATIVE", "prob_eye_imaging": 0.1,
                       "external_links": ["https://example.org/data"]})
    (tmp_path / "scrape.json").write_text(json.dumps(scrape), encoding="utf-8")
    return meta


def _all_text(*dirs) -> str:
    import gzip
    out = []
    for d in dirs:
        for p in ([Path(d)] if Path(d).is_file() else Path(d).rglob("*")):
            if p.is_file():
                data = p.read_bytes()
                if p.suffix == ".gz":
                    try:
                        data = gzip.decompress(data)
                    except OSError:
                        pass
                out.append(data.decode("utf-8", "replace"))
    return "\n".join(out)


def _run_roles(cfg, session, monkeypatch, onnx=_OnnxStub, timeout=60):
    """Producer and Processor of the pipeline in two threads of this process
    (role locks are per open file, so each sees the other alive)."""
    from envision_eye_actionable.survey import pipeline, runner
    monkeypatch.setattr(runner, "OnnxClassifier", onnx)
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = session
    proc = pipeline.Processor(cfg)
    out = {}

    def go(name, obj):
        try:
            out[name] = obj.run()
        except Exception as e:  # noqa: BLE001
            out[name] = e

    threads = [threading.Thread(target=go, args=("fetch", prod)), threading.Thread(target=go, args=("process", proc))]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout)
    assert not any(t.is_alive() for t in threads), "roles did not finish"
    return out, prod, proc


# ---------------------------------------------------------------------------
# A. token
# ---------------------------------------------------------------------------
def test_token_default_file_env_and_explicit(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import zenodo
    assert zenodo.load_token(None) is None and zenodo.token_source(None) == "none"
    default = tmp_path / "cfg" / "zenodo_token"
    default.parent.mkdir()
    default.write_text(SECRET + "\n", encoding="utf-8")
    monkeypatch.setattr(zenodo, "DEFAULT_TOKEN_FILE", default)
    assert zenodo.load_token(None) == SECRET and zenodo.token_source(None) == "default_file"
    assert zenodo.load_token(None, use_default=False) is None
    other = tmp_path / "t2"
    other.write_text("other-token-value\n", encoding="utf-8")
    assert zenodo.load_token(other) == "other-token-value"
    with pytest.raises(ValueError):
        zenodo.load_token(tmp_path / "missing")
    monkeypatch.setenv("ZENODO_TOKEN", "env-token-value")
    assert zenodo.load_token(other) == "env-token-value" and zenodo.token_source(other) == "env"


def test_redaction_of_text_and_log_records(tmp_path, caplog):
    from envision_eye_actionable.survey import zenodo
    zenodo.register_secret(SECRET)
    assert zenodo.redact(f"x {SECRET} y") == f"x {zenodo.REDACTED} y"
    assert zenodo.redact(None) is None
    lg = logging.getLogger("envision_eye_actionable.survey.test_redact")
    handler = logging.StreamHandler(stream=(buf := __import__("io").StringIO()))
    lg.addHandler(handler)
    lg.propagate = False
    zenodo.install_log_redaction(lg)
    try:
        lg.warning("GET failed: %s", f"Bearer {SECRET}")
        try:
            raise RuntimeError(f"boom {SECRET}")
        except RuntimeError:
            lg.exception("with traceback")
    finally:
        lg.removeHandler(handler)
    text = buf.getvalue()
    assert SECRET not in text and zenodo.REDACTED in text and "boom" in text


def test_rows_events_and_predictions_never_hold_the_token(tmp_path, monkeypatch):
    """A fetch error whose text carries the token (a proxy echoing the
    Authorization header) reaches rows, events and predictions redacted."""
    pytest.importorskip("jsonschema")
    from envision_eye_actionable.survey import zenodo
    monkeypatch.setattr(zenodo.time, "sleep", lambda s: None)
    tok = tmp_path / "token"
    tok.write_text(SECRET + "\n", encoding="utf-8")
    rid = "901"
    png = _png_bytes(size=(64, 48), color=(40, 40, 40))
    # the token as a file name: a vehicle that carries it into paths, errors,
    # predictions and rows
    files = {"d.zip": _zip_bytes({f"ir/{SECRET}_{i}.png": png for i in range(4)}), f"t_{SECRET}.png": png,
             "bad.zip": b"not a zip"}
    served = dict(files)
    files[f"gone_{SECRET}.png"] = png                     # listed, not served: download_failed event
    meta = _write_records(tmp_path, {rid: files})
    scrape = json.loads((tmp_path / "scrape.json").read_text())
    scrape[0]["title"] = f"title {SECRET}"                  # into the result row
    (tmp_path / "scrape.json").write_text(json.dumps(scrape), encoding="utf-8")
    cfg = _cfg(tmp_path, meta, token_file=tok, triage_cap=0)
    session = _Multi({rid: _FakeZenodo(rid, served)}, leak=rid)
    out, prod, proc = _run_roles(cfg, session, monkeypatch)
    assert out == {"fetch": 0, "process": 0}
    assert prod.zenodo.authenticated and isinstance(prod.zenodo.session, _Multi)
    # a spool record whose file went missing: spool_invalid event with the path
    sp = proc.spool
    sp.record_dir("902").mkdir(parents=True)
    sp.write_manifest("902", {"items": [{"path": f"dl/{SECRET}.png", "size": 1, "role": "download", "key": "k"}],
                              "row": {}})
    assert proc.process_one("902") == "refetch_requested"
    per_dir = {d: _all_text(d) for d in (cfg.out_dir / "survey_results.jsonl", cfg.out_dir / "survey_events.jsonl",
                                         cfg.out_dir / "predictions", cfg.state_dir / "fetch_events.jsonl")}
    for d, text in per_dir.items():
        assert SECRET not in text, d
        assert zenodo.REDACTED in text, d               # each file had the token to redact
    # the spool is scratch; 902 (written by this test with the token in a
    # path) stays there, returned, until a producer deletes it
    spool_rest = [p for p in (cfg.spool_dir / "records").iterdir() if p.name != "902"]
    assert SECRET not in _all_text(cfg.out_dir, cfg.state_dir, cfg.metadata_cache_dir, *spool_rest)
    rows = [json.loads(x) for x in (cfg.out_dir / "survey_results.jsonl").read_text().splitlines()]
    assert rows[-1]["n_classified"] == 5, rows[-1].get("error")


# ---------------------------------------------------------------------------
# D. two-pass sampling
# ---------------------------------------------------------------------------
def test_remote_triage_sample_is_a_prefix_of_the_deep_sample(tmp_path):
    from envision_eye_actionable.survey import remote
    png = _png_bytes(size=(40, 30))
    a = _zip_bytes({f"a/{i:03d}.png": png for i in range(70)})
    b = _zip_bytes({f"b/{i:03d}.png": png for i in range(13)})
    rid = "902"
    zips = [{"key": "a.zip", "size": len(a)}, {"key": "b.zip", "size": len(b)}]
    fake = _FakeZenodo(rid, {"a.zip": a, "b.zip": b})
    c = _client(tmp_path, fake)
    tri = remote.sample_remote_zips(c, rid, zips, tmp_path / "t", 7, random.Random("42:902"))
    got = {tuple(d.split("!/", 1)) for _, d in tri.fetched}
    assert len(got) == 7 and tri.more_available
    n_calls = len(fake.calls)
    listings = [remote.listing_from_dict(remote.listing_to_dict(lst)) for lst in tri.listings]
    deep = remote.sample_remote_zips(c, rid, zips, tmp_path / "d", 40, random.Random("42:902"), listings=listings,
                                     floor_cap=7, skip_members=got,
                                     before={"bytes_received": tri.bytes_received})
    sampled = {(lst.key, m.name) for lst, m in deep.sampled}
    assert got <= sampled and len(sampled) >= 40               # the triage sample is kept
    assert deep.n_reused == 7 and len(deep.fetched) == len(sampled) - 7
    assert not {tuple(d.split("!/", 1)) for _, d in deep.fetched} & got     # never fetched twice
    new_calls = fake.calls[n_calls:]
    assert not any(u.endswith("/container") for u, _ in new_calls)          # listings reused
    assert len(new_calls) == len(deep.fetched)
    # largest-remainder rounding can give a zip fewer members at a larger
    # cap (6, 6, 2 at 10 -> 4, 4, 2; at 11 -> 5, 5, 1): floor_cap keeps them
    lsts = []
    for key, n in (("a.zip", 6), ("b.zip", 6), ("c.zip", 2)):
        lst = remote.ZipListing(key=key, source="container")
        lst.members = [remote._member(f"{key}/{i}.png", 10) for i in range(n)]
        lsts.append(lst)
    small = {(lst.key, m.name) for lst, m in remote.sample_members(lsts, 10, random.Random(5))}
    big = {(lst.key, m.name) for lst, m in remote.sample_members(lsts, 11, random.Random(5), floor_cap=10)}
    assert small <= big and sum(1 for k, _ in small if k == "c.zip") == 2
    # the same seed with a larger cap from scratch draws a superset too
    for cap in (1, 5, 9, 30):
        small = {(lst.key, m.name) for lst, m in
                 remote.sample_members(listings, cap, random.Random(3))}
        big = {(lst.key, m.name) for lst, m in
               remote.sample_members(listings, 50, random.Random(3), floor_cap=cap)}
        assert small <= big


def _ir_zip(n):
    dark = _png_bytes(size=(64, 48), color=(40, 40, 40))
    return _zip_bytes({f"ir/{i:03d}.png": dark for i in range(n)})


def test_run_mode_deep_pass_only_for_eye_records(tmp_path):
    pytest.importorskip("jsonschema")
    rid = "903"
    files = {"ir.zip": _ir_zip(60)}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=30, triage_cap=8)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "ok", row.get("error")
    assert row["sampling_pass"] == "deep" and row["n_triage_classified"] == 8
    assert row["triage_eye_fraction"] == 1.0 and row["n_classified"] == 30
    assert row["n_remote_fetched"] == 30 and row["n_remote_reused"] == 8
    member_calls = [u for u, _ in s.zenodo.session.calls if "/container/" in u]
    listing_calls = [u for u, _ in s.zenodo.session.calls if u.endswith("/container")]
    assert len(member_calls) == 30 and len(listing_calls) == 1          # triage files reused, not refetched
    assert row["_class_detail"]["IR"]["est_files"] == 60

    rid = "904"
    bright = _png_bytes(size=(64, 48), color=(220, 220, 220))
    files = {"neg.zip": _zip_bytes({f"n/{i:03d}.png": bright for i in range(60)})}
    (tmp_path / "b").mkdir()
    s = _survey(tmp_path / "b", _legacy(rid, files), files, rid, remote_cap=30, triage_cap=8)
    s.clf = _BrightNegDarkOctClf()
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["status"] == "no_eye_images" and row["sampling_pass"] == "triage"
    assert row["n_classified"] == 8 and row["triage_eye_fraction"] == 0.0
    assert "0.000 < 0.05" in row["deep_pass_reason"]
    assert sum(1 for u, _ in s.zenodo.session.calls if "/container/" in u) == 8


def test_triage_sample_holds_triage_pass_entries_only(tmp_path):
    """A record run again with its deep files keeps its triage decision: the
    triage sample never takes deep-pass images, even when the triage pass
    fetched fewer images than the triage cap."""
    from envision_eye_actionable.survey.archives import RecordWalker
    rid = "907"
    s = _survey(tmp_path, _legacy(rid, {}), {}, rid, triage_cap=5, max_images=100)
    dark = _png_bytes(size=(64, 48), color=(40, 40, 40))
    w = RecordWalker(scratch=tmp_path / "w")
    (tmp_path / "w").mkdir()
    entry_pass = {}
    for i in range(13):
        p = tmp_path / "w" / f"{i:02d}.png"
        p.write_bytes(dark)
        w.add_file(p, p.name)
        entry_pass[id(w.entries[-1])] = "triage" if i < 3 else "deep"
    row = {}
    out = s._classify(row, rid, w, entry_pass=entry_pass,
                      triage={"cap": 5, "fetch_available": False, "allow_request": True})
    assert out is None and row["n_triage_classified"] == 3 and row["n_classified"] == 13
    assert row["sampling_pass"] == "deep"


def test_local_files_get_the_deep_sample_without_a_fetch(tmp_path):
    pytest.importorskip("jsonschema")
    rid = "905"
    data = _ir_zip(40)
    s = _survey(tmp_path, _legacy(rid, {"local.zip": data}), {}, rid, max_images=25, triage_cap=5)
    (s.cfg.downloads_dir / rid).mkdir(parents=True)
    (s.cfg.downloads_dir / rid / "local.zip").write_bytes(data)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["sampling_mode"] == "local" and row["sampling_pass"] == "deep"
    assert row["n_triage_classified"] == 5 and row["n_classified"] == 25
    assert (s.cfg.downloads_dir / rid / "local.zip").read_bytes() == data      # original untouched


def test_top_level_images_two_passes(tmp_path):
    pytest.importorskip("jsonschema")
    rid = "906"
    dark = _png_bytes(size=(64, 48), color=(40, 40, 40))
    files = {f"img_{i:02d}.png": dark for i in range(20)}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=12, triage_cap=4)
    row = s.process_record({"source_id": rid, "title": "T"})
    assert row["sampling_pass"] == "deep" and row["n_classified"] == 12, row.get("error")
    assert row["n_triage_classified"] == 4
    assert row["n_files_to_download"] == 12 and row["n_toplevel_images_not_fetched"] == 8
    downloads = [u for u, _ in s.zenodo.session.calls if u.endswith("/content")]
    assert len(downloads) == 12                                            # 4 triage + 8 deep, none twice


# ---------------------------------------------------------------------------
# spool
# ---------------------------------------------------------------------------
def test_spool_guards_states_and_validation(tmp_path):
    from envision_eye_actionable.survey import spool as sp
    dl = tmp_path / "downloads"
    dl.mkdir()
    for bad in (dl, dl / "x", tmp_path):
        with pytest.raises(ValueError, match="overlaps"):
            sp.Spool(bad, downloads_dir=dl)
    with pytest.raises(ValueError, match="overlaps output"):
        sp.Spool(tmp_path / "out" / "spool", out_dir=tmp_path / "out")
    data = tmp_path / "data"
    data.mkdir()
    (data / "precious.txt").write_text("keep", encoding="utf-8")
    with pytest.raises(ValueError, match="not created by the survey"):
        sp.Spool(data)
    s = sp.Spool(tmp_path / "spool", downloads_dir=dl)
    with pytest.raises(ValueError):
        s.record_dir("../etc")
    assert s.state("7") == "absent"
    d = s.record_dir("7")
    (d / "dl").mkdir(parents=True)
    (d / "dl" / "a.png").write_bytes(b"x" * 10)
    assert s.state("7") == "fetching" and s.ready_records() == []
    m = {"items": [{"path": "dl/a.png", "size": 10}, {"path": str(dl / "orig.zip"), "local": True}]}
    (dl / "orig.zip").write_bytes(b"z")
    s.write_manifest("7", m)
    assert s.state("7") == "ready" and s.ready_records() == ["7"] and s.validate("7", m) == []
    (d / "dl" / "a.png").write_bytes(b"x" * 11)
    assert any("size changed" in p for p in s.validate("7", m))
    assert any("outside" in p for p in s.validate("7", {"items": [{"path": "../../x", "size": 1}]}))
    s.mark_awaiting_deep("7")
    assert s.state("7") == "awaiting_deep" and s.ready_records() == []
    s.write_manifest("7", {"items": []}, deep=True)
    assert s.state("7") == "deep_ready" and s.ready_records() == ["7"]
    s.remove_deep("7")
    assert s.state("7") == "awaiting_deep"
    s.request("7", "deep")
    assert [r["kind"] for r in s.pending_requests()] == ["deep"]
    s.clear_request("7")
    assert s.pending_requests() == []
    assert s.bytes_used() >= 11
    # removal: only record dirs under records/, never through a symlink
    with pytest.raises(ValueError):
        s._safe_target(tmp_path / "data")
    link = s.records / "8"
    try:
        link.symlink_to(data, target_is_directory=True)
    except OSError:
        link = None
    if link is not None:
        with pytest.raises(ValueError, match="symlink"):
            s.remove("8")
        assert (data / "precious.txt").exists()
        link.unlink()
    s.remove("7")
    assert s.state("7") == "absent" and (dl / "orig.zip").exists()


# ---------------------------------------------------------------------------
# RoomGate: disk high-water mark
# ---------------------------------------------------------------------------
class _FakeSpool:
    def __init__(self, used, root):
        self.used, self.root = used, root

    def bytes_used(self):
        return self.used


def test_room_gate_waits_for_the_consumer_and_stops_without_one(tmp_path, monkeypatch):
    import shutil
    from collections import namedtuple

    from envision_eye_actionable.survey import fetch
    DU = namedtuple("DU", "total used free")
    free = {"v": 1000}
    monkeypatch.setattr(shutil, "disk_usage", lambda p: DU(10_000, 0, free["v"]))
    sp = _FakeSpool(500, tmp_path)
    gate = fetch.RoomGate(sp, spool_max_bytes=600, floor_bytes=100, poll_s=0.01)
    assert gate.acquire(50)                      # 500 + 50 <= 600, 1000 - 50 >= 100
    gate.release()
    # over the spool budget: waits until the consumer drains the spool
    waits = []

    def drain(info):
        waits.append(info)
        if len(waits) == 3:
            sp.used = 0

    gate = fetch.RoomGate(sp, 600, 100, poll_s=0.01, on_wait=drain)
    sp.used = 590
    assert gate.acquire(50) and len(waits) == 3 and gate.reserved() == 50
    gate.release()
    # deep passes are not held back by the spool budget, only by the floor
    sp.used = 590
    assert fetch.RoomGate(sp, 600, 100, poll_s=0.01, should_stop=lambda: True).acquire(50, deep=True)
    # below the floor with the spool empty: no room at all
    sp.used = 0
    free["v"] = 120
    assert not fetch.RoomGate(sp, 600, 100, poll_s=0.01).acquire(50)
    # consumer gone while the spool is full: stop instead of waiting for ever
    free["v"] = 1000
    sp.used = 590
    gate = fetch.RoomGate(sp, 600, 100, poll_s=0.01, should_stop=lambda: True)
    assert not gate.acquire(50) and gate.stopped
    # a record larger than the budget goes once nothing else is spooled or in flight
    sp.used = 0
    assert fetch.RoomGate(sp, 600, 100, poll_s=0.01).acquire(800)


# ---------------------------------------------------------------------------
# C. producer / consumer
# ---------------------------------------------------------------------------
def test_pipeline_roles_fetch_classify_and_clean_the_spool(tmp_path, monkeypatch):
    pytest.importorskip("jsonschema")
    dark = _png_bytes(size=(64, 48), color=(40, 40, 40))
    bright = _png_bytes(size=(64, 48), color=(220, 220, 220))
    records = {
        "1001": {"ir.zip": _zip_bytes({f"ir/{i:03d}.png": dark for i in range(30)})},      # eye: deep pass
        "1002": {"neg.zip": _zip_bytes({f"n/{i:03d}.png": bright for i in range(30)})},  # NEG: triage only
        "1003": {"paper.pdf": b"%PDF-1.4 x"},                                             # nothing to read
        "1004": {"local.zip": b""},                                                       # local original
    }
    local = _zip_bytes({f"l/{i}.png": dark for i in range(3)})
    records["1004"] = {"local.zip": local}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, triage_cap=5, remote_cap=12)
    (cfg.downloads_dir / "1004").mkdir(parents=True)
    orig = cfg.downloads_dir / "1004" / "local.zip"
    orig.write_bytes(local)
    mtime = orig.stat().st_mtime_ns
    session = _Multi({rid: _FakeZenodo(rid, files) for rid, files in records.items()})
    out, prod, proc = _run_roles(cfg, session, monkeypatch, onnx=_NegOnnxStub)
    assert out == {"fetch": 0, "process": 0}
    from envision_eye_actionable.survey.results import load_results
    rows = load_results(cfg.out_dir / "survey_results.jsonl")
    assert set(rows) == set(records)
    assert rows["1001"]["status"] == "ok" and rows["1001"]["sampling_pass"] == "deep"
    assert rows["1001"]["n_classified"] == 12 and rows["1001"]["n_remote_reused"] == 5
    assert rows["1002"]["status"] == "no_eye_images" and rows["1002"]["sampling_pass"] == "triage"
    assert rows["1002"]["n_classified"] == 5
    assert rows["1003"]["status"] == "no_image_files"
    assert rows["1004"]["sampling_mode"] == "local" and rows["1004"]["n_classified"] == 3
    assert (cfg.out_dir / "cmds" / "1001" / "dataset_description.json").is_file()
    assert (cfg.out_dir / "predictions" / "1001.jsonl.gz").is_file()
    assert not list((cfg.out_dir / "predictions").glob("*.part"))
    # the spool is empty, the local original untouched, no request left
    assert list((cfg.spool_dir / "records").iterdir()) == []
    assert list((cfg.spool_dir / "requests").iterdir()) == []
    assert orig.read_bytes() == local and orig.stat().st_mtime_ns == mtime
    lines = [json.loads(x) for x in (cfg.out_dir / "survey_results.jsonl").read_text().splitlines()]
    assert [x["status"] for x in lines if x["record_id"] == "1001"][:2] == ["started", "awaiting_deep"]
    ev = [json.loads(x)["event"] for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    assert ev.count("fetched_deep") == 1 and ev[-1] == "fetch_end"
    st = json.loads((cfg.state_dir / "process_status.json").read_text())
    assert st["state"] == "finished" and st["counters"]["row"] == 4


def test_consumer_deletes_the_spool_only_after_the_row_is_written(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import pipeline, runner
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {"1101": {"a.pdf": b"x"}})
    cfg = _cfg(tmp_path, meta)
    proc = pipeline.Processor(cfg)
    s = proc.spool
    s.record_dir("1101").mkdir(parents=True)
    s.write_manifest("1101", {"items": [], "terminal": {"status": "no_files", "finish": "without_images"},
                              "row": {}})

    def boom(row):
        if row.get("status") not in ("started",):
            raise OSError("disk full")
        return orig_append(row)

    orig_append = proc.survey._append_row
    monkeypatch.setattr(proc.survey, "_append_row", boom)
    with pytest.raises(OSError):
        proc.process_one("1101")
    assert s.state("1101") == "ready"                     # still there to process again
    monkeypatch.setattr(proc.survey, "_append_row", orig_append)
    assert proc.process_one("1101") == "row" and s.state("1101") == "absent"


def test_consumer_crash_loop_guard_and_invalid_spool_refetch(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import pipeline, runner
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {"1201": {"a.png": b"x"}, "1202": {"b.png": b"y"}})
    cfg = _cfg(tmp_path, meta, max_attempts=2)
    proc = pipeline.Processor(cfg)
    s = proc.spool
    for rid in ("1201", "1202"):
        s.record_dir(rid).mkdir(parents=True)
    (s.record_dir("1201") / "dl").mkdir()
    (s.record_dir("1201") / "dl" / "a.png").write_bytes(b"x")
    s.write_manifest("1201", {"items": [{"path": "dl/a.png", "size": 1, "role": "download", "key": "a.png"}],
                              "row": {}})
    (s.record_dir("1201") / ".attempts").write_text("2\n")   # died inside it twice already
    assert proc.process_one("1201") == "crashed"
    rows = {json.loads(x)["record_id"]: json.loads(x) for x in
            (cfg.out_dir / "survey_results.jsonl").read_text().splitlines()}
    assert rows["1201"]["status"] == "crashed" and s.state("1201") == "absent"
    # a spooled file that went missing: the record is dropped and asked for again
    s.write_manifest("1202", {"items": [{"path": "dl/b.png", "size": 1, "role": "download", "key": "b.png"}],
                              "row": {}})
    assert proc.process_one("1202") == "refetch_requested"
    # back to the producer: the consumer leaves the dir (RETURNED, no marker) for it to delete
    assert s.state("1202") == "fetching" and s.returned("1202")
    assert [r["kind"] for r in s.pending_requests()] == ["refetch"]


def test_producer_skips_spooled_and_finished_records_and_redoes_partial_ones(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import pipeline
    png = _png_bytes(size=(64, 48))
    records = {"1301": {"a.png": png}, "1302": {"b.png": png}, "1303": {"c.png": png}, "1304": {"d.png": png}}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, consumer_grace_s=0.0)
    cfg.out_dir.mkdir(parents=True)
    (cfg.out_dir / "survey_results.jsonl").write_text(
        json.dumps({"record_id": "1301", "status": "ok"}) + "\n"
        + json.dumps({"record_id": "1304", "status": "awaiting_deep"}) + "\n", encoding="utf-8")
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = _Multi({rid: _FakeZenodo(rid, f) for rid, f in records.items()})
    prod.spool.record_dir("1302").mkdir(parents=True)
    prod.spool.write_manifest("1302", {"items": [], "row": {}})
    partial = prod.spool.record_dir("1303")
    partial.mkdir(parents=True)
    (partial / "junk.part").write_bytes(b"half")
    assert prod.run() == 0
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    fetched = sorted(e["record"] for e in ev if e["event"] == "fetched")
    assert fetched == ["1303", "1304"]                 # 1301 done, 1302 already spooled
    assert any(e["event"] == "partial_spool_removed" and e["record"] == "1303" for e in ev)
    assert not (partial / "junk.part").exists() and prod.spool.state("1303") == "ready"


def test_producer_stops_when_the_consumer_is_gone_and_the_spool_is_full(tmp_path):
    from envision_eye_actionable.survey import pipeline
    png = _png_bytes(size=(64, 48))
    records = {"1401": {"a.png": png}, "1402": {"b.png": png}}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, consumer_grace_s=0.0, spool_max_gb=1e-9, fetch_workers=1)
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = _Multi({rid: _FakeZenodo(rid, f) for rid, f in records.items()})
    assert prod.run() == pipeline.EXIT_CONSUMER_GONE
    states = prod.spool.states()
    assert list(states.values()).count("ready") == 1 and len(states) == 1   # the second was not written
    st = json.loads((cfg.state_dir / "fetch_status.json").read_text())
    assert st["state"] == "stopped"


def test_monitor_snapshot_counts_rates_and_eta(tmp_path):
    from envision_eye_actionable.survey import pipeline
    meta = _write_records(tmp_path, {str(1500 + i): {"a.pdf": b"x"} for i in range(4)})
    cfg = _cfg(tmp_path, meta)
    cfg.out_dir.mkdir(parents=True)
    res = cfg.out_dir / "survey_results.jsonl"
    now = time.strftime("%Y-%m-%dT%H:%M:%S+00:00", time.gmtime())
    res.write_text(json.dumps({"record_id": "1500", "status": "ok", "finished_at": now}) + "\n"
                   + json.dumps({"record_id": "1501", "status": "started"}) + "\n", encoding="utf-8")
    mon = pipeline.Monitor(cfg)
    snap = mon.snapshot()
    assert snap["results"]["n_final"] == 1 and snap["results"]["in_progress"] == 1
    assert snap["results"]["remaining"] == 3 and snap["results"]["eta_hours"] == 3.0
    assert snap["roles"]["fetch"]["alive"] is False and snap["spool"]["bytes_gb"] == 0
    # a role killed while running never wrote its end state: shown as dead
    (cfg.state_dir / "fetch_status.json").write_text(json.dumps({"state": "running"}), encoding="utf-8")
    assert mon.snapshot()["roles"]["fetch"]["state"] == "dead (last state running)"
    with open(res, "a", encoding="utf-8") as fp:           # incremental: only new lines are read
        fp.write(json.dumps({"record_id": "1501", "status": "no_files", "finished_at": now}) + "\n"
                 + '{"record_id": "15')                    # a line being written
    snap = mon.snapshot()
    assert snap["results"]["n_final"] == 2 and snap["results"]["in_progress"] == 0
    with open(res, "a", encoding="utf-8") as fp:           # ... finished later
        fp.write('02", "status": "ok", "finished_at": "%s"}\n' % now)
    assert mon.snapshot()["results"]["n_final"] == 3
    assert mon.run(0.01, once=True) == 0
    status = json.loads((cfg.state_dir / "status.json").read_text())
    assert status["results"]["by_status"] == {"ok": 2, "no_files": 1}


def test_pipeline_supervisor_restarts_a_crashed_role_up_to_the_cap(tmp_path, monkeypatch):
    """process cannot start (no such model): each crash is logged and
    process is restarted --max-role-restarts times, then given up; fetch
    goes on and stops once the consumer has been gone the grace time."""
    from envision_eye_actionable.survey import pipeline
    meta = _write_records(tmp_path, {"1601": {"a.pdf": b"x"}})
    cfg = _cfg(tmp_path, meta, offline=True, max_role_restarts=2, restart_backoff_s=0.1, poll_s=0.2)
    argv = ["--scrape", str(tmp_path / "scrape.json"), "--model", str(tmp_path / "missing.onnx"),
            "--metadata-dir", str(meta), "--out-dir", str(cfg.out_dir), "--spool-dir", str(cfg.spool_dir),
            "--state-dir", str(cfg.state_dir), "--downloads-dir", str(cfg.downloads_dir), "--offline",
            "--poll-s", "0.2", "--consumer-grace-s", "1", "--interval", "0.5", "--no-link-check"]
    monkeypatch.setenv("HOME", str(tmp_path))          # the children never see a real token file
    code = pipeline.run_pipeline(cfg, argv, 0.5)
    assert code == 1
    status = json.loads((cfg.state_dir / "pipeline_status.json").read_text())
    assert status["finished"] and "process" in status["note"] and "given up: process" in status["note"]
    assert status["roles"]["process"]["exit_code"] not in (0, None) and status["roles"]["fetch"]["exit_code"] == 0
    ev = [json.loads(x) for x in (cfg.state_dir / "pipeline_events.jsonl").read_text().splitlines()]
    assert sum(1 for e in ev if e["event"] == "role_started" and e["role"] == "process") == 3   # 1 + 2 restarts
    assert sum(1 for e in ev if e["event"] == "role_crashed" and e["role"] == "process") == 3
    assert any(e["event"] == "role_given_up" and e["role"] == "process" for e in ev)
    assert (cfg.state_dir / "logs" / "process.log").read_text()


def test_pipeline_restarts_a_killed_consumer_and_the_record_that_kills_it_gets_crashed(tmp_path, monkeypatch):
    """The consumer dies (os._exit, as an OOM kill) inside record 1801 on
    every start: the supervisor restarts it, the attempts guard writes a
    crashed row for 1801 on the third start, 1802 is classified, and the
    pipeline ends on its own."""
    import sys
    from envision_eye_actionable.survey import pipeline
    from envision_eye_actionable.survey.results import load_results
    meta = _write_records(tmp_path, {"1801": {"a.pdf": b"x"}, "1802": {"b.pdf": b"y"}})
    model = tmp_path / "m.onnx"
    model.write_bytes(b"stub")
    cfg = _cfg(tmp_path, meta, offline=True, max_role_restarts=5, restart_backoff_s=0.1, poll_s=0.2,
               max_attempts=2)
    script = tmp_path / "role.py"
    script.write_text(
        "import os, sys\n"
        f"sys.path.insert(0, {str(Path(__file__).parent)!r})\n"
        "from envision_eye_actionable.survey import cli, runner\n"
        "import test_pipeline as tp\n"
        "runner.OnnxClassifier = tp._OnnxStub\n"
        "if sys.argv[1] == 'process':\n"
        "    orig = runner.Survey.finish_record\n"
        "    def finish(self, rec, *a, **k):\n"
        "        if rec['source_id'] == '1801':\n"
        "            os._exit(137)\n"
        "        return orig(self, rec, *a, **k)\n"
        "    runner.Survey.finish_record = finish\n"
        "sys.exit(cli.main(sys.argv[1:]))\n", encoding="utf-8")
    argv = ["--scrape", str(tmp_path / "scrape.json"), "--model", str(model),
            "--metadata-dir", str(meta), "--out-dir", str(cfg.out_dir), "--spool-dir", str(cfg.spool_dir),
            "--state-dir", str(cfg.state_dir), "--downloads-dir", str(cfg.downloads_dir), "--offline",
            "--poll-s", "0.2", "--consumer-grace-s", "60", "--interval", "0.5", "--no-link-check",
            "--max-attempts", "2"]
    monkeypatch.setenv("HOME", str(tmp_path))
    code = pipeline.run_pipeline(cfg, argv, 0.5, base_cmd=[sys.executable, str(script)])
    ev = [json.loads(x) for x in (cfg.state_dir / "pipeline_events.jsonl").read_text().splitlines()]
    rows = load_results(cfg.out_dir / "survey_results.jsonl")
    assert rows["1801"]["status"] == "crashed", rows["1801"]
    assert rows["1802"]["status"] == "no_image_files"
    assert sum(1 for e in ev if e["event"] == "role_crashed" and e["role"] == "process") == 2
    assert sum(1 for e in ev if e["event"] == "role_restarted" and e["role"] == "process") == 2
    assert code == 0 and ev[-1]["event"] == "pipeline_end" and not ev[-1]["given_up"]
    assert list((cfg.spool_dir / "records").iterdir()) == []


def _deep_stub(prod, served: list):
    def fake(rec, rdir, m, listings):
        served.append(rec["source_id"])
        return {"items": [], "row": {}, "remote": None, "unfetched_top": [], "bytes": 0, "fetch_s": 0}
    prod.fetcher.fetch_deep = fake


def _spool_record(spool, rid: str, state: str, nbytes: int = 0):
    d = spool.record_dir(rid)
    d.mkdir(parents=True)
    if nbytes:
        (d / "dl").mkdir()
        (d / "dl" / "x.bin").write_bytes(b"\0" * nbytes)
    spool.write_manifest(rid, {"items": [], "row": {}, "top_pending": []})
    if state in ("awaiting_deep", "deep_ready"):
        spool.mark_awaiting_deep(rid)
    if state == "deep_ready":
        spool.write_manifest(rid, {"items": [], "row": {}}, deep=True)


def test_producer_resume_leaves_spooled_records_to_the_consumer(tmp_path, monkeypatch):
    """Records a killed run left ready / awaiting_deep / deep_ready are not
    in the queue: when the consumer has finished one before the queue gets
    there, it is not fetched again. A queued record that got a final row
    since start is skipped too."""
    from envision_eye_actionable.survey import pipeline
    png = _png_bytes(size=(64, 48))
    rids = ["1901", "1902", "1903", "1904", "1905"]
    records = {r: {f"{r}.png": png} for r in rids}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, consumer_grace_s=0.0, fetch_workers=1, order="scrape")
    cfg.out_dir.mkdir(parents=True)
    res = cfg.out_dir / "survey_results.jsonl"
    res.write_text(json.dumps({"record_id": "1902", "status": "started"}) + "\n", encoding="utf-8")
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = _Multi({rid: _FakeZenodo(rid, f) for rid, f in records.items()})
    _spool_record(prod.spool, "1902", "ready")
    _spool_record(prod.spool, "1903", "deep_ready")
    prod.spool.request("1904", "deep")
    _spool_record(prod.spool, "1904", "awaiting_deep")
    served = []
    _deep_stub(prod, served)
    orig = prod.fetcher.fetch

    def fetch(rec, rdir):
        if rec["source_id"] == "1901":
            # the consumer finishes the spooled records meanwhile, and 1905
            # gets a final row (from a refetch) before the queue reaches it
            for rid in ("1902", "1903"):
                prod.spool.remove(rid)
            with open(res, "a", encoding="utf-8") as fp:
                for rid in ("1902", "1903", "1905"):
                    fp.write(json.dumps({"record_id": rid, "status": "ok"}) + "\n")
        return orig(rec, rdir)

    prod.fetcher.fetch = fetch
    assert prod.run() == 0
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    assert sorted(e["record"] for e in ev if e["event"] == "fetched") == ["1901"]
    assert served == ["1904"]
    start = next(e for e in ev if e["event"] == "fetch_start")
    assert start["n_todo"] == 2                        # 1901 and 1905: the spooled ones belong to the consumer


def test_awaiting_deep_without_a_request_gets_it_back(tmp_path):
    """The consumer died between the AWAITING_DEEP marker and its deep
    request: the producer restores the request and serves it."""
    from envision_eye_actionable.survey import pipeline
    meta = _write_records(tmp_path, {"2001": {"a.pdf": b"x"}})
    cfg = _cfg(tmp_path, meta, consumer_grace_s=0.0, offline=True)
    cfg.out_dir.mkdir(parents=True)
    prod = pipeline.Producer(cfg)
    _spool_record(prod.spool, "2001", "awaiting_deep")
    served = []
    _deep_stub(prod, served)
    assert prod.run() == 0
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    assert any(e["event"] == "deep_request_restored" and e["record"] == "2001" for e in ev)
    assert served == ["2001"] and prod.spool.state("2001") == "deep_ready"
    assert prod.spool.pending_requests() == []


def test_a_request_written_right_after_publishing_survives(tmp_path):
    """The consumer may process a record the moment its marker is written
    and ask for a refetch; the producer must not delete that new request."""
    from envision_eye_actionable.survey import pipeline
    png = _png_bytes(size=(64, 48))
    records = {"2101": {"a.png": png}, "2102": {"b.png": png}}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta)
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = _Multi({rid: _FakeZenodo(rid, f) for rid, f in records.items()})
    _deep_stub(prod, [])
    orig = prod.spool.write_manifest

    def publish_then_consumer_asks(rid, m, deep=False):
        orig(rid, m, deep=deep)
        prod.spool.request(rid, "refetch", {"refetches": 9})     # the consumer, right after the marker

    prod.spool.write_manifest = publish_then_consumer_asks
    _spool_record(prod.spool, "2101", "awaiting_deep")
    prod.spool.request("2101", "deep")
    prod._deep({"source_id": "2101"})
    assert [(r["record_id"], r["kind"]) for r in prod.spool.pending_requests()] == [("2101", "refetch")]
    prod.spool.clear_request("2101")
    prod.spool.request("2102", "refetch", {"refetches": 1})
    prod._refetch({"source_id": "2102", "_refetches": 1})
    assert [(r["record_id"], r.get("refetches")) for r in prod.spool.pending_requests()] == [("2102", 9)]


def test_producer_does_not_end_while_the_consumer_is_between_delete_and_request(tmp_path):
    """The consumer deletes an invalid record and then asks for it again:
    a producer that looked in between must look once more before it ends."""
    import sys
    from envision_eye_actionable.survey import pipeline
    meta = _write_records(tmp_path, {"2301": {"a.pdf": b"x"}})
    cfg = _cfg(tmp_path, meta, offline=True, consumer_grace_s=600.0)
    cfg.out_dir.mkdir(parents=True)
    (cfg.out_dir / "survey_results.jsonl").write_text(json.dumps({"record_id": "2301", "status": "ok"}) + "\n")
    prod = pipeline.Producer(cfg)
    consumer = pipeline.RoleLock(cfg.state_dir, "process")
    orig = prod._consumer_gone
    calls = {"run": 0}

    def consumer_gone():
        if sys._getframe(1).f_code.co_name == "_run":
            calls["run"] += 1
            if calls["run"] == 1:                  # the first end check found nothing left ...
                prod.spool.request("2301", "refetch", {"refetches": 1})   # ... and now the request lands
        return orig()

    prod._consumer_gone = consumer_gone
    stop = threading.Event()

    def fake_consumer():
        while not stop.is_set():
            for rid in prod.spool.ready_records():
                prod.spool.remove(rid)
            time.sleep(0.02)

    t = threading.Thread(target=fake_consumer, daemon=True)
    t.start()
    try:
        assert prod.run() == 0
    finally:
        stop.set()
        consumer.close()
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    assert any(e["event"] == "fetched" and e["record"] == "2301" for e in ev)
    assert prod.spool.pending_requests() == []


def test_room_gate_reserves_atomically_and_yields_to_requests(tmp_path, monkeypatch):
    import shutil
    from collections import namedtuple

    from envision_eye_actionable.survey import fetch
    DU = namedtuple("DU", "total used free")
    monkeypatch.setattr(shutil, "disk_usage", lambda p: DU(10_000, 0, 5_000))

    class SlowSpool(_FakeSpool):
        def bytes_used(self):
            time.sleep(0.05)                       # both threads read before either reserves
            return self.used

    gate = fetch.RoomGate(SlowSpool(0, tmp_path), spool_max_bytes=100, floor_bytes=0, poll_s=0.02)
    peak, got = [], []
    lock = threading.Lock()

    def job():
        assert gate.acquire(60)
        with lock:
            got.append(1)
            peak.append(gate.reserved())
        time.sleep(0.2)
        gate.release()

    threads = [threading.Thread(target=job) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(10)
    assert len(got) == 2 and max(peak) == 60       # never 120 reserved against a budget of 100
    # a triage job waiting for room gives way to a request; a deep pass never does
    sp = _FakeSpool(500, tmp_path)
    gate = fetch.RoomGate(sp, 100, 0, poll_s=0.01, should_yield=lambda: True)
    with pytest.raises(fetch.YieldToRequests):
        gate.acquire(50)
    assert gate.waiting == 0 and gate.reserved() == 0
    assert gate.acquire(50, deep=True)


def test_waiting_triage_jobs_give_their_slot_to_a_deep_request(tmp_path):
    """All fetch slots wait for room behind a spool full of a record that
    awaits its deep pass; the deep request must still be served (it frees
    the spool), or both roles wait for ever."""
    from envision_eye_actionable.survey import pipeline
    png = _png_bytes(size=(64, 48))
    records = {"2201": {"a.png": png}, "2202": {"b.png": png}}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, fetch_workers=1, spool_max_gb=5e-6, consumer_grace_s=600.0)   # 5 kB
    cfg.out_dir.mkdir(parents=True)
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = _Multi({"2202": _FakeZenodo("2202", records["2202"])})
    _spool_record(prod.spool, "2201", "ready", nbytes=20_000)
    served = []
    _deep_stub(prod, served)
    consumer = pipeline.RoleLock(cfg.state_dir, "process")      # the consumer is alive
    stop = threading.Event()

    def fake_consumer():
        while not stop.is_set():
            for rid in prod.spool.ready_records():
                if rid == "2201" and prod.spool.state(rid) == "ready":
                    if prod.gate.waiting:                          # every slot waits for room: now
                        prod.spool.mark_awaiting_deep(rid)         # its triage sample wants more
                        prod.spool.request(rid, "deep")
                    continue
                prod.spool.remove(rid)                             # "processed"
            time.sleep(0.02)

    t = threading.Thread(target=fake_consumer, daemon=True)
    t.start()
    out = {}
    runner = threading.Thread(target=lambda: out.setdefault("code", prod.run()), daemon=True)
    runner.start()
    runner.join(30)
    stop.set()
    consumer.close()
    assert not runner.is_alive(), "producer livelocked behind the spool budget"
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    assert served == ["2201"] and any(e["event"] == "triage_yielded" and e["record"] == "2202" for e in ev)
    assert any(e["event"] == "fetched" and e["record"] == "2202" for e in ev) and out["code"] == 0


# ---------------------------------------------------------------------------
# B. metadata prefetch
# ---------------------------------------------------------------------------
def test_metadata_prefetch_caches_and_resumes(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import pipeline, zenodo
    meta = _write_records(tmp_path, {"1701": {"a.pdf": b"x"}})
    empty = tmp_path / "nometa"
    empty.mkdir()
    cfg = _cfg(tmp_path, empty)
    calls = []

    class Sess:
        headers = {}
        auth = None

        def get(self, url, headers=None, timeout=None, **k):
            calls.append((url, (headers or {}).get("Accept")))
            body = json.loads((meta / "1701.json").read_text())
            return _HttpResp(200, json.dumps(body).encode())

    real = zenodo.ZenodoClient.__init__

    def init(self, *a, **k):
        real(self, *a, **k)
        self.session = Sess()

    monkeypatch.setattr(zenodo.ZenodoClient, "__init__", init)
    out = pipeline.prefetch_metadata(cfg, threads=1)
    assert out["fetched"] == 1 and len(calls) == 2
    assert (cfg.metadata_cache_dir / "legacy" / "1701.json").is_file()
    assert (cfg.metadata_cache_dir / "datacite" / "1701.json").is_file()
    out = pipeline.prefetch_metadata(cfg, threads=1)
    assert out["cached"] == 1 and len(calls) == 2                 # resumed: no request


def test_order_puts_cheap_records_first(tmp_path):
    from envision_eye_actionable.survey import pipeline
    cfg = _cfg(tmp_path, None)
    big = {"source_id": "1", "file_names": ["scans.tar.gz"], "size_mb": 5000}
    zips = {"source_id": "2", "file_names": ["a.zip"], "size_mb": 90000}
    pdf = {"source_id": "3", "file_names": ["paper.pdf"], "size_mb": 3}
    assert pipeline.fetch_cost(pdf, None, cfg.downloads_dir, cfg) == 0
    assert pipeline.fetch_cost(zips, None, cfg.downloads_dir, cfg) == 50e6
    assert pipeline.fetch_cost(big, None, cfg.downloads_dir, cfg) == 5e9


# ---------------------------------------------------------------------------
# F. streamed workbook
# ---------------------------------------------------------------------------
def test_streamed_workbook_formats_sheet_and_cmds_paths(tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    from envision_eye_actionable.survey import excel
    out = tmp_path / "res"
    rows = []
    for i in range(1, 6):
        (out / "cmds" / str(i)).mkdir(parents=True)
        (out / "cmds" / str(i) / "dataset_description.json").write_text('{"a": %d}' % i, encoding="utf-8")
        rows.append({"record_id": str(i), "status": "ok" if i % 2 else "no_image_files", "cmds_dir": f"cmds/{i}",
                     "setfit_label": "NEGATIVE", "setfit_prob_eye_imaging": 0.1, "sampling_pass": "triage",
                     "formats_classified": {"png": 3}, "conversion_counts": {"rgb": 3},
                     "formats_unread": {"tiff": 1} if i == 2 else {}})
    res = out / "survey_results.jsonl"
    res.write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
    monkeypatch.setattr(excel, "EMBED_CMDS_MAX_ROWS", 2)
    stats = excel.build_workbook(res, out / "w.xlsx")
    assert stats["records"] == 5 and stats["cmds_json"] == "paths"
    wb = load_workbook(out / "w.xlsx", read_only=True)
    assert wb.sheetnames == ["README", "Records", "Record_Classes", "Weblinks", "Archive_Probes", "DICOM_Mapping",
                             "Schema_Fields", "CMDS_JSON", "Formats"]
    head = [c.value for c in next(wb["Records"].iter_rows(max_row=1))]
    assert any("NOT used to select" in (h or "") for h in head)
    recs = {r[0]: r for r in wb["Records"].iter_rows(min_row=2, values_only=True)}
    assert recs["3"][head.index("Survey status")] == "ok"
    cj = list(wb["CMDS_JSON"].iter_rows(min_row=1, values_only=True))
    assert "dataset_description.json file" in cj[0] and cj[1][4].endswith("cmds/1/dataset_description.json")
    fo = list(wb["Formats"].iter_rows(min_row=2, values_only=True))
    assert ("classified images by source format", "png", 15, 5) in fo
    assert ("files that gave no image, by source format", "tiff", 1, 1) in fo
    stats = excel.build_workbook(res, out / "w2.xlsx", cmds_json="embed")
    wb = load_workbook(out / "w2.xlsx", read_only=True)
    assert list(wb["CMDS_JSON"].iter_rows(min_row=2, max_row=2, values_only=True))[0][4] == '{"a": 1}'


def test_is_final_and_merge_treat_awaiting_deep_as_unfinished(tmp_path):
    from envision_eye_actionable.survey import excel
    from envision_eye_actionable.survey.results import is_final
    assert not is_final({"status": "awaiting_deep"}) and not is_final({"status": "started"})
    assert is_final({"status": "ok"}) and not is_final({"status": "error"}, ["error"]) and not is_final(None)
    a, b = tmp_path / "a.jsonl", tmp_path / "b.jsonl"
    a.write_text(json.dumps({"record_id": "1", "status": "ok"}) + "\n", encoding="utf-8")
    b.write_text(json.dumps({"record_id": "1", "status": "awaiting_deep"}) + "\n"
                 + json.dumps({"record_id": "2", "status": "started"}) + "\n", encoding="utf-8")
    idx = excel.merge_index([a, b])
    assert idx["1"][0] == 0 and idx["1"][2] == "ok" and idx["2"][2] == "started"
    assert {k: v["status"] for k, v in excel.merge_results([a, b]).items()} == {k: v[2] for k, v in idx.items()}


def test_cli_pipeline_options_and_token_file_check(tmp_path, monkeypatch, capsys):
    from envision_eye_actionable.survey import cli
    with pytest.raises(SystemExit):
        cli.main(["fetch", "--scrape", "x.json"])                 # needs --spool-dir and --state-dir
    assert "needs --spool-dir and --state-dir" in capsys.readouterr().err
    with pytest.raises(SystemExit):
        cli.main(["fetch", "--spool-dir", str(tmp_path / "s"), "--state-dir", str(tmp_path / "t"),
                  "--token-file", str(tmp_path / "nope")])
    assert "--token-file" in capsys.readouterr().err
    seen = {}
    from envision_eye_actionable.survey import pipeline

    class P:
        def __init__(self, cfg):
            seen["cfg"] = cfg

        def run(self):
            return 7

    monkeypatch.setattr(pipeline, "Producer", P)
    assert cli.main(["fetch", "--spool-dir", str(tmp_path / "s"), "--state-dir", str(tmp_path / "t"),
                     "--triage-cap", "9", "--spool-max-gb", "12", "--fetch-workers", "3"]) == 7
    cfg = seen["cfg"]
    assert (cfg.triage_cap, cfg.spool_max_gb, cfg.fetch_workers, cfg.zenodo_per_minute) == (9, 12.0, 3, 120)
    assert np is not None and _FakeClf is not None


def test_consumer_gives_up_after_repeated_refetches(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import pipeline, runner
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {"1801": {"a.png": b"x"}})
    cfg = _cfg(tmp_path, meta)
    proc = pipeline.Processor(cfg)
    s = proc.spool
    s.record_dir("1801").mkdir(parents=True)
    s.write_manifest("1801", {"items": [{"path": "dl/a.png", "size": 1, "role": "download", "key": "a.png"}],
                              "row": {}, "refetches": pipeline.MAX_REFETCHES})
    assert proc.process_one("1801") == "row"
    row = json.loads((cfg.out_dir / "survey_results.jsonl").read_text().splitlines()[-1])
    assert row["status"] == "error" and "refetches" in row["error"]
    assert s.pending_requests() == [] and s.state("1801") == "absent"


def test_downloads_with_one_base_name_keep_both_files(tmp_path):
    from envision_eye_actionable.survey.fetch import RecordFetcher
    from envision_eye_actionable.survey.runner import SurveyConfig

    class Sess:
        def get(self, url, headers=None, stream=False, timeout=None):
            return _HttpResp(200, url.encode())

    cfg = SurveyConfig(scrape_path=tmp_path / "s", model_path=tmp_path / "m", disk_floor_gb=0, download_workers=2)
    zen = _client(tmp_path, Sess())
    f = RecordFetcher(cfg, zen)
    files = [{"key": "a/x.png", "size": len(b"u1"), "url": "u1"}, {"key": "b/x.png", "size": len(b"u2"), "url": "u2"}]
    failures, items = f._download(files, tmp_path / "rec" / "dl", "1")
    assert not failures and len(items) == 2
    got = {it["key"]: (tmp_path / "rec" / it["path"]).read_bytes() for it in items}
    assert got == {"a/x.png": b"u1", "b/x.png": b"u2"}


def test_a_record_with_its_deep_pass_never_asks_for_it_again(tmp_path):
    """finish_record with the deep manifest does not request the deep pass
    again, even when the caller allows requests."""
    pytest.importorskip("jsonschema")
    rid = "1901"
    files = {"ir.zip": _ir_zip(30)}
    s = _survey(tmp_path, _legacy(rid, files), files, rid, remote_cap=12, triage_cap=4)
    f = s._fetcher()
    rdir = s.records_scratch / rid
    m = f.fetch({"source_id": rid, "title": "T"}, rdir)
    listings = m.pop("_listings")
    row, want = s.finish_record({"source_id": rid}, rdir, m, allow_deep_request=True)
    assert row is None and want
    dm = f.fetch_deep({"source_id": rid}, rdir, m, listings)
    row, want = s.finish_record({"source_id": rid}, rdir, m, dm, allow_deep_request=True)
    assert not want and row["sampling_pass"] == "deep" and row["n_classified"] == 12


# ---------------------------------------------------------------------------
# D. review fixes: concurrency, crash windows, budgets, memory
# ---------------------------------------------------------------------------
def test_disk_reservation_writes_are_serialised_and_never_stale(tmp_path, caplog):
    """Every fetch thread publishes the gate total: the published file is
    always whole JSON, no replace fails, and the last total wins."""
    from envision_eye_actionable.survey.locks import DiskReservations
    state = tmp_path / "state"
    disk = DiskReservations(state, tmp_path)
    total = {"n": 0}
    tlock = threading.Lock()
    bad = []
    stop = threading.Event()

    def reader():
        while not stop.is_set():
            try:
                json.loads(disk._file.read_text(encoding="utf-8"))
            except FileNotFoundError:
                pass
            except ValueError:
                bad.append(1)

    def worker():
        for _ in range(300):
            with tlock:
                total["n"] += 1
            disk.reserve(lambda: total["n"])

    r = threading.Thread(target=reader)
    r.start()
    with caplog.at_level(logging.WARNING, logger="envision_eye_actionable.survey.locks"):
        ts = [threading.Thread(target=worker) for _ in range(8)]
        for t in ts:
            t.start()
        for t in ts:
            t.join(60)
    stop.set()
    r.join(10)
    assert not [x for x in caplog.records if "could not write the disk reservation" in x.getMessage()]
    assert not bad
    assert json.loads(disk._file.read_text(encoding="utf-8"))["bytes"] == 2400 == disk.bytes
    disk.close()


def test_spool_is_locked_per_role_and_bound_to_its_state_dir(tmp_path):
    from envision_eye_actionable.survey import pipeline
    spool = tmp_path / "spool"
    spool.mkdir()
    a = pipeline.RoleLock(tmp_path / "stA", "fetch", spool)
    with pytest.raises(ValueError, match="belongs to state dir"):
        pipeline.RoleLock(tmp_path / "stB", "fetch", spool)
    with pytest.raises(ValueError, match="belongs to state dir"):
        pipeline.RoleLock(tmp_path / "stB", "process", spool)
    (spool / pipeline.SPOOL_OWNER).unlink()                  # even without the owner file ...
    with pytest.raises(ValueError, match="with spool dir"):  # ... the spool's own fetch lock is held
        pipeline.RoleLock(tmp_path / "stB", "fetch", spool)
    a.close()
    b = pipeline.RoleLock(tmp_path / "stA", "fetch", spool)  # freed: the same state dir starts again
    b.close()


def test_producers_with_one_spool_and_two_state_dirs_refuse(tmp_path):
    from envision_eye_actionable.survey import pipeline
    meta = _write_records(tmp_path, {"2501": {"a.pdf": b"x"}})
    cfg = _cfg(tmp_path, meta, offline=True)
    first = pipeline.Producer(cfg)
    cfg2 = _cfg(tmp_path, meta, offline=True, state_dir=tmp_path / "state2")
    with pytest.raises(ValueError, match="belongs to state dir"):
        pipeline.Producer(cfg2)
    first.lock.close()


def test_invalid_spool_goes_back_without_an_orphan_window(tmp_path, monkeypatch):
    """The consumer takes the markers away before it writes the request and
    never deletes the dir itself; a kill between the two steps leaves a dir
    whose request the producer restores; the producer then refetches it."""
    from envision_eye_actionable.survey import pipeline, runner
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    png = _png_bytes(size=(64, 48))
    records = {"2601": {"a.png": png}, "2602": {"b.png": png}}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, consumer_grace_s=0.0)
    proc = pipeline.Processor(cfg)
    s = proc.spool
    for rid in records:
        s.record_dir(rid).mkdir(parents=True)
        s.write_manifest(rid, {"items": [{"path": "dl/gone.png", "size": 1, "role": "download", "key": "x"}],
                               "row": {}, "refetches": 1})
    assert proc.process_one("2601") == "refetch_requested"
    assert s.state("2601") == "fetching"                              # the dir stays, for the producer
    assert [(r["record_id"], r["refetches"]) for r in s.pending_requests()] == [("2601", 2)]
    s.clear_request("2601")
    # killed after the markers went and before the request: a RETURNED dir without a request
    s.unmark("2602")
    proc.lock.close()
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = _Multi({rid: _FakeZenodo(rid, f) for rid, f in records.items()})
    prod._repair_requests()
    assert sorted((r["record_id"], r["refetches"]) for r in s.pending_requests()) == [("2601", 2), ("2602", 2)]
    assert prod.run() == 0
    assert s.state("2601") == "ready" and s.state("2602") == "ready" and s.pending_requests() == []
    assert s.read_manifest("2602")["refetches"] == 2
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    assert {e["record"] for e in ev if e["event"] == "refetch_request_restored"} == {"2601", "2602"}


def test_error_row_is_written_before_the_dir_goes_and_a_done_record_is_not_redone(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import pipeline, runner
    from envision_eye_actionable.survey.results import load_results
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {"2701": {"a.png": b"x"}, "2702": {"b.pdf": b"y"}, "2703": {"c.pdf": b"z"}})
    cfg = _cfg(tmp_path, meta)
    proc = pipeline.Processor(cfg)
    s = proc.spool
    s.record_dir("2701").mkdir(parents=True)
    s.write_manifest("2701", {"items": [{"path": "dl/a.png", "size": 1, "role": "download", "key": "a.png"}],
                              "row": {}, "refetches": pipeline.MAX_REFETCHES})
    orig_remove = s.remove

    def killed(rid):
        raise KeyboardInterrupt("killed")                   # the kill lands right after the row

    monkeypatch.setattr(s, "remove", killed)
    with pytest.raises(KeyboardInterrupt):
        proc.process_one("2701")
    monkeypatch.setattr(s, "remove", orig_remove)
    assert load_results(cfg.out_dir / "survey_results.jsonl")["2701"]["status"] == "error"
    assert s.state("2701") == "ready"
    n_lines = len((cfg.out_dir / "survey_results.jsonl").read_text().splitlines())
    assert proc.process_one("2701") == "already_done" and s.state("2701") == "absent"
    assert len((cfg.out_dir / "survey_results.jsonl").read_text().splitlines()) == n_lines

    # a normal record: the row is written, the kill hits before the dir goes
    for rid in ("2702", "2703"):
        s.record_dir(rid).mkdir(parents=True)
        s.write_manifest(rid, {"items": [], "terminal": {"status": "no_files", "finish": "without_images"},
                               "row": {}})
    monkeypatch.setattr(s, "remove", killed)
    with pytest.raises(KeyboardInterrupt):
        proc.process_one("2702")
    monkeypatch.setattr(s, "remove", orig_remove)
    n_lines = len((cfg.out_dir / "survey_results.jsonl").read_text().splitlines())
    assert proc.process_one("2702") == "already_done"
    assert len((cfg.out_dir / "survey_results.jsonl").read_text().splitlines()) == n_lines
    assert not (s.record_dir("2702")).exists()
    # a final row older than the marker (an earlier run, --retry-status) does not count
    with open(cfg.out_dir / "survey_results.jsonl", "a", encoding="utf-8") as fp:
        fp.write(json.dumps({"record_id": "2703", "status": "error", "finished_at": "2020-01-01T00:00:00+00:00"})
                 + "\n")
    assert proc.process_one("2703") == "row"


def test_producer_start_removes_partial_dirs_outside_its_scope(tmp_path):
    from envision_eye_actionable.survey import pipeline
    png = _png_bytes(size=(64, 48))
    records = {"2801": {"a.png": png}, "2802": {"b.png": png}}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, consumer_grace_s=0.0, ids=["2801"])
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = _Multi({rid: _FakeZenodo(rid, f) for rid, f in records.items()})
    partial = prod.spool.record_dir("2802")
    (partial / "dl").mkdir(parents=True)
    (partial / "dl" / "half.part").write_bytes(b"\0" * 5000)
    (partial / "manifest.json").write_text("{}", encoding="utf-8")   # killed between manifest and READY
    assert prod.run() == 0
    assert not partial.exists()
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    assert any(e["event"] == "partial_spool_removed" and e["record"] == "2802" for e in ev)


def test_consumer_registers_the_token_for_its_outputs(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import pipeline, runner, zenodo
    monkeypatch.setattr(zenodo, "_SECRETS", set())
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    tok = tmp_path / "tok"
    tok.write_text(SECRET + "\n", encoding="utf-8")
    meta = _write_records(tmp_path, {"2901": {"a.pdf": b"x"}})
    cfg = _cfg(tmp_path, meta, token_file=tok)
    proc = pipeline.Processor(cfg)
    assert zenodo.redact(f"x {SECRET} y") == f"x {zenodo.REDACTED} y"
    s = proc.spool
    s.record_dir("2901").mkdir(parents=True)
    # the producer's manifest carries it (an error text); the spool is
    # scratch, the consumer's outputs must not
    s.write_manifest("2901", {"items": [], "row": {"archive_probe_detail": f"Bearer {SECRET}",
                                                   "size_note": f"proxy said {SECRET}"},
                              "terminal": {"status": "no_files", "finish": "without_images",
                                           "error": f"Bearer {SECRET}"}})
    assert proc.process_one("2901") == "row"
    text = _all_text(cfg.out_dir)
    assert SECRET not in text and zenodo.REDACTED in text


class _Resp:
    def __init__(self, url, status=200, history=(), headers=None):
        self.url, self.status_code, self.history = url, status, list(history)
        self.headers = headers or {}

    def close(self):
        pass


def test_zenodo_link_checks_count_every_hop_and_fallback_get(tmp_path):
    from envision_eye_actionable.survey.weblinks import LinkChecker

    class Zen:
        def __init__(self):
            self.throttles, self.observed = 0, []

        def throttle(self):
            self.throttles += 1

        def observe(self, r):
            self.observed.append(r.url)

        def backoff(self, s):
            pass

    zen = Zen()
    lc = LinkChecker(tmp_path / "links.json", min_interval=0, zenodo=zen)
    doi = "https://doi.org/10.5281/zenodo.42"
    rec = "https://zenodo.org/records/42"
    page = "https://zenodo.org/records/42/files"

    class Sess:
        def head(self, url, **k):
            # doi.org -> zenodo.org/records/42 -> .../files, refused
            return _Resp(page, 405, history=[_Resp(doi, 302), _Resp(rec, 302)])

        def get(self, url, **k):
            return _Resp(page, 200, history=[_Resp(doi, 302), _Resp(rec, 302)])

    lc.session = Sess()
    assert lc.check(doi)["http_status"] == 200
    # HEAD: 3 requests, GET: 3 requests; each one takes a slot, each Zenodo answer is observed
    assert zen.throttles == 6
    assert zen.observed == [doi, rec, page, doi, rec, page]
    # a non-Zenodo link that lands on Zenodo counts only its Zenodo hops
    zen.throttles, zen.observed = 0, []

    class Sess2:
        def head(self, url, **k):
            return _Resp(rec, 200, history=[_Resp("https://example.org/x", 301)])

    lc.session = Sess2()
    lc.check("https://example.org/x")
    assert zen.throttles == 1 and zen.observed == [rec]


def test_yielded_triage_reuses_its_archive_probe(tmp_path):
    from types import SimpleNamespace

    from envision_eye_actionable.survey.fetch import RecordFetcher
    cfg = SimpleNamespace(archive_probe=True, offline=False)
    f = RecordFetcher(cfg, zenodo=None)
    calls = []

    def run_probe(row, rid, files, to_fetch):
        calls.append(rid)
        row.update({"archive_probe_bytes": 123, "n_archives_probed": 1, "archive_probes": ["x"]})
        return "PROBE"

    f._run_probe = run_probe
    rec = {"source_id": "3001"}
    to_fetch = [{"key": "a.tar.gz", "size": 10}]
    row1, row2 = {}, {}
    assert f._probe_archives(row1, "3001", to_fetch, to_fetch, rec) == "PROBE"
    assert f._probe_archives(row2, "3001", to_fetch, to_fetch, rec) == "PROBE"      # the retry after a yield
    assert calls == ["3001"]
    assert row2["archive_probe_bytes"] == 123 and row2["n_archives_probed"] == 1 and row2["archive_probe_reused"]
    # another file set is probed again
    assert f._probe_archives({}, "3001", to_fetch + [{"key": "b.rar", "size": 5}], to_fetch, rec) is not None
    f._probe_archives({}, "3001", to_fetch, to_fetch + [{"key": "b.rar", "size": 5}], rec)
    assert len(calls) == 2


def test_consumer_stop_signal_restores_attempts_and_exits_cleanly(tmp_path, monkeypatch):
    import os
    import signal as _signal
    from envision_eye_actionable.survey import pipeline, runner
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {"3101": {"a.pdf": b"x"}, "3102": {"b.pdf": b"y"}})
    cfg = _cfg(tmp_path, meta, max_attempts=1)
    proc = pipeline.Processor(cfg)
    s = proc.spool
    for rid in ("3101", "3102"):
        s.record_dir(rid).mkdir(parents=True)
        s.write_manifest(rid, {"items": [], "terminal": {"status": "no_files", "finish": "without_images"},
                               "row": {}})
    seen = {}
    orig = proc.survey.finish_record

    def finish(rec, *a, **k):
        # the pipeline stops mid-record: SIGTERM to this (main) thread
        os.kill(os.getpid(), _signal.SIGTERM)
        time.sleep(0.05)
        seen["attempts"] = s.attempts(rec["source_id"])
        return orig(rec, *a, **k)

    monkeypatch.setattr(proc.survey, "finish_record", finish)
    fetch_alive = pipeline.RoleLock(cfg.state_dir, "fetch")      # the producer is up
    before = _signal.getsignal(_signal.SIGTERM)
    try:
        assert proc.run() == 0
    finally:
        fetch_alive.close()
    assert _signal.getsignal(_signal.SIGTERM) is before           # handlers restored
    assert seen["attempts"] == 0                                   # the stop is not a start inside the record
    states = s.states()
    assert list(states.values()).count("ready") == 1                # finished the current one, left the next
    st = json.loads((cfg.state_dir / "process_status.json").read_text())
    assert st["state"] == "stopped"


def test_clean_consumer_exits_do_not_use_up_its_restarts(tmp_path, monkeypatch):
    """fetch crashes more than max_role_restarts / 2 times; process exits
    cleanly each time fetch is gone: process is never given up and the
    pipeline ends with 0 once fetch finishes."""
    import sys
    from envision_eye_actionable.survey import pipeline
    meta = _write_records(tmp_path, {"3201": {"a.pdf": b"x"}})
    cfg = _cfg(tmp_path, meta, offline=True, max_role_restarts=4, restart_backoff_s=0.1, poll_s=0.1)
    script = tmp_path / "roles.py"
    script.write_text(
        "import sys, time\n"
        f"sys.path.insert(0, {str(Path(__file__).parent.parent)!r})\n"
        "from pathlib import Path\n"
        "from envision_eye_actionable.survey import pipeline as pl\n"
        f"state = Path({str(cfg.state_dir)!r})\n"
        "role = sys.argv[1]\n"
        "if role == 'monitor':\n"
        "    sys.exit(0)\n"
        "lock = pl.RoleLock(state, role)\n"
        "if role == 'fetch':\n"
        "    n_file = state / 'fetch_starts'\n"
        "    n = int(n_file.read_text()) if n_file.exists() else 0\n"
        "    n_file.write_text(str(n + 1))\n"
        "    time.sleep(0.8)\n"
        "    sys.exit(1 if n < 4 else 0)\n"
        "gone = 0\n"
        "while gone < 3:\n"
        "    gone = 0 if pl.role_alive(state, 'fetch') else gone + 1\n"
        "    time.sleep(0.1)\n"
        "sys.exit(0)\n", encoding="utf-8")
    monkeypatch.setenv("HOME", str(tmp_path))
    code = pipeline.run_pipeline(cfg, [], 0.5, base_cmd=[sys.executable, str(script)])
    ev = [json.loads(x) for x in (cfg.state_dir / "pipeline_events.jsonl").read_text().splitlines()]
    assert not [e for e in ev if e["event"] == "role_given_up"], ev
    assert sum(1 for e in ev if e["event"] == "role_crashed" and e["role"] == "fetch") == 4
    status = json.loads((cfg.state_dir / "pipeline_status.json").read_text())
    assert status["restarts"].get("process", 0) == 0 and status["clean_restarts"].get("process", 0) >= 1
    assert code == 0


@pytest.mark.skipif(not __import__("sys").platform.startswith("linux"), reason="PR_SET_PDEATHSIG is Linux only")
def test_roles_stop_when_the_supervisor_is_killed(tmp_path):
    """SIGKILL of the pipeline supervisor (no chance to forward a signal):
    its fetch, process and monitor children must not keep running alone."""
    import os
    import signal
    import subprocess
    import sys
    state = tmp_path / "state"
    state.mkdir()
    roles = tmp_path / "roles.py"
    roles.write_text(
        "import os, sys, time\n"
        f"sys.path.insert(0, {str(Path(__file__).parent.parent)!r})\n"
        "from pathlib import Path\n"
        "from envision_eye_actionable.survey import pipeline as pl\n"
        f"state = Path({str(state)!r})\n"
        "role = sys.argv[1]\n"
        "lock = pl.RoleLock(state, role) if role != 'monitor' else None\n"
        "(state / f'{role}.pid').write_text(str(os.getpid()))\n"
        "time.sleep(120)\n", encoding="utf-8")
    driver = tmp_path / "driver.py"
    driver.write_text(
        "import sys, types\n"
        f"sys.path.insert(0, {str(Path(__file__).parent.parent)!r})\n"
        "from pathlib import Path\n"
        "from envision_eye_actionable.survey import pipeline as pl\n"
        f"cfg = types.SimpleNamespace(state_dir=Path({str(state)!r}), poll_s=0.2, max_role_restarts=0,\n"
        "                            restart_backoff_s=0.1)\n"
        f"pl.run_pipeline(cfg, [], 0.5, base_cmd=[sys.executable, {str(roles)!r}])\n", encoding="utf-8")
    sup = subprocess.Popen([sys.executable, str(driver)])
    pids = {}
    try:
        deadline = time.monotonic() + 60
        while len(pids) < 3 and time.monotonic() < deadline:
            for r in ("fetch", "process", "monitor"):
                p = state / f"{r}.pid"
                if r not in pids and p.exists() and p.read_text():
                    pids[r] = int(p.read_text())
            time.sleep(0.1)
        assert len(pids) == 3, pids
        sup.send_signal(signal.SIGKILL)
        sup.wait(10)

        def alive(pid):
            try:
                st = Path(f"/proc/{pid}/stat").read_text()
            except OSError:
                return False
            return st.rsplit(")", 1)[1].split()[0] != "Z"

        deadline = time.monotonic() + 15
        while any(alive(p) for p in pids.values()) and time.monotonic() < deadline:
            time.sleep(0.1)
        assert not [r for r, p in pids.items() if alive(p)], "roles outlived the supervisor"
    finally:
        if sup.poll() is None:
            sup.kill()
        for p in pids.values():
            try:
                os.kill(p, signal.SIGKILL)
            except OSError:
                pass


def test_svg_memory_is_bounded(tmp_path, monkeypatch):
    import base64
    import gzip
    import io
    import tracemalloc

    from envision_eye_actionable.survey import images
    monkeypatch.setattr(images, "SVG_MAX_BYTES", 1 << 20)
    # gzip bomb: 64 MB of zeros that compress to about 64 kB
    buf = io.BytesIO()
    with gzip.GzipFile(fileobj=buf, mode="wb") as g:
        chunk = b"\0" * (1 << 20)
        for _ in range(64):
            g.write(chunk)
    bomb = buf.getvalue()
    tracemalloc.start()
    ld = images.load_frame("vector", ".svgz", bomb, None)
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert ld.error and ld.error.startswith("too_large"), ld.error
    assert peak < 16 << 20, peak                     # never the 64 MB it expands to
    # a large file on disk is refused before it is read
    big = tmp_path / "big.svg"
    big.write_bytes(b"<svg>" + b" " * (2 << 20) + b"</svg>")

    def no_read(self):
        raise AssertionError("read the whole file")

    monkeypatch.setattr(Path, "read_bytes", no_read)
    ld = images.load_frame("vector", ".svg", None, big)
    assert ld.error and ld.error.startswith("too_large"), ld.error
    monkeypatch.undo()
    # an embedded image over the pixel budget is never decoded
    photo = _png_bytes(size=(100, 100))
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="200" '
           f'height="200"><image width="200" height="200" xlink:href="data:image/png;base64,'
           f'{base64.b64encode(photo).decode()}"/></svg>').encode()
    ld = images.load_frame("vector", ".svg", svg, None, max_pixels=5000)
    assert ld.image is None and ld.error.startswith("too_large: embedded image 100x100"), ld.error


# ---------------------------------------------------------------------------
# keep dir: retention of eye-positive records' fetched files (keep.py)
# ---------------------------------------------------------------------------
def _keep_files(root: Path) -> set[str]:
    return {p.relative_to(root).as_posix() for p in Path(root).rglob("*") if p.is_file()}


def test_pipeline_keeps_eye_records_and_deletes_the_rest(tmp_path, monkeypatch):
    """Eye-positive records (deep pass, remote zip; local original) are moved
    into the keep dir with triage and deep files and their manifests, the
    rest is deleted; the local original is only referenced."""
    dark = _png_bytes(size=(64, 48), color=(40, 40, 40))
    bright = _png_bytes(size=(64, 48), color=(220, 220, 220))
    local = _zip_bytes({f"l/{i}.png": dark for i in range(3)})
    records = {
        "3001": {"ir.zip": _zip_bytes({f"ir/{i:03d}.png": dark for i in range(30)})},      # eye: deep pass
        "3002": {"neg.zip": _zip_bytes({f"n/{i:03d}.png": bright for i in range(30)})},  # NEG only
        "3003": {"paper.pdf": b"%PDF-1.4 x"},
        "3004": {"local.zip": local},
    }
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, triage_cap=5, remote_cap=12, keep_dir=tmp_path / "keep", keep_max_gb=10,
               spool_max_gb=1)
    (cfg.downloads_dir / "3004").mkdir(parents=True)
    orig = cfg.downloads_dir / "3004" / "local.zip"
    orig.write_bytes(local)
    mtime = orig.stat().st_mtime_ns
    session = _Multi({rid: _FakeZenodo(rid, files) for rid, files in records.items()})
    out, prod, proc = _run_roles(cfg, session, monkeypatch, onnx=_NegOnnxStub)
    assert out == {"fetch": 0, "process": 0}
    from envision_eye_actionable.survey.results import load_results
    rows = load_results(cfg.out_dir / "survey_results.jsonl")
    kroot = (tmp_path / "keep").resolve()
    r1 = rows["3001"]
    assert r1["status"] == "ok" and r1["sampling_pass"] == "deep" and r1["n_eye_images"] == 12
    assert r1["kept"] is True and r1["kept_reason"] == "" and r1["kept_path"] == str(kroot / "3001")
    kept = _keep_files(kroot / "3001")
    members = {k for k in kept if k.startswith("remote/")}
    deep_members = {k for k in kept if k.startswith("deep/")}
    assert len(members) == 5 and len(deep_members) == 7                 # triage files and deep files
    assert {"manifest.json", "manifest_deep.json", "listings.json", "KEPT.json"} <= kept
    assert not kept & {"READY", "DEEP_READY", "AWAITING_DEEP", ".attempts"}
    assert r1["kept_files"] == 12 and r1["kept_bytes"] == sum(
        (kroot / "3001" / k).stat().st_size for k in members | deep_members)
    marker = json.loads((kroot / "3001" / "KEPT.json").read_text())
    assert marker["files"] == 12 and marker["bytes"] == r1["kept_bytes"]
    # NEG only and nothing to read: not kept, deleted as before
    for rid in ("3002", "3003"):
        assert rows[rid]["kept"] is False and rows[rid]["kept_reason"] == "no eye image"
        assert not (kroot / rid).exists()
    # local original: referenced, never moved or copied
    r4 = rows["3004"]
    assert r4["kept"] is True and r4["kept_files"] == 0 and r4["kept_local_files"] == 1
    assert r4["kept_local_paths"] == [str(orig.resolve())]
    assert json.loads((kroot / "3004" / "KEPT.json").read_text())["local_originals"] == [str(orig.resolve())]
    assert not any(k.endswith(".zip") for k in _keep_files(kroot / "3004"))
    assert orig.read_bytes() == local and orig.stat().st_mtime_ns == mtime
    assert list((cfg.spool_dir / "records").iterdir()) == []
    ks = json.loads((cfg.state_dir / "keep_status.json").read_text())
    assert ks["records"] == 2 and ks["bytes"] == sum(
        p.stat().st_size for p in kroot.rglob("*") if p.is_file() and p.name != "KEPT.json"
        and p.parent != kroot)
    from envision_eye_actionable.survey import pipeline
    snap = pipeline.Monitor(cfg).snapshot()
    assert snap["keep"]["records"] == 2 and snap["keep"]["kept_this_run"] == 2
    assert snap["disk"]["keep_gb"] == ks["bytes_gb"]


def _eye_spool_record(spool, rid: str, n: int = 4, extra: bool = True):
    """A ready spool record of n dark (IR) downloaded PNGs, with the scratch
    and temporary files the keep plan must leave out."""
    d = spool.record_dir(rid)
    (d / "dl").mkdir(parents=True)
    items = []
    for i in range(n):
        data = _png_bytes(size=(64, 48), color=(40, 40, 40))
        (d / "dl" / f"{i}.png").write_bytes(data)
        items.append({"path": f"dl/{i}.png", "size": len(data), "role": "download", "key": f"{i}.png"})
    if extra:
        (d / "work").mkdir()
        (d / "work" / "extracted.png").write_bytes(b"dup")
        (d / "dl" / "half.part").write_bytes(b"x")
    spool.write_manifest(rid, {"items": items, "row": {}})
    return {f"dl/{i}.png" for i in range(n)}


def test_keep_move_survives_a_kill_without_losing_or_duplicating_files(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import keep, pipeline, runner
    from envision_eye_actionable.survey.results import load_results
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {"3101": {"0.png": b"x"}, "3102": {"0.png": b"x"}})
    cfg = _cfg(tmp_path, meta, keep_dir=tmp_path / "keep", spool_max_gb=0.001)
    proc = pipeline.Processor(cfg)
    s = proc.spool
    kroot = (tmp_path / "keep").resolve()
    files = _eye_spool_record(s, "3101")
    real = keep._rename
    calls = {"n": 0}

    def killed_after_two(a, b):
        if calls["n"] == 2:
            raise KeyboardInterrupt("killed mid-move")
        calls["n"] += 1
        return real(a, b)

    monkeypatch.setattr(keep, "_rename", killed_after_two)
    with pytest.raises(KeyboardInterrupt):
        proc.process_one("3101")
    monkeypatch.setattr(keep, "_rename", real)
    row = load_results(cfg.out_dir / "survey_results.jsonl")["3101"]
    assert row["status"] == "ok" and row["kept"] is True and row["kept_files"] == 4
    in_keep = {k for k in _keep_files(kroot / "3101") if k.startswith("dl/")}
    in_spool = {k for k in _keep_files(s.record_dir("3101")) if k in files}
    assert in_keep | in_spool == files and not in_keep & in_spool       # nothing lost, nothing twice
    assert len(in_keep) == 2 and not (kroot / "3101" / "KEPT.json").exists()
    assert s.state("3101") == "ready"
    n_lines = len((cfg.out_dir / "survey_results.jsonl").read_text().splitlines())
    # restart (a fresh results tail): the row is newer than the marker and
    # says kept: the move is finished, the record is not classified again
    proc.results = pipeline.ResultsTail(proc.survey.results_path)
    assert proc.process_one("3101") == "already_done"
    assert len((cfg.out_dir / "survey_results.jsonl").read_text().splitlines()) == n_lines
    got = _keep_files(kroot / "3101")
    assert {k for k in got if k.startswith("dl/")} == files               # the .part file stayed out
    assert "work/extracted.png" not in got and "READY" not in got
    assert json.loads((kroot / "3101" / "KEPT.json").read_text())["files"] == 4
    assert s.state("3101") == "absent"
    # a kill after every file moved but before the spool dir went: the
    # marker is written again, nothing moves twice
    _eye_spool_record(s, "3102", extra=False)
    orig_remove = s.remove

    def killed(rid):
        raise KeyboardInterrupt("killed")

    monkeypatch.setattr(s, "remove", killed)
    with pytest.raises(KeyboardInterrupt):
        proc.process_one("3102")
    monkeypatch.setattr(s, "remove", orig_remove)
    assert (kroot / "3102" / "KEPT.json").exists()
    assert proc.process_one("3102") == "already_done" and s.state("3102") == "absent"
    assert json.loads((kroot / "3102" / "KEPT.json").read_text())["files"] == 4


def _release_processor(proc):
    """What a process role exit frees: its role lock and the survey's out and
    scratch dir locks (a restart in the same test process)."""
    proc.lock.close()
    for fp in getattr(proc.survey, "_locks", None) or []:
        if fp is not None:
            fp.close()
    disk = getattr(proc.survey, "disk", None)
    if disk is not None:
        disk.close()


def test_keep_move_failures_park_the_record_and_never_delete_its_files(tmp_path, monkeypatch):
    """A move that keeps failing: backoff between tries, then the record is
    parked (keep_failed) with its unmoved files; the files that did move are
    counted; the next consumer start finishes the move."""
    from envision_eye_actionable.survey import keep, pipeline, runner
    from envision_eye_actionable.survey.results import load_results
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {"3151": {"0.png": b"x"}, "3152": {"0.png": b"x"}})
    cfg = _cfg(tmp_path, meta, keep_dir=tmp_path / "keep", spool_max_gb=0.001)
    proc = pipeline.Processor(cfg)
    s = proc.spool
    kroot = (tmp_path / "keep").resolve()
    files = _eye_spool_record(s, "3151")
    real = keep._rename
    calls = {"n": 0}

    def eio_after_one(a, b):
        calls["n"] += 1
        if calls["n"] == 1:
            return real(a, b)
        raise OSError(5, "Input/output error")

    sleeps = []
    monkeypatch.setattr(pipeline.time, "sleep", lambda t: sleeps.append(t))
    monkeypatch.setattr(keep, "_rename", eio_after_one)
    assert proc.process_one("3151") == "row"
    row = load_results(cfg.out_dir / "survey_results.jsonl")["3151"]
    assert row["kept"] is True
    n_lines = len((cfg.out_dir / "survey_results.jsonl").read_text().splitlines())
    assert s.state("3151") == "ready"
    assert proc.process_one("3151") == "keep_move_retry"
    assert proc.process_one("3151") == "keep_move_retry"            # third failure: parked
    assert sleeps == [2.0, 4.0]
    assert s.state("3151") == "keep_failed" and "3151" not in s.ready_records()
    assert proc.process_one("3151") == "skipped"
    in_keep = {k for k in _keep_files(kroot / "3151") if k.startswith("dl/")}
    in_spool = {k for k in _keep_files(s.record_dir("3151")) if k in files}
    assert len(in_keep) == 1 and in_keep | in_spool == files and not in_keep & in_spool   # nothing deleted
    moved = (kroot / "3151" / next(iter(in_keep))).stat().st_size
    assert proc.keeper.per_record["3151"] == moved and proc.keeper.total_bytes == moved
    assert json.loads((cfg.state_dir / "keep_status.json").read_text())["bytes"] == moved
    ev = [json.loads(x) for x in (cfg.out_dir / "survey_events.jsonl").read_text().splitlines()]
    assert [e["event"] for e in ev].count("keep_move_failed") == 3
    assert [e for e in ev if e["event"] == "keep_move_gave_up"][0]["files_left_in_spool"] == len(in_spool) + 1
    # the producer never refetches or deletes a parked record
    prod = pipeline.Producer.__new__(pipeline.Producer)
    prod.spool = s
    assert prod._triage({"source_id": "3151"}) == "skipped_keep_failed"
    # the next consumer start (rename works again) finishes the move
    monkeypatch.setattr(keep, "_rename", real)
    _release_processor(proc)
    proc2 = pipeline.Processor(cfg)
    assert s.state("3151") == "ready"
    assert proc2.process_one("3151") == "already_done"
    assert {k for k in _keep_files(kroot / "3151") if k.startswith("dl/")} == files
    assert json.loads((kroot / "3151" / "KEPT.json").read_text())["files"] == 4
    assert s.state("3151") == "absent"
    assert len((cfg.out_dir / "survey_results.jsonl").read_text().splitlines()) == n_lines
    # a consumer without --keep-dir parks a record whose move a kill cut short
    files2 = _eye_spool_record(s, "3152", extra=False)
    calls["n"] = 0

    def killed_after_one(a, b):
        calls["n"] += 1
        if calls["n"] > 1:
            raise KeyboardInterrupt("killed mid-move")
        return real(a, b)

    monkeypatch.setattr(keep, "_rename", killed_after_one)
    with pytest.raises(KeyboardInterrupt):
        proc2.process_one("3152")
    monkeypatch.setattr(keep, "_rename", real)
    _release_processor(proc2)
    cfg.keep_dir = None
    proc3 = pipeline.Processor(cfg)
    assert proc3.keeper is None
    assert proc3.process_one("3152") == "keep_parked" and s.state("3152") == "keep_failed"
    assert {k for k in _keep_files(s.record_dir("3152")) if k in files2} | {
        k for k in _keep_files(kroot / "3152") if k.startswith("dl/")} == files2
    _release_processor(proc3)
    cfg.keep_dir = tmp_path / "keep"
    proc4 = pipeline.Processor(cfg)
    assert proc4.process_one("3152") == "already_done" and s.state("3152") == "absent"
    assert {k for k in _keep_files(kroot / "3152") if k.startswith("dl/")} == files2


def test_keep_guards_refuse_without_deleting_kept_records(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import keep, pipeline, runner
    from envision_eye_actionable.survey.results import load_results
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {r: {"0.png": b"x"} for r in ("3201", "3202", "3203")})
    kdir = tmp_path / "keep"
    kdir.mkdir()
    (kdir / keep.KEEP_MARKER).write_text("x")
    (kdir / "3299").mkdir()
    (kdir / "3299" / "old.png").write_bytes(b"\0" * 5000)          # kept earlier, no marker: measured
    cfg = _cfg(tmp_path, meta, keep_dir=kdir, keep_max_gb=6e-6, spool_max_gb=0.001)       # 6 kB
    proc = pipeline.Processor(cfg)
    assert proc.keeper.total_bytes == 5000
    _eye_spool_record(proc.spool, "3201", extra=False)
    assert proc.process_one("3201") == "row"
    row = load_results(cfg.out_dir / "survey_results.jsonl")["3201"]
    assert row["kept"] is False and "--keep-max-gb" in row["kept_reason"] and row["kept_files"] == 0
    assert row["n_eye_images"] == 4
    assert not (kdir / "3201").exists() and proc.spool.state("3201") == "absent"
    assert (kdir / "3299" / "old.png").stat().st_size == 5000          # never deleted
    # the disk guard: free space with an empty spool under floor + spool max
    proc.keeper.max_bytes = None
    monkeypatch.setattr(proc.keeper, "_free_if_spool_empty", lambda need: 10)
    proc.keeper.floor_bytes, proc.keeper.spool_max_bytes = 5, 6
    _eye_spool_record(proc.spool, "3202", extra=False)
    assert proc.process_one("3202") == "row"
    row = load_results(cfg.out_dir / "survey_results.jsonl")["3202"]
    assert row["kept"] is False and row["kept_reason"].startswith("disk:")
    monkeypatch.setattr(proc.keeper, "_free_if_spool_empty", lambda need: 11)
    _eye_spool_record(proc.spool, "3203", extra=False)
    assert proc.process_one("3203") == "row"
    assert load_results(cfg.out_dir / "survey_results.jsonl")["3203"]["kept"] is True
    ev = [json.loads(x)["event"] for x in (cfg.out_dir / "survey_events.jsonl").read_text().splitlines()]
    assert ev.count("keep_refused") == 2 and ev.count("kept") == 1


def test_keep_dir_must_be_its_own_and_on_the_spool_filesystem(tmp_path, monkeypatch):
    from envision_eye_actionable.survey import keep
    from envision_eye_actionable.survey.spool import Spool
    (tmp_path / "dl").mkdir()
    spool = Spool(tmp_path / "spool", tmp_path / "dl", tmp_path / "out")
    for bad in (tmp_path / "spool" / "k", tmp_path / "dl" / "k", tmp_path / "out", tmp_path):
        with pytest.raises(ValueError, match="overlaps"):
            keep.Keeper(bad, spool, tmp_path / "dl", tmp_path / "out")
    # the scratch dir (whose records/ the survey sweeps of all-digit dirs)
    # and the state dir
    for bad in (tmp_path / "scratch" / "records", tmp_path / "scratch", tmp_path / "state" / "k"):
        with pytest.raises(ValueError, match="overlaps (scratch|state) dir"):
            keep.Keeper(bad, spool, tmp_path / "dl", tmp_path / "out", scratch_dir=tmp_path / "scratch",
                        state_dir=tmp_path / "state")
    assert not (tmp_path / "scratch").exists()
    # the process role passes its scratch and state dirs
    from envision_eye_actionable.survey import pipeline, runner
    monkeypatch.setattr(runner, "OnnxClassifier", _OnnxStub)
    meta = _write_records(tmp_path, {"3301": {"0.png": b"x"}})
    for i, (bad, what) in enumerate((("scratch/records", "scratch"), ("st1/keep", "state"))):
        cfg = _cfg(tmp_path, meta, keep_dir=tmp_path / bad, spool_dir=tmp_path / f"spool{i}",
                   state_dir=tmp_path / f"st{i}")
        with pytest.raises(ValueError, match=f"overlaps {what} dir"):
            pipeline.Processor(cfg)
    foreign = tmp_path / "data"
    foreign.mkdir()
    (foreign / "x.txt").write_text("mine")
    with pytest.raises(ValueError, match="not created by the survey"):
        keep.Keeper(foreign, spool)
    real_dev = keep._st_dev
    target = (tmp_path / "k2").resolve()
    monkeypatch.setattr(keep, "_st_dev", lambda p: real_dev(p) + (Path(p) == target))
    with pytest.raises(ValueError, match="filesystem"):
        keep.Keeper(tmp_path / "k2", spool)
    monkeypatch.setattr(keep, "_st_dev", real_dev)
    assert keep.Keeper(tmp_path / "k3", spool).total_bytes == 0


def test_keep_eye_count_uses_thresholded_eye_labels_only():
    from envision_eye_actionable.survey import keep
    assert keep.n_eye_images({"n_IR": 2, "n_MASK": 9, "n_NEG": 4, "n_UNCERTAIN": 3, "n_OCTA": 1}) == 3
    assert keep.n_eye_images({"n_MASK": 5, "n_NEG": 1, "argmax_IR": 4}) == 0
    assert keep.n_eye_images({f"n_{c}": 1 for c in ("CFP", "IR", "PSC", "FAF", "OCT", "OCTA")}) == 6


def test_keep_plan_leaves_out_markers_scratch_and_temporaries(tmp_path):
    from envision_eye_actionable.survey import keep
    d = tmp_path / "123"
    for rel in ("READY", "DEEP_READY", "AWAITING_DEEP", ".attempts", "RETURNED", "manifest.json",
                "manifest.json.77.tmp", "dl/a.zip", "dl/b.png.part", "work/x.png", ".fetchwalk/y",
                "remote/m/1.png", "deep/dl/c.tif", "deep/work/keep.png"):
        (d / rel).parent.mkdir(parents=True, exist_ok=True)
        (d / rel).write_bytes(b"1")
    assert keep.keep_plan(d) == ["deep/dl/c.tif", "deep/work/keep.png", "dl/a.zip", "manifest.json",
                                 "remote/m/1.png"]


def test_backfill_ids_and_refetch_keep_ids(tmp_path, monkeypatch, capsys):
    import gzip
    from envision_eye_actionable.survey import cli, pipeline
    png = _png_bytes(size=(64, 48))
    rids = ["3301", "3302", "3303", "3304", "3305", "3306"]
    records = {r: {f"{r}.png": png} for r in rids}
    meta = _write_records(tmp_path, records)
    cfg = _cfg(tmp_path, meta, consumer_grace_s=0.0, fetch_workers=1, order="scrape")
    cfg.out_dir.mkdir(parents=True)
    rows = [{"record_id": "3301", "status": "ok", "n_IR": 3, "n_classified": 3},        # pre-retention eye
            {"record_id": "3302", "status": "no_eye_images", "n_NEG": 5, "n_classified": 5},
            {"record_id": "3303", "status": "ok", "n_CFP": 2, "n_classified": 2, "kept": True},
            {"record_id": "3304", "status": "mask_dominated", "n_classified": 2},        # eye only in predictions
            {"record_id": "3305", "status": "ok", "n_IR": 1, "n_classified": 1},
            {"record_id": "3305", "status": "started"},                                 # in progress: not final
            {"record_id": "3306", "status": "ok", "n_OCT": 4, "n_classified": 4},      # already kept
            {"record_id": "3307", "status": "ok", "n_OCT": 4, "n_classified": 4, "spool_bytes": 0}]  # local only
    (cfg.out_dir / "survey_results.jsonl").write_text("".join(json.dumps(r) + "\n" for r in rows), encoding="utf-8")
    (cfg.out_dir / "predictions").mkdir()
    with gzip.open(cfg.out_dir / "predictions" / "3304.jsonl.gz", "wt") as fp:
        fp.write(json.dumps({"path": "a", "label": "OCT"}) + "\n" + json.dumps({"path": "b", "label": "MASK"}) + "\n")
    with gzip.open(cfg.out_dir / "predictions" / "3302.jsonl.gz", "wt") as fp:
        fp.write(json.dumps({"path": "a", "label": "NEG"}) + "\n")
    kdir = tmp_path / "keep"
    (kdir / "3306").mkdir(parents=True)
    (kdir / "3306" / "KEPT.json").write_text("{}")
    ids_file = tmp_path / "state" / "backfill_keep_ids.txt"
    assert cli.main(["keep-backfill-ids", "--out-dir", str(cfg.out_dir), "--keep-dir", str(kdir),
                     "--output", str(ids_file)]) == 0
    assert ids_file.read_text().split() == ["3301", "3304"]
    summary = json.loads(capsys.readouterr().out)
    assert summary["eye_by_rows"] == 2 and summary["predictions_only"] == ["3304"] and summary["already_kept"] == 1
    assert summary["eye_local_only_not_listed"] == ["3307"]
    # the producer redoes the listed pre-retention records (first), never a
    # record whose row was written with retention, plus the unfinished 3305
    cfg.refetch_ids = ["3301", "3303", "3304"]
    prod = pipeline.Producer(cfg)
    prod.zenodo.session = _Multi({rid: _FakeZenodo(rid, f) for rid, f in records.items()})
    assert prod.run() == 0
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    fetched = [e["record"] for e in ev if e["event"] == "fetched"]
    assert fetched == ["3301", "3304", "3305"]
    assert next(e for e in ev if e["event"] == "fetch_start")["n_refetch_keep"] == 2
    # the refetched rows replace the old ones (last line wins), and a row
    # written with retention is not redone on the next start
    with open(cfg.out_dir / "survey_results.jsonl", "a", encoding="utf-8") as fp:
        for rid in ("3301", "3304", "3305"):
            fp.write(json.dumps({"record_id": rid, "status": "ok", "n_IR": 1, "kept": True}) + "\n")
    from envision_eye_actionable.survey.excel import merge_index
    assert merge_index([cfg.out_dir / "survey_results.jsonl"])["3304"][2] == "ok"
    for rid in ("3301", "3304", "3305"):
        prod.spool.remove(rid)
    prod2 = pipeline.Producer(cfg)
    prod2.zenodo.session = _Multi({rid: _FakeZenodo(rid, f) for rid, f in records.items()})
    assert prod2.run() == 0
    ev = [json.loads(x) for x in (cfg.state_dir / "fetch_events.jsonl").read_text().splitlines()]
    starts = [e for e in ev if e["event"] == "fetch_start"]
    assert starts[-1]["n_refetch_keep"] == 0 and starts[-1]["n_todo"] == 0


def test_cli_keep_options(tmp_path):
    from envision_eye_actionable.survey import cli, excel
    keys = {k for k, _ in excel.RECORD_COLUMNS}
    assert {"kept", "kept_reason", "kept_path", "kept_files", "kept_bytes", "n_eye_images"} <= keys
    model = tmp_path / "m.onnx"
    model.write_bytes(b"x")
    with pytest.raises(SystemExit):
        cli.main(["run", "--model", str(model), "--keep-dir", str(tmp_path / "k")])
    with pytest.raises(SystemExit):
        cli.main(["fetch", "--model", str(model), "--spool-dir", str(tmp_path / "s"), "--state-dir",
                  str(tmp_path / "st"), "--refetch-keep-ids", str(tmp_path / "missing.txt")])
